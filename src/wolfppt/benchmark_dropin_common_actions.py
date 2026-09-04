"""Common drop-in benchmark edit actions."""

from __future__ import annotations

from datetime import datetime, timezone
from typing import Any

from .benchmark_cases import (
    CREATE_NOTES_SLIDE_EXPECTED,
    MIXED_WORKLOAD_EXPECTED,
    NOTES_BACKGROUND_EXPECTED_RGB,
    NOTES_PLACEHOLDER_HYPERLINK_EXPECTED_ADDRESS,
    NOTES_RUN_HYPERLINK_EXPECTED_ADDRESS,
    NOTES_TEXT_EXPECTED,
    SLIDE_BACKGROUND_EXPECTED_RGB,
    SLIDE_SIZE_EXPECTED_CX,
    SLIDE_SIZE_EXPECTED_CY,
    SLIDE_NAME_EXPECTED,
    TEMPLATE_BACKGROUND_EXPECTED_RGB,
)
from .presentation import Presentation as WolfPresentation
from .shape_core_facade import Shape
from .text_facade_inspection import _shape_run_count_payload


def _apply_python_pptx_dropin_edit(fixture_id: str, prs: Any) -> None:
    if fixture_id == "text_basic/title_body_bullets":
        prs.slides[0].shapes[1].text_frame.paragraphs[1].text = "Benchmark paragraph"
        return
    if fixture_id in _MIXED_WORKLOAD_FIXTURES:
        prs.slides[0].shapes[1].text = "Benchmark paragraph"
        return
    if fixture_id == "tables/simple_table":
        prs.slides[0].shapes[0].table.cell(1, 1).text = "2"
        return
    raise RuntimeError(f"drop-in edit benchmark does not support fixture {fixture_id}")


def _apply_wolfppt_dropin_edit(fixture_id: str, prs: WolfPresentation) -> None:
    if fixture_id == "text_basic/title_body_bullets":
        _shape_at(prs, 0, 1).text_frame.paragraphs[1].text = "Benchmark paragraph"
        return
    if fixture_id in _MIXED_WORKLOAD_FIXTURES:
        _shape_at(prs, 0, 1).text = "Benchmark paragraph"
        return
    if fixture_id == "tables/simple_table":
        _shape_at(prs, 0, 0).table.cell(1, 1).text = "2"
        return
    raise RuntimeError(f"drop-in edit benchmark does not support fixture {fixture_id}")


def _apply_python_pptx_dropin_multi_edit(fixture_id: str, prs: Any) -> None:
    if fixture_id == "text_basic/title_body_bullets":
        prs.slides[0].shapes[0].text = "WolfPPT Native"
        prs.slides[0].shapes[1].text_frame.paragraphs[1].text = "Second"
        return
    if fixture_id == "tables/simple_table":
        table = prs.slides[0].shapes[0].table
        table.cell(1, 0).text = "Rows"
        table.cell(1, 1).text = "2"
        return
    if fixture_id == "workloads/multi_edit_table":
        table = prs.slides[0].shapes[0].table
        for row_idx in range(1, 6):
            for col_idx in range(6):
                table.cell(row_idx, col_idx).text = f"E{row_idx}-{col_idx}"
        return
    raise RuntimeError(f"drop-in multi-edit benchmark does not support fixture {fixture_id}")


def _apply_wolfppt_dropin_multi_edit(fixture_id: str, prs: WolfPresentation) -> None:
    if fixture_id == "text_basic/title_body_bullets":
        _shape_at(prs, 0, 0).text = "WolfPPT Native"
        _shape_at(prs, 0, 1).text_frame.paragraphs[1].text = "Second"
        return
    if fixture_id == "tables/simple_table":
        table = _shape_at(prs, 0, 0).table
        table.cell(1, 0).text = "Rows"
        table.cell(1, 1).text = "2"
        return
    if fixture_id == "workloads/multi_edit_table":
        table = _shape_at(prs, 0, 0).table
        for row_idx in range(1, 6):
            for col_idx in range(6):
                table.cell(row_idx, col_idx).text = f"E{row_idx}-{col_idx}"
        return
    raise RuntimeError(f"drop-in multi-edit benchmark does not support fixture {fixture_id}")


def _apply_python_pptx_dropin_slide_name_edit(fixture_id: str, prs: Any) -> None:
    if fixture_id not in _REAL_WORLD_SAFE_TEXT_FIXTURES:
        raise RuntimeError(
            f"drop-in slide name benchmark does not support fixture {fixture_id}"
        )
    prs.slides[0].name = SLIDE_NAME_EXPECTED


def _apply_wolfppt_dropin_slide_name_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in _REAL_WORLD_SAFE_TEXT_FIXTURES:
        raise RuntimeError(
            f"drop-in slide name benchmark does not support fixture {fixture_id}"
        )
    prs.slides[0].name = SLIDE_NAME_EXPECTED


def _apply_python_pptx_dropin_slide_size_edit(fixture_id: str, prs: Any) -> None:
    if fixture_id not in _REAL_WORLD_SAFE_TEXT_FIXTURES:
        raise RuntimeError(
            f"drop-in slide size benchmark does not support fixture {fixture_id}"
        )
    prs.slide_width = SLIDE_SIZE_EXPECTED_CX
    prs.slide_height = SLIDE_SIZE_EXPECTED_CY


def _apply_python_pptx_dropin_slide_layout_remove(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id != "slides/two_slide_text":
        raise RuntimeError(
            f"drop-in slide-layout remove benchmark does not support fixture {fixture_id}"
        )
    layout = prs.slide_layouts.get_by_name("Title Slide")
    if layout is None:
        raise RuntimeError("Title Slide layout is unavailable")
    prs.slide_layouts.remove(layout)


def _apply_wolfppt_dropin_slide_size_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in _REAL_WORLD_SAFE_TEXT_FIXTURES:
        raise RuntimeError(
            f"drop-in slide size benchmark does not support fixture {fixture_id}"
        )
    prs.slide_width = SLIDE_SIZE_EXPECTED_CX
    prs.slide_height = SLIDE_SIZE_EXPECTED_CY


def _apply_wolfppt_dropin_slide_layout_remove(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id != "slides/two_slide_text":
        raise RuntimeError(
            f"drop-in slide-layout remove benchmark does not support fixture {fixture_id}"
        )
    layout = prs.slide_layouts.get_by_name("Title Slide")
    if layout is None:
        raise RuntimeError("Title Slide layout is unavailable")
    prs.slide_layouts.remove(layout)


def _apply_python_pptx_dropin_slide_background_edit(fixture_id: str, prs: Any) -> None:
    if fixture_id not in _REAL_WORLD_SAFE_TEXT_FIXTURES:
        raise RuntimeError(
            "drop-in slide background benchmark does not support fixture "
            f"{fixture_id}"
        )
    from pptx.dml.color import RGBColor

    prs.slides[0].background.fill.solid()
    prs.slides[0].background.fill.fore_color.rgb = RGBColor.from_string(
        SLIDE_BACKGROUND_EXPECTED_RGB
    )


def _apply_wolfppt_dropin_slide_background_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in _REAL_WORLD_SAFE_TEXT_FIXTURES:
        raise RuntimeError(
            "drop-in slide background benchmark does not support fixture "
            f"{fixture_id}"
        )
    prs.slides[0].background.fill.solid()
    prs.slides[0].background.fill.fore_color.rgb = SLIDE_BACKGROUND_EXPECTED_RGB


def _apply_python_pptx_dropin_notes_background_edit(fixture_id: str, prs: Any) -> None:
    if fixture_id not in _NOTES_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in notes background benchmark does not support fixture "
            f"{fixture_id}"
        )
    from pptx.dml.color import RGBColor

    prs.slides[0].notes_slide.background.fill.solid()
    prs.slides[0].notes_slide.background.fill.fore_color.rgb = RGBColor.from_string(
        NOTES_BACKGROUND_EXPECTED_RGB
    )


def _apply_wolfppt_dropin_notes_background_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in _NOTES_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in notes background benchmark does not support fixture "
            f"{fixture_id}"
        )
    prs.slides[0].notes_slide.background.fill.solid()
    prs.slides[0].notes_slide.background.fill.fore_color.rgb = (
        NOTES_BACKGROUND_EXPECTED_RGB
    )


def _apply_python_pptx_dropin_template_background_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in _REAL_WORLD_SAFE_TEXT_FIXTURES:
        raise RuntimeError(
            "drop-in template background benchmark does not support fixture "
            f"{fixture_id}"
        )
    from pptx.dml.color import RGBColor

    color = RGBColor.from_string(TEMPLATE_BACKGROUND_EXPECTED_RGB)
    prs.slide_layouts[0].background.fill.solid()
    prs.slide_layouts[0].background.fill.fore_color.rgb = color
    prs.slide_master.background.fill.solid()
    prs.slide_master.background.fill.fore_color.rgb = color


def _apply_wolfppt_dropin_template_background_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in _REAL_WORLD_SAFE_TEXT_FIXTURES:
        raise RuntimeError(
            "drop-in template background benchmark does not support fixture "
            f"{fixture_id}"
        )
    prs.slide_layouts[0].background.fill.solid()
    prs.slide_layouts[0].background.fill.fore_color.rgb = (
        TEMPLATE_BACKGROUND_EXPECTED_RGB
    )
    prs.slide_master.background.fill.solid()
    prs.slide_master.background.fill.fore_color.rgb = TEMPLATE_BACKGROUND_EXPECTED_RGB


def _apply_python_pptx_dropin_core_properties_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in _REAL_WORLD_SAFE_TEXT_FIXTURES:
        raise RuntimeError(
            "drop-in core properties benchmark does not support fixture "
            f"{fixture_id}"
        )
    _apply_core_properties(prs)


def _apply_wolfppt_dropin_core_properties_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in _REAL_WORLD_SAFE_TEXT_FIXTURES:
        raise RuntimeError(
            "drop-in core properties benchmark does not support fixture "
            f"{fixture_id}"
        )
    _apply_core_properties(prs)


def _apply_core_properties(prs: Any) -> None:
    core = prs.core_properties
    core.title = "Quarterly Review"
    core.subject = "Board"
    core.author = "Wolf"
    core.keywords = "qoe,deck"
    core.comments = "Updated by WolfPPT"
    core.last_modified_by = "Wolf"
    core.category = "Finance"
    core.content_status = "Draft"
    core.identifier = "deck-001"
    core.language = "en-US"
    core.version = "v1"
    core.revision = 7
    core.created = datetime(2020, 1, 2, 3, 4, 5, tzinfo=timezone.utc)
    core.modified = datetime(2020, 2, 3, 4, 5, 6)
    core.last_printed = datetime(2020, 3, 4, 5, 6, 7)


def _apply_python_pptx_dropin_mixed_workload_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in _MIXED_WORKLOAD_FIXTURES:
        raise RuntimeError(
            f"drop-in mixed workload benchmark does not support fixture {fixture_id}"
        )
    _apply_mixed_workload_edit(prs)


def _apply_wolfppt_dropin_mixed_workload_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in _MIXED_WORKLOAD_FIXTURES:
        raise RuntimeError(
            f"drop-in mixed workload benchmark does not support fixture {fixture_id}"
        )
    _apply_mixed_workload_edit(prs)


def _apply_mixed_workload_edit(prs: Any) -> None:
    from pptx.chart.data import CategoryChartData

    expected = MIXED_WORKLOAD_EXPECTED
    _shape_at(prs, 0, 0).text = expected["title"]
    _shape_at(prs, 1, 1).table.cell(1, 3).text = expected["table_revenue_q3"]
    chart_data = CategoryChartData()
    chart_data.categories = expected["chart_categories"]
    for series in expected["chart_series"]:
        chart_data.add_series(series["name"], tuple(series["values"]))
    _shape_at(prs, 2, 1).chart.replace_data(chart_data)
    _shape_at(prs, 3, 1).text = expected["risk"]
    _shape_at(prs, 4, 1).text = "\n".join(expected["actions"])
    prs.slides[4].notes_slide.notes_text_frame.text = expected["notes"]


def _shape_at(prs: Any, slide_index: int, shape_index: int) -> Any:
    slide = prs.slides[slide_index]
    payload = getattr(slide, "_payload", None)
    if isinstance(payload, dict):
        hints = payload.get("shape_shell_hints")
        if isinstance(hints, list):
            table_index = 0
            run_start_index = 0
            for hint in hints:
                if not isinstance(hint, dict):
                    continue
                hint_index = int(hint.get("_shape_index", -1))
                if hint_index == shape_index:
                    return Shape(
                        slide,
                        shape_index,
                        hint,
                        table_index=int(hint.get("_table_index", table_index)),
                        run_start_index=run_start_index,
                    )
                table_index += len(hint.get("tables") or [])
                run_start_index += _shape_run_count_payload(hint)
    return slide.shapes[shape_index]


def _shape_collection_at(prs: Any, slide_index: int) -> Any:
    slide = prs.slides[slide_index]
    payload = getattr(slide, "_payload", None)
    if isinstance(prs, WolfPresentation) and isinstance(payload, dict):
        hints = payload.get("shape_shell_hints")
        if isinstance(hints, list):
            from .shape_facade import ShapeCollection

            shapes: list[Shape] = []
            table_index = 0
            run_start_index = 0
            for hint in hints:
                if not isinstance(hint, dict):
                    continue
                shape_index = int(hint.get("_shape_index", len(shapes)))
                shapes.append(
                    Shape(
                        slide,
                        shape_index,
                        hint,
                        table_index=int(hint.get("_table_index", table_index)),
                        run_start_index=run_start_index,
                    )
                )
                table_index += len(hint.get("tables") or [])
                run_start_index += _shape_run_count_payload(hint)
            return ShapeCollection(slide, shapes)
    return slide.shapes


def _apply_python_pptx_dropin_notes_text_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in _NOTES_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in notes text benchmark does not support fixture {fixture_id}"
        )
    prs.slides[0].notes_slide.notes_text_frame.text = NOTES_TEXT_EXPECTED


def _apply_python_pptx_dropin_create_notes_slide(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id != "text_basic/title_body_bullets":
        raise RuntimeError(
            f"drop-in create notes-slide benchmark does not support fixture {fixture_id}"
        )
    prs.slides[0].notes_slide.notes_text_frame.text = CREATE_NOTES_SLIDE_EXPECTED


def _apply_wolfppt_dropin_create_notes_slide(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id != "text_basic/title_body_bullets":
        raise RuntimeError(
            f"drop-in create notes-slide benchmark does not support fixture {fixture_id}"
        )
    prs.slides[0].notes_slide.notes_text_frame.text = CREATE_NOTES_SLIDE_EXPECTED


def _apply_wolfppt_dropin_notes_text_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in _NOTES_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in notes text benchmark does not support fixture {fixture_id}"
        )
    prs.slides[0].notes_slide.notes_text_frame.text = NOTES_TEXT_EXPECTED


def _apply_python_pptx_dropin_notes_text_frame_flow_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE

    if fixture_id not in _NOTES_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in notes text-frame flow benchmark does not support fixture "
            f"{fixture_id}"
        )
    frame = prs.slides[0].notes_slide.notes_text_frame
    _apply_notes_text_frame_flow(frame, MSO_ANCHOR, MSO_AUTO_SIZE)


def _apply_wolfppt_dropin_notes_text_frame_flow_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE

    if fixture_id not in _NOTES_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in notes text-frame flow benchmark does not support fixture "
            f"{fixture_id}"
        )
    frame = prs.slides[0].notes_slide.notes_text_frame
    _apply_notes_text_frame_flow(frame, MSO_ANCHOR, MSO_AUTO_SIZE)


def _apply_notes_text_frame_flow(
    frame: Any,
    anchor_enum: Any,
    auto_size_enum: Any,
) -> None:
    from .benchmark_cases import NOTES_TEXT_FRAME_FLOW_EXPECTED

    frame.margin_left = NOTES_TEXT_FRAME_FLOW_EXPECTED["margin_left"]
    frame.margin_right = NOTES_TEXT_FRAME_FLOW_EXPECTED["margin_right"]
    frame.margin_top = NOTES_TEXT_FRAME_FLOW_EXPECTED["margin_top"]
    frame.margin_bottom = NOTES_TEXT_FRAME_FLOW_EXPECTED["margin_bottom"]
    frame.word_wrap = NOTES_TEXT_FRAME_FLOW_EXPECTED["word_wrap"]
    frame.vertical_anchor = anchor_enum.MIDDLE
    frame.auto_size = auto_size_enum.TEXT_TO_FIT_SHAPE


def _apply_python_pptx_dropin_notes_hyperlink_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in _NOTES_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in notes hyperlink benchmark does not support fixture "
            f"{fixture_id}"
        )
    notes = prs.slides[0].notes_slide
    notes.notes_placeholder.click_action.hyperlink.address = (
        NOTES_PLACEHOLDER_HYPERLINK_EXPECTED_ADDRESS
    )
    notes.notes_text_frame.paragraphs[0].runs[0].hyperlink.address = (
        NOTES_RUN_HYPERLINK_EXPECTED_ADDRESS
    )


def _apply_wolfppt_dropin_notes_hyperlink_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in _NOTES_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in notes hyperlink benchmark does not support fixture "
            f"{fixture_id}"
        )
    notes = prs.slides[0].notes_slide
    notes.notes_placeholder.click_action.hyperlink.address = (
        NOTES_PLACEHOLDER_HYPERLINK_EXPECTED_ADDRESS
    )
    notes.notes_text_frame.paragraphs[0].runs[0].hyperlink.address = (
        NOTES_RUN_HYPERLINK_EXPECTED_ADDRESS
    )


def _apply_python_pptx_dropin_notes_paragraph_run_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.lang import MSO_LANGUAGE_ID
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt

    if fixture_id not in _NOTES_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in notes paragraph/run benchmark does not support fixture "
            f"{fixture_id}"
        )
    notes_frame = prs.slides[0].notes_slide.notes_text_frame
    paragraph = notes_frame.paragraphs[0]
    paragraph.runs[0].text = "Run note update"
    paragraph.add_run().text = " appended"
    paragraph.runs[0].font.bold = True
    paragraph.runs[0].font.italic = True
    paragraph.runs[0].font.underline = True
    paragraph.runs[0].font.size = Pt(18)
    paragraph.runs[0].font.name = "Aptos"
    paragraph.runs[0].font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    paragraph.runs[0].font.language_id = MSO_LANGUAGE_ID.GERMAN
    paragraph.font.bold = True
    paragraph.font.italic = True
    paragraph.font.underline = True
    paragraph.font.size = Pt(18)
    paragraph.font.name = "Aptos"
    paragraph.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    paragraph.font.language_id = MSO_LANGUAGE_ID.GERMAN
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.level = 1
    paragraph.space_before = Pt(12)
    paragraph.space_after = Pt(6)
    paragraph.line_spacing = 1.25
    notes_frame.add_paragraph().text = "Second note"


def _apply_python_pptx_dropin_notes_placeholder_clone(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in _NOTES_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in notes placeholder clone benchmark does not support fixture "
            f"{fixture_id}"
        )
    notes = prs.slides[0].notes_slide
    notes.clone_master_placeholders(prs.notes_master)
    notes.shapes.clone_placeholder(prs.notes_master.placeholders[0])


def _apply_wolfppt_dropin_notes_paragraph_run_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.lang import MSO_LANGUAGE_ID
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt

    if fixture_id not in _NOTES_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in notes paragraph/run benchmark does not support fixture "
            f"{fixture_id}"
        )
    notes_frame = prs.slides[0].notes_slide.notes_text_frame
    paragraph = notes_frame.paragraphs[0]
    paragraph.runs[0].text = "Run note update"
    paragraph.add_run().text = " appended"
    paragraph.runs[0].font.bold = True
    paragraph.runs[0].font.italic = True
    paragraph.runs[0].font.underline = True
    paragraph.runs[0].font.size = Pt(18)
    paragraph.runs[0].font.name = "Aptos"
    paragraph.runs[0].font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    paragraph.runs[0].font.language_id = MSO_LANGUAGE_ID.GERMAN
    paragraph.font.bold = True
    paragraph.font.italic = True
    paragraph.font.underline = True
    paragraph.font.size = Pt(18)
    paragraph.font.name = "Aptos"
    paragraph.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    paragraph.font.language_id = MSO_LANGUAGE_ID.GERMAN
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.level = 1
    paragraph.space_before = Pt(12)
    paragraph.space_after = Pt(6)
    paragraph.line_spacing = 1.25
    notes_frame.add_paragraph().text = "Second note"


def _apply_wolfppt_dropin_notes_placeholder_clone(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in _NOTES_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in notes placeholder clone benchmark does not support fixture "
            f"{fixture_id}"
        )
    notes = prs.slides[0].notes_slide
    notes.clone_master_placeholders(prs.notes_master)
    notes.shapes.clone_placeholder(prs.notes_master.placeholders[0])


_MIXED_WORKLOAD_FIXTURES = {
    "workloads/mixed_real_world_deck",
    "workloads/mixed_deal_review_deck",
}

_NOTES_EDIT_FIXTURES = {
    "notes/speaker_notes",
    *_MIXED_WORKLOAD_FIXTURES,
}

_REAL_WORLD_SAFE_TEXT_FIXTURES = {
    "text_basic/title_body_bullets",
    *_MIXED_WORKLOAD_FIXTURES,
}
