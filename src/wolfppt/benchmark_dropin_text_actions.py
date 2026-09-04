"""Text and text-frame drop-in benchmark edit actions."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from .benchmark_cases import (
    ADD_CHART_EXPECTED,
    ADD_CHART_EXPECTED_TRANSFORM,
    ADD_FREEFORM_EXPECTED_TEXT,
    ADD_FREEFORM_EXPECTED_TRANSFORM,
    ADD_MOVIE_EXPECTED_TRANSFORM,
    ADD_OLE_OBJECT_EXPECTED_TRANSFORM,
    ADD_PICTURE_NATIVE_EXPECTED_TRANSFORM,
    ADD_SHAPE_EXPECTED_TEXT,
    ADD_SHAPE_EXPECTED_TRANSFORM,
    APPENDED_PARAGRAPH_TEXT,
    CHART_DATA_EDIT_CASES,
    CHART_READ_EXPECTED_METADATA,
    CHART_TITLE_EXPECTED_TEXT,
    CONNECTOR_EXPECTED_TRANSFORM,
    CONNECTOR_LINE_STYLE_EXPECTED,
    MULTI_FORMAT_RUN_COUNT,
    PARAGRAPH_LEVEL_EXPECTED,
    PARAGRAPH_SPACING_EXPECTED,
    PICTURE_CROP_EXPECTED,
    REPLACED_PARAGRAPH_TEXT,
    RICH_FORMATTING_SIZE,
    SHAPE_HYPERLINK_EXPECTED_ADDRESS,
    SHAPE_NAME_EXPECTED,
    SHAPE_ROTATION_EXPECTED,
    SHAPE_SHADOW_EXPECTED_INHERIT,
    TABLE_CELL_TEXT_FRAME_EXPECTED,
    TABLE_STYLE_FLAGS_EXPECTED,
    TEXT_FRAME_FIT_TEXT_EXPECTED,
    TEXT_FRAME_MARGIN_EXPECTED,
    TEXT_FRAME_WORD_WRAP_EXPECTED,
    TEXT_RUN_HYPERLINK_EXPECTED_ADDRESS,
)
from .benchmark_fit_text_fonts import find_fit_text_font
from .benchmark_dropin_common_actions import _shape_at
from .presentation import Presentation as WolfPresentation


def _dropin_body_text_shape(fixture_id: str, prs: Any, operation: str) -> Any:
    if fixture_id == "text_basic/title_body_bullets":
        if isinstance(prs, WolfPresentation):
            return _shape_at(prs, 0, 1)
        return prs.slides[0].shapes[1]
    if fixture_id == "shapes/grouped_shapes":
        if isinstance(prs, WolfPresentation):
            return _shape_at(prs, 0, 0).shapes[0]
        return prs.slides[0].shapes[0].shapes[0]
    raise RuntimeError(
        f"drop-in {operation} benchmark does not support fixture {fixture_id}"
    )


def _dropin_body_text_paragraph(fixture_id: str, prs: Any, operation: str) -> Any:
    return _dropin_body_text_shape(fixture_id, prs, operation).text_frame.paragraphs[0]


def _apply_python_pptx_dropin_formatting_edit(fixture_id: str, prs: Any) -> None:
    from pptx.dml.color import RGBColor

    if fixture_id == "text_basic/title_body_bullets":
        font = prs.slides[0].shapes[1].text_frame.paragraphs[0].runs[0].font
        font.bold = True
        font.italic = True
        font.underline = True
        font.size = RICH_FORMATTING_SIZE
        font.name = "Aptos"
        font.color.rgb = RGBColor(0x12, 0x34, 0x56)
        return
    if fixture_id == "workloads/multi_format_runs":
        paragraph = prs.slides[0].shapes[0].text_frame.paragraphs[0]
        for run_idx, run in enumerate(paragraph.runs):
            if run_idx >= MULTI_FORMAT_RUN_COUNT:
                break
            font = run.font
            font.bold = True
            font.italic = True
            font.underline = True
            font.size = RICH_FORMATTING_SIZE
            font.name = "Aptos"
            font.color.rgb = RGBColor(0x12, 0x34, 0x56)
        return
    if fixture_id == "shapes/grouped_shapes":
        font = (
            prs.slides[0]
            .shapes[0]
            .shapes[0]
            .text_frame.paragraphs[0]
            .runs[0]
            .font
        )
        font.bold = True
        font.italic = True
        font.underline = True
        font.size = RICH_FORMATTING_SIZE
        font.name = "Aptos"
        font.color.rgb = RGBColor(0x12, 0x34, 0x56)
        return
    raise RuntimeError(f"drop-in formatting benchmark does not support fixture {fixture_id}")


def _apply_python_pptx_dropin_font_language_edit(fixture_id: str, prs: Any) -> None:
    from pptx.enum.lang import MSO_LANGUAGE_ID

    paragraph = _dropin_body_text_paragraph(fixture_id, prs, "font language")
    paragraph.runs[0].font.language_id = MSO_LANGUAGE_ID.FRENCH


def _apply_python_pptx_dropin_font_fill_edit(fixture_id: str, prs: Any) -> None:
    from pptx.dml.color import RGBColor

    font = _dropin_body_text_paragraph(fixture_id, prs, "font fill").runs[0].font
    font.fill.gradient()
    font.fill.gradient_angle = 45.0
    font.fill.gradient_stops[0].position = 0.25
    font.fill.gradient_stops[0].color.rgb = RGBColor(0x12, 0x34, 0x56)


def _apply_python_pptx_dropin_add_run_formatting_edit(fixture_id: str, prs: Any) -> None:
    from pptx.dml.color import RGBColor

    run = _dropin_body_text_paragraph(fixture_id, prs, "add-run formatting").add_run()
    run.text = " Decks"
    font = run.font
    font.bold = True
    font.italic = True
    font.underline = True
    font.size = RICH_FORMATTING_SIZE
    font.name = "Aptos"
    font.color.rgb = RGBColor(0x12, 0x34, 0x56)


def _apply_python_pptx_dropin_add_paragraph_run_formatting_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.dml.color import RGBColor

    shape = _dropin_body_text_shape(fixture_id, prs, "add-paragraph run formatting")
    paragraph = shape.text_frame.add_paragraph()
    run = paragraph.add_run()
    run.text = APPENDED_PARAGRAPH_TEXT
    font = run.font
    font.bold = True
    font.italic = True
    font.underline = True
    font.size = RICH_FORMATTING_SIZE
    font.name = "Aptos"
    font.color.rgb = RGBColor(0x12, 0x34, 0x56)


def _apply_python_pptx_dropin_replace_run_formatting_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.dml.color import RGBColor

    shape = _dropin_body_text_shape(fixture_id, prs, "replace-run formatting")
    paragraph_index = 1 if fixture_id == "text_basic/title_body_bullets" else 0
    paragraph = shape.text_frame.paragraphs[paragraph_index]
    paragraph.text = REPLACED_PARAGRAPH_TEXT
    font = paragraph.runs[0].font
    font.bold = True
    font.italic = True
    font.underline = True
    font.size = RICH_FORMATTING_SIZE
    font.name = "Aptos"
    font.color.rgb = RGBColor(0x12, 0x34, 0x56)


def _apply_python_pptx_dropin_paragraph_format_edit(fixture_id: str, prs: Any) -> None:
    from pptx.enum.text import PP_ALIGN

    _dropin_body_text_paragraph(fixture_id, prs, "paragraph format").alignment = (
        PP_ALIGN.CENTER
    )


def _apply_python_pptx_dropin_paragraph_level_edit(fixture_id: str, prs: Any) -> None:
    paragraph = _dropin_body_text_paragraph(fixture_id, prs, "paragraph level")
    paragraph.level = PARAGRAPH_LEVEL_EXPECTED


def _apply_python_pptx_dropin_paragraph_spacing_edit(fixture_id: str, prs: Any) -> None:
    paragraph = _dropin_body_text_paragraph(fixture_id, prs, "paragraph spacing")
    paragraph.space_before = PARAGRAPH_SPACING_EXPECTED["space_before"]
    paragraph.space_after = PARAGRAPH_SPACING_EXPECTED["space_after"]
    paragraph.line_spacing = PARAGRAPH_SPACING_EXPECTED["line_spacing"]


def _apply_python_pptx_dropin_paragraph_clear_edit(fixture_id: str, prs: Any) -> None:
    _dropin_body_text_paragraph(fixture_id, prs, "paragraph clear").clear()


def _apply_python_pptx_dropin_paragraph_line_break_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    paragraph = _dropin_body_text_paragraph(fixture_id, prs, "paragraph line-break")
    paragraph.add_line_break()
    paragraph.add_run().text = "After"


def _apply_python_pptx_dropin_paragraph_font_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.lang import MSO_LANGUAGE_ID
    from pptx.util import Pt

    font = _dropin_body_text_paragraph(fixture_id, prs, "paragraph font").font
    font.bold = True
    font.italic = True
    font.underline = True
    font.size = Pt(18)
    font.name = "Aptos"
    font.color.rgb = RGBColor(0x0C, 0x22, 0x38)
    font.language_id = MSO_LANGUAGE_ID.GERMAN


def _apply_python_pptx_dropin_add_paragraph_format_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.enum.text import PP_ALIGN

    shape = _dropin_body_text_shape(fixture_id, prs, "add-paragraph format")
    paragraph = shape.text_frame.add_paragraph()
    paragraph.text = APPENDED_PARAGRAPH_TEXT
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.level = PARAGRAPH_LEVEL_EXPECTED
    paragraph.space_before = PARAGRAPH_SPACING_EXPECTED["space_before"]
    paragraph.space_after = PARAGRAPH_SPACING_EXPECTED["space_after"]
    paragraph.line_spacing = PARAGRAPH_SPACING_EXPECTED["line_spacing"]


def _apply_wolfppt_dropin_formatting_edit(fixture_id: str, prs: WolfPresentation) -> None:
    if fixture_id == "text_basic/title_body_bullets":
        font = _dropin_body_text_paragraph(fixture_id, prs, "formatting").runs[0].font
        font.bold = True
        font.italic = True
        font.underline = True
        font.size = RICH_FORMATTING_SIZE
        font.name = "Aptos"
        font.color.rgb = "123456"
        return
    if fixture_id == "workloads/multi_format_runs":
        paragraph = _shape_at(prs, 0, 0).text_frame.paragraphs[0]
        for run_idx, run in enumerate(paragraph.runs):
            if run_idx >= MULTI_FORMAT_RUN_COUNT:
                break
            font = run.font
            font.bold = True
            font.italic = True
            font.underline = True
            font.size = RICH_FORMATTING_SIZE
            font.name = "Aptos"
            font.color.rgb = "123456"
        return
    if fixture_id == "shapes/grouped_shapes":
        font = _dropin_body_text_paragraph(fixture_id, prs, "formatting").runs[0].font
        font.bold = True
        font.italic = True
        font.underline = True
        font.size = RICH_FORMATTING_SIZE
        font.name = "Aptos"
        font.color.rgb = "123456"
        return
    raise RuntimeError(f"drop-in formatting benchmark does not support fixture {fixture_id}")


def _apply_wolfppt_dropin_font_language_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.enum.lang import MSO_LANGUAGE_ID

    paragraph = _dropin_body_text_paragraph(fixture_id, prs, "font language")
    paragraph.runs[0].font.language_id = MSO_LANGUAGE_ID.FRENCH


def _apply_wolfppt_dropin_font_fill_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.dml.color import RGBColor

    font = _dropin_body_text_paragraph(fixture_id, prs, "font fill").runs[0].font
    font.fill.gradient()
    font.fill.gradient_angle = 45.0
    font.fill.gradient_stops[0].position = 0.25
    font.fill.gradient_stops[0].color.rgb = RGBColor(0x12, 0x34, 0x56)


def _apply_wolfppt_dropin_add_run_formatting_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    run = _dropin_body_text_paragraph(fixture_id, prs, "add-run formatting").add_run()
    run.text = " Decks"
    font = run.font
    font.bold = True
    font.italic = True
    font.underline = True
    font.size = RICH_FORMATTING_SIZE
    font.name = "Aptos"
    font.color.rgb = "123456"


def _apply_wolfppt_dropin_add_paragraph_run_formatting_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    shape = _dropin_body_text_shape(fixture_id, prs, "add-paragraph run formatting")
    paragraph = shape.text_frame.add_paragraph()
    run = paragraph.add_run()
    run.text = APPENDED_PARAGRAPH_TEXT
    font = run.font
    font.bold = True
    font.italic = True
    font.underline = True
    font.size = RICH_FORMATTING_SIZE
    font.name = "Aptos"
    font.color.rgb = "123456"


def _apply_wolfppt_dropin_replace_run_formatting_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    shape = _dropin_body_text_shape(fixture_id, prs, "replace-run formatting")
    paragraph_index = 1 if fixture_id == "text_basic/title_body_bullets" else 0
    paragraph = shape.text_frame.paragraphs[paragraph_index]
    paragraph.text = REPLACED_PARAGRAPH_TEXT
    font = paragraph.runs[0].font
    font.bold = True
    font.italic = True
    font.underline = True
    font.size = RICH_FORMATTING_SIZE
    font.name = "Aptos"
    font.color.rgb = "123456"


def _apply_wolfppt_dropin_paragraph_format_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.enum.text import PP_ALIGN

    _dropin_body_text_paragraph(fixture_id, prs, "paragraph format").alignment = (
        PP_ALIGN.CENTER
    )


def _apply_wolfppt_dropin_paragraph_level_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    paragraph = _dropin_body_text_paragraph(fixture_id, prs, "paragraph level")
    paragraph.level = PARAGRAPH_LEVEL_EXPECTED


def _apply_wolfppt_dropin_paragraph_spacing_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    paragraph = _dropin_body_text_paragraph(fixture_id, prs, "paragraph spacing")
    paragraph.space_before = PARAGRAPH_SPACING_EXPECTED["space_before"]
    paragraph.space_after = PARAGRAPH_SPACING_EXPECTED["space_after"]
    paragraph.line_spacing = PARAGRAPH_SPACING_EXPECTED["line_spacing"]


def _apply_wolfppt_dropin_paragraph_clear_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _dropin_body_text_paragraph(fixture_id, prs, "paragraph clear").clear()


def _apply_wolfppt_dropin_paragraph_line_break_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    paragraph = _dropin_body_text_paragraph(fixture_id, prs, "paragraph line-break")
    paragraph.add_line_break()
    paragraph.add_run().text = "After"


def _apply_wolfppt_dropin_paragraph_font_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.enum.lang import MSO_LANGUAGE_ID
    from pptx.util import Pt

    font = _dropin_body_text_paragraph(fixture_id, prs, "paragraph font").font
    font.bold = True
    font.italic = True
    font.underline = True
    font.size = Pt(18)
    font.name = "Aptos"
    font.color.rgb = "0C2238"
    font.language_id = MSO_LANGUAGE_ID.GERMAN


def _apply_wolfppt_dropin_add_paragraph_format_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.enum.text import PP_ALIGN

    shape = _dropin_body_text_shape(fixture_id, prs, "add-paragraph format")
    paragraph = shape.text_frame.add_paragraph()
    paragraph.text = APPENDED_PARAGRAPH_TEXT
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.level = PARAGRAPH_LEVEL_EXPECTED
    paragraph.space_before = PARAGRAPH_SPACING_EXPECTED["space_before"]
    paragraph.space_after = PARAGRAPH_SPACING_EXPECTED["space_after"]
    paragraph.line_spacing = PARAGRAPH_SPACING_EXPECTED["line_spacing"]


def _apply_python_pptx_dropin_text_run_hyperlink_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id == "text_basic/title_body_bullets":
        run = prs.slides[0].shapes[1].text_frame.paragraphs[0].runs[0]
    elif fixture_id == "shapes/grouped_shapes":
        run = prs.slides[0].shapes[0].shapes[0].text_frame.paragraphs[0].runs[0]
    else:
        raise RuntimeError(
            f"drop-in text run hyperlink benchmark does not support fixture {fixture_id}"
        )
    run.hyperlink.address = TEXT_RUN_HYPERLINK_EXPECTED_ADDRESS


def _apply_wolfppt_dropin_text_run_hyperlink_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    run = _dropin_body_text_paragraph(fixture_id, prs, "text run hyperlink").runs[0]
    run.hyperlink.address = TEXT_RUN_HYPERLINK_EXPECTED_ADDRESS


def _dropin_text_frame_target(fixture_id: str, prs: Any, operation: str) -> Any:
    return _dropin_text_frame_shape_target(fixture_id, prs, operation).text_frame


def _dropin_text_frame_shape_target(fixture_id: str, prs: Any, operation: str) -> Any:
    if fixture_id == "text_basic/title_body_bullets":
        if isinstance(prs, WolfPresentation):
            return _shape_at(prs, 0, 0)
        return prs.slides[0].shapes[0]
    if fixture_id == "shapes/grouped_shapes":
        if isinstance(prs, WolfPresentation):
            return _shape_at(prs, 0, 0).shapes[0]
        return prs.slides[0].shapes[0].shapes[0]
    raise RuntimeError(
        f"drop-in text frame {operation} benchmark does not support fixture {fixture_id}"
    )


def _apply_python_pptx_dropin_text_frame_margin_edit(fixture_id: str, prs: Any) -> None:
    text_frame = _dropin_text_frame_target(fixture_id, prs, "margin")
    text_frame.margin_left = TEXT_FRAME_MARGIN_EXPECTED["margin_left"]
    text_frame.margin_right = TEXT_FRAME_MARGIN_EXPECTED["margin_right"]
    text_frame.margin_top = TEXT_FRAME_MARGIN_EXPECTED["margin_top"]
    text_frame.margin_bottom = TEXT_FRAME_MARGIN_EXPECTED["margin_bottom"]


def _apply_wolfppt_dropin_text_frame_margin_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    text_frame = _dropin_text_frame_target(fixture_id, prs, "margin")
    text_frame.margin_left = TEXT_FRAME_MARGIN_EXPECTED["margin_left"]
    text_frame.margin_right = TEXT_FRAME_MARGIN_EXPECTED["margin_right"]
    text_frame.margin_top = TEXT_FRAME_MARGIN_EXPECTED["margin_top"]
    text_frame.margin_bottom = TEXT_FRAME_MARGIN_EXPECTED["margin_bottom"]


def _apply_python_pptx_dropin_text_frame_word_wrap_edit(fixture_id: str, prs: Any) -> None:
    text_frame = _dropin_text_frame_target(fixture_id, prs, "word-wrap")
    text_frame.word_wrap = TEXT_FRAME_WORD_WRAP_EXPECTED


def _apply_wolfppt_dropin_text_frame_word_wrap_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    text_frame = _dropin_text_frame_target(fixture_id, prs, "word-wrap")
    text_frame.word_wrap = TEXT_FRAME_WORD_WRAP_EXPECTED


def _apply_python_pptx_dropin_text_frame_vertical_anchor_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.enum.text import MSO_ANCHOR

    text_frame = _dropin_text_frame_target(fixture_id, prs, "vertical-anchor")
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE


def _apply_wolfppt_dropin_text_frame_vertical_anchor_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.enum.text import MSO_ANCHOR

    text_frame = _dropin_text_frame_target(fixture_id, prs, "vertical-anchor")
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE


def _apply_python_pptx_dropin_text_frame_auto_size_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.enum.text import MSO_AUTO_SIZE

    text_frame = _dropin_text_frame_target(fixture_id, prs, "auto-size")
    text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def _apply_wolfppt_dropin_text_frame_auto_size_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.enum.text import MSO_AUTO_SIZE

    text_frame = _dropin_text_frame_target(fixture_id, prs, "auto-size")
    text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def _apply_python_pptx_dropin_text_frame_fit_text_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.util import Inches

    font_family, font_file = find_fit_text_font()
    shape = _dropin_text_frame_shape_target(fixture_id, prs, "fit-text")
    if fixture_id == "shapes/grouped_shapes":
        shape.width = Inches(2)
        shape.height = Inches(1)
    text_frame = shape.text_frame
    text_frame.fit_text(
        font_family=font_family,
        max_size=18,
        bold=TEXT_FRAME_FIT_TEXT_EXPECTED["bold"],
        italic=TEXT_FRAME_FIT_TEXT_EXPECTED["italic"],
        font_file=font_file,
    )


def _apply_wolfppt_dropin_text_frame_fit_text_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.util import Inches

    font_family, font_file = find_fit_text_font()
    shape = _dropin_text_frame_shape_target(fixture_id, prs, "fit-text")
    if fixture_id == "shapes/grouped_shapes":
        shape.width = Inches(2)
        shape.height = Inches(1)
    text_frame = shape.text_frame
    text_frame.fit_text(
        font_family=font_family,
        max_size=18,
        bold=TEXT_FRAME_FIT_TEXT_EXPECTED["bold"],
        italic=TEXT_FRAME_FIT_TEXT_EXPECTED["italic"],
        font_file=font_file,
    )


def _fit_text_font_file() -> str:
    return find_fit_text_font()[1]
