"""Chart title and legend drop-in benchmark edit actions."""

from __future__ import annotations

from typing import Any

from .benchmark_cases import (
    CHART_TITLE_PARAGRAPH_FONT_EXPECTED,
    CHART_TITLE_PARAGRAPH_FORMAT_EXPECTED,
    CHART_TITLE_TEXT_FRAME_FLOW_EXPECTED,
    CHART_TITLE_TEXT_FRAME_EXPECTED,
    CHART_TITLE_EXPECTED_TEXT,
)
from .benchmark_dropin_chart_readers import (
    CHART_EDIT_FIXTURES,
    _first_chart_shape,
)
from .presentation import Presentation as WolfPresentation


def _apply_python_pptx_dropin_chart_title_edit(fixture_id: str, prs: Any) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_PATTERN_TYPE
    from pptx.util import Pt

    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart title benchmark does not support fixture {fixture_id}"
        )
    chart_title = _first_chart_shape(prs).chart.chart_title
    chart_title.text_frame.text = CHART_TITLE_EXPECTED_TEXT
    chart_title.format.fill.background()
    chart_title.format.line.fill.patterned()
    chart_title.format.line.fill.pattern = MSO_PATTERN_TYPE.DIVOT
    chart_title.format.line.fill.fore_color.rgb = RGBColor(0x65, 0x43, 0x21)
    chart_title.format.line.fill.back_color.rgb = RGBColor(0x10, 0x20, 0x30)
    chart_title.format.line.width = Pt(2)
    chart_title.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def _apply_wolfppt_dropin_chart_title_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_PATTERN_TYPE
    from pptx.util import Pt

    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart title benchmark does not support fixture {fixture_id}"
        )
    chart_title = _first_chart_shape(prs).chart.chart_title
    chart_title.text_frame.text = CHART_TITLE_EXPECTED_TEXT
    chart_title.format.fill.background()
    chart_title.format.line.fill.patterned()
    chart_title.format.line.fill.pattern = MSO_PATTERN_TYPE.DIVOT
    chart_title.format.line.fill.fore_color.rgb = RGBColor(0x65, 0x43, 0x21)
    chart_title.format.line.fill.back_color.rgb = RGBColor(0x10, 0x20, 0x30)
    chart_title.format.line.width = Pt(2)
    chart_title.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def _apply_python_pptx_dropin_chart_title_remove(fixture_id: str, prs: Any) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart title removal benchmark does not support fixture {fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.has_title = False


def _apply_wolfppt_dropin_chart_title_remove(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart title removal benchmark does not support fixture {fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.has_title = False


def _apply_python_pptx_dropin_chart_legend_edit(fixture_id: str, prs: Any) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_LEGEND_POSITION
    from pptx.util import Pt

    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart legend benchmark does not support fixture {fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.bold = True
    chart.legend.font.italic = True
    chart.legend.font.underline = True
    chart.legend.font.size = Pt(11)
    chart.legend.font.fill.gradient()
    chart.legend.font.fill.gradient_angle = 45.0
    chart.legend.font.fill.gradient_stops[0].position = 0.25
    chart.legend.font.fill.gradient_stops[0].color.rgb = RGBColor(0xAB, 0xCD, 0xEF)
    chart.legend.font.name = "Aptos"


def _apply_wolfppt_dropin_chart_legend_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_LEGEND_POSITION
    from pptx.util import Pt

    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart legend benchmark does not support fixture {fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.has_legend = True
    if chart.legend is None:
        raise AttributeError("chart legend is unavailable")
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.bold = True
    chart.legend.font.italic = True
    chart.legend.font.underline = True
    chart.legend.font.size = Pt(11)
    chart.legend.font.fill.gradient()
    chart.legend.font.fill.gradient_angle = 45.0
    chart.legend.font.fill.gradient_stops[0].position = 0.25
    chart.legend.font.fill.gradient_stops[0].color.rgb = RGBColor(0xAB, 0xCD, 0xEF)
    chart.legend.font.name = "Aptos"


def _apply_python_pptx_dropin_chart_axis_title_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.util import Pt

    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart axis title benchmark does not support fixture {fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    category_title = chart.category_axis.axis_title
    category_title.text_frame.text = "Quarter"
    category_title.format.fill.solid()
    category_title.format.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    category_title.format.line.fill.solid()
    category_title.format.line.color.rgb = RGBColor(0x65, 0x43, 0x21)
    category_title.format.line.width = Pt(2)
    category_title.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    value_title = chart.value_axis.axis_title
    value_title.text_frame.text = "Revenue"
    value_title.format.fill.solid()
    value_title.format.fill.fore_color.rgb = RGBColor(0xAB, 0xCD, 0xEF)
    value_title.format.line.fill.solid()
    value_title.format.line.color.rgb = RGBColor(0xFE, 0xDC, 0xBA)
    value_title.format.line.width = Pt(1)
    value_title.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH_DOT


def _apply_wolfppt_dropin_chart_axis_title_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.util import Pt

    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart axis title benchmark does not support fixture {fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    category_title = chart.category_axis.axis_title
    category_title.text_frame.text = "Quarter"
    category_title.format.fill.solid()
    category_title.format.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    category_title.format.line.fill.solid()
    category_title.format.line.color.rgb = RGBColor(0x65, 0x43, 0x21)
    category_title.format.line.width = Pt(2)
    category_title.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    value_title = chart.value_axis.axis_title
    value_title.text_frame.text = "Revenue"
    value_title.format.fill.solid()
    value_title.format.fill.fore_color.rgb = RGBColor(0xAB, 0xCD, 0xEF)
    value_title.format.line.fill.solid()
    value_title.format.line.color.rgb = RGBColor(0xFE, 0xDC, 0xBA)
    value_title.format.line.width = Pt(1)
    value_title.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH_DOT


def _apply_python_pptx_dropin_chart_axis_title_remove(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in chart axis-title removal benchmark does not support fixture "
            f"{fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.category_axis.has_title = False
    chart.value_axis.has_title = False


def _apply_wolfppt_dropin_chart_axis_title_remove(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in chart axis-title removal benchmark does not support fixture "
            f"{fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.category_axis.has_title = False
    chart.value_axis.has_title = False


def _apply_python_pptx_dropin_chart_title_text_frame_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in chart title text-frame benchmark does not support fixture "
            f"{fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.has_title = True
    _apply_rich_chart_title_text_frame(chart.chart_title.text_frame)
    chart.category_axis.has_title = True
    _apply_rich_chart_title_text_frame(chart.category_axis.axis_title.text_frame)


def _apply_wolfppt_dropin_chart_title_text_frame_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in chart title text-frame benchmark does not support fixture "
            f"{fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.has_title = True
    _apply_rich_chart_title_text_frame(chart.chart_title.text_frame)
    chart.category_axis.has_title = True
    _apply_rich_chart_title_text_frame(chart.category_axis.axis_title.text_frame)


def _apply_python_pptx_dropin_chart_title_text_frame_flow_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in chart title text-frame flow benchmark does not support fixture "
            f"{fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.has_title = True
    _apply_chart_title_text_frame_flow(chart.chart_title.text_frame)
    chart.value_axis.has_title = True
    _apply_chart_title_text_frame_flow(chart.value_axis.axis_title.text_frame)


def _apply_wolfppt_dropin_chart_title_text_frame_flow_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in chart title text-frame flow benchmark does not support fixture "
            f"{fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.has_title = True
    _apply_chart_title_text_frame_flow(chart.chart_title.text_frame)
    chart.value_axis.has_title = True
    _apply_chart_title_text_frame_flow(chart.value_axis.axis_title.text_frame)


def _apply_python_pptx_dropin_chart_title_paragraph_format_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in chart title paragraph format benchmark does not support fixture "
            f"{fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.has_title = True
    _apply_chart_title_paragraph_format(chart.chart_title.text_frame)
    chart.category_axis.has_title = True
    _apply_chart_title_paragraph_format(chart.category_axis.axis_title.text_frame)


def _apply_wolfppt_dropin_chart_title_paragraph_format_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in chart title paragraph format benchmark does not support fixture "
            f"{fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.has_title = True
    _apply_chart_title_paragraph_format(chart.chart_title.text_frame)
    chart.category_axis.has_title = True
    _apply_chart_title_paragraph_format(chart.category_axis.axis_title.text_frame)


def _apply_python_pptx_dropin_chart_title_paragraph_font_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in chart title paragraph font benchmark does not support fixture "
            f"{fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.has_title = True
    _apply_chart_title_paragraph_font(chart.chart_title.text_frame)
    chart.value_axis.has_title = True
    _apply_chart_title_paragraph_font(chart.value_axis.axis_title.text_frame)


def _apply_wolfppt_dropin_chart_title_paragraph_font_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            "drop-in chart title paragraph font benchmark does not support fixture "
            f"{fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.has_title = True
    _apply_chart_title_paragraph_font(chart.chart_title.text_frame)
    chart.value_axis.has_title = True
    _apply_chart_title_paragraph_font(chart.value_axis.axis_title.text_frame)


def _apply_rich_chart_title_text_frame(text_frame: Any) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_PATTERN_TYPE
    from pptx.enum.lang import MSO_LANGUAGE_ID
    from pptx.util import Pt

    text_frame.text = "Revenue"
    run = text_frame.paragraphs[0].add_run()
    run.text = CHART_TITLE_TEXT_FRAME_EXPECTED["run_text"]
    run.font.bold = True
    run.font.italic = True
    run.font.underline = True
    run.font.size = Pt(14)
    run.font.fill.patterned()
    run.font.fill.pattern = MSO_PATTERN_TYPE.DIVOT
    run.font.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    run.font.fill.back_color.rgb = RGBColor(0xAB, 0xCD, 0xEF)
    run.font.name = CHART_TITLE_TEXT_FRAME_EXPECTED["font_name"]
    run.font.language_id = MSO_LANGUAGE_ID.FRENCH
    paragraph = text_frame.add_paragraph()
    paragraph.text = "Q4 outlook"


def _apply_chart_title_paragraph_format(text_frame: Any) -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt

    text_frame.text = CHART_TITLE_PARAGRAPH_FORMAT_EXPECTED["text"]
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.level = CHART_TITLE_PARAGRAPH_FORMAT_EXPECTED["level"]
    paragraph.space_before = Pt(12)
    paragraph.space_after = Pt(6)
    paragraph.line_spacing = CHART_TITLE_PARAGRAPH_FORMAT_EXPECTED["line_spacing"]


def _apply_chart_title_paragraph_font(text_frame: Any) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.lang import MSO_LANGUAGE_ID
    from pptx.util import Pt

    text_frame.text = CHART_TITLE_PARAGRAPH_FONT_EXPECTED["text"]
    paragraph = text_frame.paragraphs[0]
    paragraph.font.bold = True
    paragraph.font.italic = True
    paragraph.font.underline = True
    paragraph.font.size = Pt(18)
    paragraph.font.fill.gradient()
    paragraph.font.fill.gradient_angle = 45.0
    paragraph.font.fill.gradient_stops[0].position = 0.25
    paragraph.font.fill.gradient_stops[0].color.rgb = RGBColor(0x0C, 0x22, 0x38)
    paragraph.font.name = CHART_TITLE_PARAGRAPH_FONT_EXPECTED["font_name"]
    paragraph.font.language_id = MSO_LANGUAGE_ID.GERMAN


def _apply_chart_title_text_frame_flow(text_frame: Any) -> None:
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.util import Inches

    text_frame.margin_left = Inches(0.2)
    text_frame.margin_right = Inches(0.3)
    text_frame.margin_top = Inches(0.1)
    text_frame.margin_bottom = Inches(0.15)
    text_frame.word_wrap = True
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    text_frame.text = CHART_TITLE_TEXT_FRAME_FLOW_EXPECTED["text"]
