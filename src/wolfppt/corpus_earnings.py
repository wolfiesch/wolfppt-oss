"""Earnings board appendix workload fixture builder."""

from __future__ import annotations

from io import BytesIO

from .corpus_slide_helpers import add_textbox as _add_textbox
from .corpus_slide_helpers import png_1x1 as _png_1x1


def add_earnings_board_appendix_pack(prs, slide) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches

    _add_textbox(
        slide,
        "Earnings Board Appendix Pack",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    _add_textbox(
        slide,
        "Quarterly KPI review, guidance bridge, segment risk, and disclosure actions.",
        Inches(0.7),
        Inches(1.05),
        Inches(8.2),
        Inches(0.8),
    )
    slide.shapes.add_picture(
        BytesIO(_png_1x1()),
        Inches(8.85),
        Inches(0.35),
        Inches(0.4),
        Inches(0.4),
    )
    slide.notes_slide.notes_text_frame.text = (
        "Synthetic earnings board appendix for real-corpus preservation evidence."
    )

    kpi_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        kpi_slide,
        "Quarterly KPI Scorecard",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    kpis = kpi_slide.shapes.add_table(
        6,
        5,
        Inches(0.45),
        Inches(1.0),
        Inches(9.1),
        Inches(3.0),
    ).table
    for row_idx, row in enumerate(
        [
            ["Metric", "Q1", "Q2", "Q3", "Q4 Outlook"],
            ["Revenue", "$118m", "$126m", "$133m", "$141m"],
            ["Gross margin", "71%", "72%", "73%", "73%"],
            ["Net retention", "112%", "114%", "115%", "116%"],
            ["Cash burn", "$8m", "$6m", "$5m", "$4m"],
            ["Rule of 40", "32", "36", "39", "42"],
        ]
    ):
        for col_idx, value in enumerate(row):
            kpis.cell(row_idx, col_idx).text = value
    kpi_slide.notes_slide.notes_text_frame.text = (
        "KPI scorecard table should preserve quarterly values and outlook labels."
    )

    guidance_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        guidance_slide,
        "Guidance Range Bridge",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    chart_data = CategoryChartData()
    chart_data.categories = ["Base", "Expansion", "FX", "Churn", "Upside"]
    chart_data.add_series("Low case", (130, 8, -3, -5, 0))
    chart_data.add_series("High case", (132, 12, -2, -3, 6))
    guidance_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.7),
        Inches(1.0),
        Inches(8.4),
        Inches(4.0),
        chart_data,
    )
    guidance_slide.notes_slide.notes_text_frame.text = (
        "Guidance bridge chart workbook and notes should survive round-trip."
    )

    segment_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        segment_slide,
        "Segment Risk Bridge",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    segment = segment_slide.shapes.add_table(
        5,
        4,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(2.45),
    ).table
    for row_idx, row in enumerate(
        [
            ["Segment", "Signal", "Risk", "Action"],
            ["Enterprise", "Pipeline slip", "Medium", "Deal desk review"],
            ["Mid-market", "Discount pressure", "High", "Packaging refresh"],
            ["SMB", "Support load", "Medium", "Self-serve enablement"],
            ["International", "FX drag", "Low", "Hedge disclosure"],
        ]
    ):
        for col_idx, value in enumerate(row):
            segment.cell(row_idx, col_idx).text = value
    segment_slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(1.0),
        Inches(3.75),
        Inches(8.4),
        Inches(3.75),
    )
    segment_slide.notes_slide.notes_text_frame.text = (
        "Segment risk table and connector line should remain intact."
    )

    actions_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        actions_slide,
        "Disclosure Owner Actions",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    group = actions_slide.shapes.add_group_shape()
    for idx, (label, owner) in enumerate(
        (
            ("Finalize guidance language", "FP&A"),
            ("Revenue recognition memo", "Accounting"),
            ("Q&A risk review", "IR"),
            ("Board appendix QA", "Legal"),
        )
    ):
        tile = group.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.25 + (idx % 2) * 2.85),
            Inches(0.25 + (idx // 2) * 1.05),
            Inches(2.5),
            Inches(0.78),
        )
        tile.text = f"{label}\n{owner}"
    actions_slide.notes_slide.notes_text_frame.text = (
        "Grouped disclosure action tiles should keep child text and transforms."
    )

    appendix_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        appendix_slide,
        "Appendix: Evidence and Draft Sources",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    appendix = appendix_slide.shapes.add_table(
        5,
        3,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(2.6),
    ).table
    for row_idx, row in enumerate(
        [
            ["Artifact", "Source", "Owner"],
            ["Guidance bridge", "guidance_bridge.xlsx", "FP&A"],
            ["Segment risks", "segment_risk_register.csv", "Strategy"],
            ["Disclosure checklist", "disclosure_checklist.docx", "Legal"],
            ["Q&A log", "earnings_qa_log.xlsx", "IR"],
        ]
    ):
        for col_idx, value in enumerate(row):
            appendix.cell(row_idx, col_idx).text = value
    appendix_slide.notes_slide.notes_text_frame.text = (
        "Evidence index should preserve source filenames and owners."
    )
