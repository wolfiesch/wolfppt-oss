"""Basic shape drop-in benchmark operation implementations."""

from __future__ import annotations

from pathlib import Path
from time import perf_counter
from typing import Any

from .benchmark_dropin_actions import (
    _apply_python_pptx_dropin_add_shape,
    _apply_python_pptx_dropin_add_slide,
    _apply_python_pptx_dropin_add_table,
    _apply_python_pptx_dropin_add_textbox,
    _apply_python_pptx_dropin_add_title_slide,
    _apply_python_pptx_dropin_group_child_adjustment_edit,
    _apply_python_pptx_dropin_group_child_geometry_edit,
    _apply_python_pptx_dropin_shape_adjustment_edit,
    _apply_python_pptx_dropin_shape_geometry_edit,
    _apply_python_pptx_dropin_shape_line_element_edit,
    _apply_wolfppt_dropin_add_shape,
    _apply_wolfppt_dropin_add_slide,
    _apply_wolfppt_dropin_add_table,
    _apply_wolfppt_dropin_add_textbox,
    _apply_wolfppt_dropin_add_title_slide,
    _apply_wolfppt_dropin_group_child_adjustment_edit,
    _apply_wolfppt_dropin_group_child_geometry_edit,
    _apply_wolfppt_dropin_shape_adjustment_edit,
    _apply_wolfppt_dropin_shape_geometry_edit,
    _apply_wolfppt_dropin_shape_line_element_edit,
)
from .benchmark_dropin_shape_details import (
    _dropin_add_shape_details,
    _dropin_add_slide_details,
    _dropin_add_table_details,
    _dropin_add_textbox_details,
    _dropin_add_title_slide_details,
    _dropin_group_child_adjustment_edit_details,
    _dropin_group_child_geometry_edit_details,
    _dropin_shape_adjustment_edit_details,
    _dropin_shape_geometry_edit_details,
    _dropin_shape_line_element_edit_details,
)
from .benchmark_runtime import elapsed_ms as _elapsed_ms, stamp as _stamp
from .presentation import Presentation as WolfPresentation


def _bench_python_pptx_dropin_add_shape(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-add-shape-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_shape(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_shape_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_shape(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-add-shape-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_shape(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_shape_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_textbox(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-add-textbox-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_textbox(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_textbox_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_textbox(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-add-textbox-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_textbox(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_textbox_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_table(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-add-table-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_table(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_table_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_table(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-add-table-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_table(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_table_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_slide(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-add-slide-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_slide(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_slide_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_slide(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-add-slide-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_slide(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_slide_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_title_slide(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-add-title-slide-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_title_slide(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_title_slide_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_title_slide(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-add-title-slide-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_title_slide(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_title_slide_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_shape_adjustment_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-shape-adjustment-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_shape_adjustment_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_shape_adjustment_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_shape_adjustment_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-shape-adjustment-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_shape_adjustment_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_shape_adjustment_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_shape_geometry_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-shape-geometry-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_shape_geometry_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_shape_geometry_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_shape_geometry_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-shape-geometry-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_shape_geometry_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_shape_geometry_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_group_child_geometry_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-group-child-geometry-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_group_child_geometry_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_group_child_geometry_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_group_child_geometry_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-group-child-geometry-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_group_child_geometry_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_group_child_geometry_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_group_child_adjustment_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-group-child-adjustment-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_group_child_adjustment_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_group_child_adjustment_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_group_child_adjustment_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-group-child-adjustment-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_group_child_adjustment_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_group_child_adjustment_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_shape_line_element_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-shape-line-element-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_shape_line_element_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_shape_line_element_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_shape_line_element_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-shape-line-element-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_shape_line_element_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_shape_line_element_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )
