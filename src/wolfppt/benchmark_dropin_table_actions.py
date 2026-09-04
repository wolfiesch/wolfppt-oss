"""Table drop-in benchmark edit actions."""

from __future__ import annotations

from typing import Any

from .benchmark_cases import (
    TABLE_CELL_PARAGRAPH_LINE_BREAK_EXPECTED,
    TABLE_CELL_PARAGRAPH_FONT_EXPECTED,
    TABLE_CELL_PARAGRAPH_RUNS_EXPECTED,
    TABLE_CELL_RUN_FONT_EXPECTED,
    TABLE_CELL_RUN_HYPERLINK_EXPECTED,
    TABLE_CELL_TEXT_FRAME_EXPECTED,
    TABLE_CELL_TEXT_FRAME_FLOW_EXPECTED,
    TABLE_DIMENSIONS_EXPECTED,
    TABLE_STYLE_FLAGS_EXPECTED,
)
from .presentation import Presentation as WolfPresentation
from .shape_core_facade import Shape

TABLE_CELL_FIXTURES = (
    "tables/simple_table",
    "workloads/mixed_real_world_deck",
    "workloads/mixed_deal_review_deck",
)


def _first_table_shape(fixture_id: str, prs: Any, operation: str) -> Any:
    if fixture_id not in TABLE_CELL_FIXTURES:
        raise RuntimeError(
            f"drop-in {operation} benchmark does not support fixture {fixture_id}"
        )
    for slide in prs.slides:
        shell_shape = _first_shell_table_shape(slide)
        if shell_shape is not None:
            return shell_shape
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                return shape
    raise RuntimeError(f"fixture {fixture_id} does not contain a table shape")


def _first_shell_table_shape(slide: Any) -> Any | None:
    payload = getattr(slide, "_payload", None)
    if not isinstance(payload, dict):
        return None
    hints = payload.get("table_shape_hints")
    if not isinstance(hints, list):
        return None
    for hint in hints:
        if not isinstance(hint, dict) or not hint.get("tables"):
            continue
        return Shape(
            slide,
            int(hint.get("_shape_index", 0)),
            hint,
            table_index=int(hint.get("_table_index", 0)),
        )
    return None


def _apply_python_pptx_dropin_table_cell_text_frame_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_table_cell_text_frame_case(
        _first_table_shape(fixture_id, prs, "table-cell text-frame").table
    )


def _apply_wolfppt_dropin_table_cell_text_frame_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_table_cell_text_frame_case(
        _first_table_shape(fixture_id, prs, "table-cell text-frame").table
    )


def _apply_table_cell_text_frame_case(table: Any) -> None:
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches

    cell = table.cell(0, 0)
    cell.text_frame.text = TABLE_CELL_TEXT_FRAME_EXPECTED["text"]
    cell.text_frame.margin_left = Inches(0.2)
    cell.margin_right = Inches(0.3)
    cell.margin_top = Inches(0.1)
    cell.text_frame.margin_bottom = Inches(0.15)
    cell.vertical_anchor = MSO_ANCHOR.BOTTOM
    cell.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE


def _apply_python_pptx_dropin_table_cell_text_frame_flow_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_table_cell_text_frame_flow_case(
        _first_table_shape(fixture_id, prs, "table-cell text-frame flow").table
    )


def _apply_wolfppt_dropin_table_cell_text_frame_flow_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_table_cell_text_frame_flow_case(
        _first_table_shape(fixture_id, prs, "table-cell text-frame flow").table
    )


def _apply_table_cell_text_frame_flow_case(table: Any) -> None:
    from pptx.enum.text import MSO_AUTO_SIZE

    frame = table.cell(0, 0).text_frame
    frame.clear()
    frame.add_paragraph().text = TABLE_CELL_TEXT_FRAME_FLOW_EXPECTED["paragraphs"][1]
    frame.word_wrap = TABLE_CELL_TEXT_FRAME_FLOW_EXPECTED["word_wrap"]
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def _apply_python_pptx_dropin_table_cell_paragraph_runs_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_table_cell_paragraph_runs_case(
        _first_table_shape(fixture_id, prs, "table-cell paragraph-runs").table
    )


def _apply_wolfppt_dropin_table_cell_paragraph_runs_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_table_cell_paragraph_runs_case(
        _first_table_shape(fixture_id, prs, "table-cell paragraph-runs").table
    )


def _apply_table_cell_paragraph_runs_case(table: Any) -> None:
    paragraph = table.cell(0, 0).text_frame.paragraphs[0]
    paragraph.runs[0].text = TABLE_CELL_PARAGRAPH_RUNS_EXPECTED["runs"][0]
    paragraph.add_run().text = TABLE_CELL_PARAGRAPH_RUNS_EXPECTED["runs"][1]


def _apply_python_pptx_dropin_table_cell_paragraph_line_break_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_table_cell_paragraph_line_break_case(
        _first_table_shape(fixture_id, prs, "table-cell paragraph-line-break").table
    )


def _apply_wolfppt_dropin_table_cell_paragraph_line_break_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_table_cell_paragraph_line_break_case(
        _first_table_shape(fixture_id, prs, "table-cell paragraph-line-break").table
    )


def _apply_table_cell_paragraph_line_break_case(table: Any) -> None:
    paragraph = table.cell(0, 0).text_frame.paragraphs[0]
    paragraph.runs[0].text = TABLE_CELL_PARAGRAPH_LINE_BREAK_EXPECTED["runs"][0]
    paragraph.add_line_break()
    paragraph.add_run().text = TABLE_CELL_PARAGRAPH_LINE_BREAK_EXPECTED["runs"][1]


def _apply_python_pptx_dropin_table_cell_run_font_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_table_cell_run_font_case(
        _first_table_shape(fixture_id, prs, "table-cell run-font").table
    )


def _apply_wolfppt_dropin_table_cell_run_font_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_table_cell_run_font_case(
        _first_table_shape(fixture_id, prs, "table-cell run-font").table
    )


def _apply_table_cell_run_font_case(table: Any) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.enum.lang import MSO_LANGUAGE_ID
    from pptx.util import Pt

    paragraph = table.cell(0, 0).text_frame.paragraphs[0]
    paragraph.runs[0].text = TABLE_CELL_RUN_FONT_EXPECTED["runs"][0]
    paragraph.runs[0].font.color.theme_color = MSO_THEME_COLOR.ACCENT_2
    run = paragraph.add_run()
    run.text = TABLE_CELL_RUN_FONT_EXPECTED["runs"][1]
    run.font.bold = TABLE_CELL_RUN_FONT_EXPECTED["bold"]
    run.font.italic = TABLE_CELL_RUN_FONT_EXPECTED["italic"]
    run.font.underline = TABLE_CELL_RUN_FONT_EXPECTED["underline"]
    run.font.size = Pt(14)
    run.font.name = TABLE_CELL_RUN_FONT_EXPECTED["font_name"]
    run.font.language_id = MSO_LANGUAGE_ID.FRENCH
    run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)


def _apply_python_pptx_dropin_table_cell_run_hyperlink_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_table_cell_run_hyperlink_case(
        _first_table_shape(fixture_id, prs, "table-cell run-hyperlink").table
    )


def _apply_wolfppt_dropin_table_cell_run_hyperlink_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_table_cell_run_hyperlink_case(
        _first_table_shape(fixture_id, prs, "table-cell run-hyperlink").table
    )


def _apply_table_cell_run_hyperlink_case(table: Any) -> None:
    paragraph = table.cell(0, 0).text_frame.paragraphs[0]
    paragraph.runs[0].text = TABLE_CELL_RUN_HYPERLINK_EXPECTED["runs"][0]
    run = paragraph.add_run()
    run.text = TABLE_CELL_RUN_HYPERLINK_EXPECTED["runs"][1]
    run.hyperlink.address = TABLE_CELL_RUN_HYPERLINK_EXPECTED[
        "run_hyperlink_address"
    ]


def _apply_python_pptx_dropin_table_cell_paragraph_font_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_table_cell_paragraph_font_case(
        _first_table_shape(fixture_id, prs, "table-cell paragraph-font").table
    )


def _apply_wolfppt_dropin_table_cell_paragraph_font_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_table_cell_paragraph_font_case(
        _first_table_shape(fixture_id, prs, "table-cell paragraph-font").table
    )


def _apply_table_cell_paragraph_font_case(table: Any) -> None:
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.enum.lang import MSO_LANGUAGE_ID
    from pptx.util import Pt

    paragraph = table.cell(0, 0).text_frame.paragraphs[0]
    paragraph.text = TABLE_CELL_PARAGRAPH_FONT_EXPECTED["text"]
    font = paragraph.font
    font.bold = TABLE_CELL_PARAGRAPH_FONT_EXPECTED["bold"]
    font.italic = TABLE_CELL_PARAGRAPH_FONT_EXPECTED["italic"]
    font.underline = TABLE_CELL_PARAGRAPH_FONT_EXPECTED["underline"]
    font.size = Pt(14)
    font.language_id = MSO_LANGUAGE_ID.FRENCH
    font.name = TABLE_CELL_PARAGRAPH_FONT_EXPECTED["font_name"]
    font.color.theme_color = MSO_THEME_COLOR.ACCENT_3


def _apply_python_pptx_dropin_table_cell_fill_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_table_cell_fill_case(
        _first_table_shape(fixture_id, prs, "table-cell fill").table
    )


def _apply_wolfppt_dropin_table_cell_fill_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_table_cell_fill_case(
        _first_table_shape(fixture_id, prs, "table-cell fill").table
    )


def _apply_table_cell_fill_case(table: Any) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_PATTERN_TYPE, MSO_THEME_COLOR

    table.cell(0, 0).fill.solid()
    table.cell(0, 0).fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    table.cell(1, 0).fill.patterned()
    table.cell(1, 0).fill.pattern = MSO_PATTERN_TYPE.DIVOT
    table.cell(1, 0).fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_3
    table.cell(1, 0).fill.back_color.theme_color = MSO_THEME_COLOR.ACCENT_4
    gradient = table.cell(0, 1).fill
    gradient.gradient()
    gradient.gradient_angle = 45.0
    gradient.gradient_stops[0].position = 0.25
    gradient.gradient_stops[0].color.rgb = RGBColor(0x65, 0x43, 0x21)
    table.cell(1, 1).fill.background()


def _apply_python_pptx_dropin_table_cell_merge_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_table_cell_merge_case(
        _first_table_shape(fixture_id, prs, "table-cell merge").table
    )


def _apply_wolfppt_dropin_table_cell_merge_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_table_cell_merge_case(
        _first_table_shape(fixture_id, prs, "table-cell merge").table
    )


def _apply_table_cell_merge_case(table: Any) -> None:
    table.cell(1, 1).merge(table.cell(0, 0))


def _apply_python_pptx_dropin_table_dimensions_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_table_dimensions_case(
        _first_table_shape(fixture_id, prs, "table dimensions").table
    )


def _apply_wolfppt_dropin_table_dimensions_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_table_dimensions_case(
        _first_table_shape(fixture_id, prs, "table dimensions").table
    )


def _apply_table_dimensions_case(table: Any) -> None:
    table.rows[0].height = TABLE_DIMENSIONS_EXPECTED["row_heights"][0]
    table.columns[0].width = TABLE_DIMENSIONS_EXPECTED["column_widths"][0]


def _apply_python_pptx_dropin_table_style_flags_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_table_style_flags_case(
        _first_table_shape(fixture_id, prs, "table-style flags").table
    )


def _apply_wolfppt_dropin_table_style_flags_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_table_style_flags_case(
        _first_table_shape(fixture_id, prs, "table-style flags").table
    )


def _apply_table_style_flags_case(table: Any) -> None:
    for name, value in TABLE_STYLE_FLAGS_EXPECTED["flags"].items():
        setattr(table, name, value)
