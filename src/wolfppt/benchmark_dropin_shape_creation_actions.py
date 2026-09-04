"""Top-level shape creation drop-in benchmark edit actions."""

from __future__ import annotations

from typing import Any

from .benchmark_cases import (
    ADD_SHAPE_EXPECTED_TEXT,
    ADD_SHAPE_EXPECTED_TRANSFORM,
    ADD_TABLE_EXPECTED_TRANSFORM,
    ADD_TEXTBOX_EXPECTED_TEXT,
    ADD_TEXTBOX_EXPECTED_TRANSFORM,
)
from .benchmark_dropin_common_actions import _shape_collection_at
from .presentation import Presentation as WolfPresentation

TOP_LEVEL_CREATION_FIXTURES = (
    "text_basic/title_body_bullets",
    "workloads/mixed_real_world_deck",
    "workloads/mixed_deal_review_deck",
)


def _apply_python_pptx_dropin_add_shape(fixture_id: str, prs: Any) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-shape benchmark does not support fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_SHAPE

    shape = prs.slides[0].shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        ADD_SHAPE_EXPECTED_TRANSFORM["x"],
        ADD_SHAPE_EXPECTED_TRANSFORM["y"],
        ADD_SHAPE_EXPECTED_TRANSFORM["cx"],
        ADD_SHAPE_EXPECTED_TRANSFORM["cy"],
    )
    shape.text = ADD_SHAPE_EXPECTED_TEXT


def _apply_wolfppt_dropin_add_shape(fixture_id: str, prs: WolfPresentation) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-shape benchmark does not support fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_SHAPE

    shape = _shape_collection_at(prs, 0).add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        ADD_SHAPE_EXPECTED_TRANSFORM["x"],
        ADD_SHAPE_EXPECTED_TRANSFORM["y"],
        ADD_SHAPE_EXPECTED_TRANSFORM["cx"],
        ADD_SHAPE_EXPECTED_TRANSFORM["cy"],
    )
    shape.text = ADD_SHAPE_EXPECTED_TEXT


def _apply_python_pptx_dropin_add_textbox(fixture_id: str, prs: Any) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-textbox benchmark does not support fixture {fixture_id}"
        )
    shape = prs.slides[0].shapes.add_textbox(
        ADD_TEXTBOX_EXPECTED_TRANSFORM["x"],
        ADD_TEXTBOX_EXPECTED_TRANSFORM["y"],
        ADD_TEXTBOX_EXPECTED_TRANSFORM["cx"],
        ADD_TEXTBOX_EXPECTED_TRANSFORM["cy"],
    )
    shape.text = ADD_TEXTBOX_EXPECTED_TEXT


def _apply_wolfppt_dropin_add_textbox(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-textbox benchmark does not support fixture {fixture_id}"
        )
    shape = _shape_collection_at(prs, 0).add_textbox(
        ADD_TEXTBOX_EXPECTED_TRANSFORM["x"],
        ADD_TEXTBOX_EXPECTED_TRANSFORM["y"],
        ADD_TEXTBOX_EXPECTED_TRANSFORM["cx"],
        ADD_TEXTBOX_EXPECTED_TRANSFORM["cy"],
    )
    shape.text = ADD_TEXTBOX_EXPECTED_TEXT


def _apply_python_pptx_dropin_add_table(fixture_id: str, prs: Any) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-table benchmark does not support fixture {fixture_id}"
        )
    table = (
        prs.slides[0]
        .shapes.add_table(
            2,
            3,
            ADD_TABLE_EXPECTED_TRANSFORM["x"],
            ADD_TABLE_EXPECTED_TRANSFORM["y"],
            ADD_TABLE_EXPECTED_TRANSFORM["cx"],
            ADD_TABLE_EXPECTED_TRANSFORM["cy"],
        )
        .table
    )
    table.cell(0, 0).text = "Metric"


def _apply_wolfppt_dropin_add_table(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-table benchmark does not support fixture {fixture_id}"
        )
    table = (
        _shape_collection_at(prs, 0)
        .add_table(
            2,
            3,
            ADD_TABLE_EXPECTED_TRANSFORM["x"],
            ADD_TABLE_EXPECTED_TRANSFORM["y"],
            ADD_TABLE_EXPECTED_TRANSFORM["cx"],
            ADD_TABLE_EXPECTED_TRANSFORM["cy"],
        )
        .table
    )
    table.cell(0, 0).text = "Metric"


def _apply_python_pptx_dropin_add_slide(fixture_id: str, prs: Any) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-slide benchmark does not support fixture {fixture_id}"
        )
    prs.slides.add_slide(prs.slide_layouts[6])


def _apply_python_pptx_dropin_add_title_slide(fixture_id: str, prs: Any) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-title-slide benchmark does not support fixture {fixture_id}"
        )
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Pipeline Review"
    slide.placeholders[1].text = "Risks and next steps"


def _apply_wolfppt_dropin_add_slide(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-slide benchmark does not support fixture {fixture_id}"
        )
    prs.slides.add_slide(prs.slide_layouts[6])


def _apply_wolfppt_dropin_add_title_slide(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-title-slide benchmark does not support fixture {fixture_id}"
        )
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Pipeline Review"
    slide.placeholders[1].text = "Risks and next steps"
