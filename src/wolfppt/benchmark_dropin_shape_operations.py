"""Shape and media drop-in benchmark operation implementations."""
# ruff: noqa: F401

from __future__ import annotations

from pathlib import Path
from time import perf_counter
from typing import Any

from .benchmark_dropin_actions import (
    _apply_python_pptx_dropin_add_group_auto_shape,
    _apply_python_pptx_dropin_add_group_chart,
    _apply_python_pptx_dropin_add_group_shape,
    _apply_python_pptx_dropin_add_group_textbox,
    _apply_python_pptx_dropin_add_deeper_nested_group_auto_shape,
    _apply_python_pptx_dropin_add_deeper_nested_group_chart,
    _apply_python_pptx_dropin_add_deeper_nested_group_shape,
    _apply_python_pptx_dropin_add_deeper_nested_group_textbox,
    _apply_python_pptx_dropin_group_existing_children,
    _apply_python_pptx_dropin_group_existing_shapes,
    _apply_python_pptx_dropin_add_nested_group_auto_shape,
    _apply_python_pptx_dropin_add_nested_group_chart,
    _apply_python_pptx_dropin_add_nested_group_shape,
    _apply_python_pptx_dropin_add_nested_group_textbox,
    _apply_python_pptx_dropin_add_shape,
    _apply_python_pptx_dropin_add_slide,
    _apply_python_pptx_dropin_add_table,
    _apply_python_pptx_dropin_add_textbox,
    _apply_python_pptx_dropin_add_title_slide,
    _apply_python_pptx_dropin_clone_layout_placeholders,
    _apply_python_pptx_dropin_group_child_adjustment_edit,
    _apply_python_pptx_dropin_group_child_geometry_edit,
    _apply_python_pptx_dropin_picture_crop_edit,
    _apply_python_pptx_dropin_shape_line_element_edit,
    _apply_python_pptx_dropin_shape_adjustment_edit,
    _apply_python_pptx_dropin_shape_geometry_edit,
    _apply_wolfppt_dropin_add_group_auto_shape,
    _apply_wolfppt_dropin_add_group_chart,
    _apply_wolfppt_dropin_add_group_shape,
    _apply_wolfppt_dropin_add_group_textbox,
    _apply_wolfppt_dropin_add_deeper_nested_group_auto_shape,
    _apply_wolfppt_dropin_add_deeper_nested_group_chart,
    _apply_wolfppt_dropin_add_deeper_nested_group_shape,
    _apply_wolfppt_dropin_add_deeper_nested_group_textbox,
    _apply_wolfppt_dropin_group_existing_children,
    _apply_wolfppt_dropin_group_existing_shapes,
    _apply_wolfppt_dropin_add_nested_group_auto_shape,
    _apply_wolfppt_dropin_add_nested_group_chart,
    _apply_wolfppt_dropin_add_nested_group_shape,
    _apply_wolfppt_dropin_add_nested_group_textbox,
    _apply_wolfppt_dropin_add_shape,
    _apply_wolfppt_dropin_add_slide,
    _apply_wolfppt_dropin_add_table,
    _apply_wolfppt_dropin_add_textbox,
    _apply_wolfppt_dropin_add_title_slide,
    _apply_wolfppt_dropin_clone_layout_placeholders,
    _apply_wolfppt_dropin_group_child_adjustment_edit,
    _apply_wolfppt_dropin_group_child_geometry_edit,
    _apply_wolfppt_dropin_picture_crop_edit,
    _apply_wolfppt_dropin_shape_line_element_edit,
    _apply_wolfppt_dropin_shape_adjustment_edit,
    _apply_wolfppt_dropin_shape_geometry_edit,
)
from .benchmark_dropin_shape_basic_operations import (
    _bench_python_pptx_dropin_add_shape,
    _bench_python_pptx_dropin_add_slide,
    _bench_python_pptx_dropin_add_table,
    _bench_python_pptx_dropin_add_textbox,
    _bench_python_pptx_dropin_add_title_slide,
    _bench_python_pptx_dropin_group_child_adjustment_edit,
    _bench_python_pptx_dropin_group_child_geometry_edit,
    _bench_python_pptx_dropin_shape_adjustment_edit,
    _bench_python_pptx_dropin_shape_geometry_edit,
    _bench_python_pptx_dropin_shape_line_element_edit,
    _bench_wolfppt_facade_dropin_add_shape,
    _bench_wolfppt_facade_dropin_add_slide,
    _bench_wolfppt_facade_dropin_add_table,
    _bench_wolfppt_facade_dropin_add_textbox,
    _bench_wolfppt_facade_dropin_add_title_slide,
    _bench_wolfppt_facade_dropin_group_child_adjustment_edit,
    _bench_wolfppt_facade_dropin_group_child_geometry_edit,
    _bench_wolfppt_facade_dropin_shape_adjustment_edit,
    _bench_wolfppt_facade_dropin_shape_geometry_edit,
    _bench_wolfppt_facade_dropin_shape_line_element_edit,
)
from .benchmark_dropin_shape_freeform_operations import (
    _bench_python_pptx_dropin_build_deeper_nested_group_freeform,
    _bench_python_pptx_dropin_build_freeform,
    _bench_python_pptx_dropin_build_group_freeform,
    _bench_python_pptx_dropin_build_nested_group_freeform,
    _bench_wolfppt_facade_dropin_build_deeper_nested_group_freeform,
    _bench_wolfppt_facade_dropin_build_freeform,
    _bench_wolfppt_facade_dropin_build_group_freeform,
    _bench_wolfppt_facade_dropin_build_nested_group_freeform,
)
from .benchmark_dropin_shape_format_operations import (
    _bench_python_pptx_dropin_shape_fill_solid_edit,
    _bench_python_pptx_dropin_shape_gradient_fill_edit,
    _bench_python_pptx_dropin_shape_hyperlink_edit,
    _bench_python_pptx_dropin_shape_line_fill_background_edit,
    _bench_python_pptx_dropin_shape_name_edit,
    _bench_python_pptx_dropin_shape_patterned_fill_edit,
    _bench_python_pptx_dropin_shape_rotation_edit,
    _bench_python_pptx_dropin_shape_shadow_edit,
    _bench_python_pptx_dropin_shape_style_edit,
    _bench_python_pptx_dropin_shape_target_new_slide_edit,
    _bench_python_pptx_dropin_shape_target_slide_edit,
    _bench_python_pptx_dropin_shape_theme_color_edit,
    _bench_wolfppt_facade_dropin_shape_fill_solid_edit,
    _bench_wolfppt_facade_dropin_shape_gradient_fill_edit,
    _bench_wolfppt_facade_dropin_shape_hyperlink_edit,
    _bench_wolfppt_facade_dropin_shape_line_fill_background_edit,
    _bench_wolfppt_facade_dropin_shape_name_edit,
    _bench_wolfppt_facade_dropin_shape_patterned_fill_edit,
    _bench_wolfppt_facade_dropin_shape_rotation_edit,
    _bench_wolfppt_facade_dropin_shape_shadow_edit,
    _bench_wolfppt_facade_dropin_shape_style_edit,
    _bench_wolfppt_facade_dropin_shape_target_new_slide_edit,
    _bench_wolfppt_facade_dropin_shape_target_slide_edit,
    _bench_wolfppt_facade_dropin_shape_theme_color_edit,
)
from .benchmark_dropin_shape_chart_operations import (
    _bench_python_pptx_dropin_add_bar_chart,
    _bench_python_pptx_dropin_add_bubble_3d_chart,
    _bench_python_pptx_dropin_add_bubble_chart,
    _bench_python_pptx_dropin_add_chart,
    _bench_python_pptx_dropin_add_chart_template_family,
    _bench_python_pptx_dropin_add_hierarchical_chart,
    _bench_python_pptx_dropin_add_line_chart,
    _bench_python_pptx_dropin_add_pie_chart,
    _bench_python_pptx_dropin_add_xy_scatter_chart,
    _bench_python_pptx_dropin_add_xy_scatter_template_family,
    _bench_wolfppt_facade_dropin_add_bar_chart,
    _bench_wolfppt_facade_dropin_add_bubble_3d_chart,
    _bench_wolfppt_facade_dropin_add_bubble_chart,
    _bench_wolfppt_facade_dropin_add_chart,
    _bench_wolfppt_facade_dropin_add_chart_template_family,
    _bench_wolfppt_facade_dropin_add_hierarchical_chart,
    _bench_wolfppt_facade_dropin_add_line_chart,
    _bench_wolfppt_facade_dropin_add_pie_chart,
    _bench_wolfppt_facade_dropin_add_xy_scatter_chart,
    _bench_wolfppt_facade_dropin_add_xy_scatter_template_family,
)
from .benchmark_dropin_shape_connector_operations import (
    _bench_python_pptx_dropin_add_connector,
    _bench_python_pptx_dropin_add_deeper_nested_group_connector,
    _bench_python_pptx_dropin_add_group_connector,
    _bench_python_pptx_dropin_add_nested_group_connector,
    _bench_python_pptx_dropin_connector_connection_edit,
    _bench_python_pptx_dropin_connector_line_style_edit,
    _bench_python_pptx_dropin_existing_connector_connection_edit,
    _bench_python_pptx_dropin_group_connector_connection_edit,
    _bench_wolfppt_facade_dropin_add_connector,
    _bench_wolfppt_facade_dropin_add_deeper_nested_group_connector,
    _bench_wolfppt_facade_dropin_add_group_connector,
    _bench_wolfppt_facade_dropin_add_nested_group_connector,
    _bench_wolfppt_facade_dropin_connector_connection_edit,
    _bench_wolfppt_facade_dropin_connector_line_style_edit,
    _bench_wolfppt_facade_dropin_existing_connector_connection_edit,
    _bench_wolfppt_facade_dropin_group_connector_connection_edit,
)
from .benchmark_dropin_group_shape_operations import (
    _bench_python_pptx_dropin_add_deeper_nested_group_auto_shape,
    _bench_python_pptx_dropin_add_deeper_nested_group_chart,
    _bench_python_pptx_dropin_add_deeper_nested_group_shape,
    _bench_python_pptx_dropin_add_deeper_nested_group_textbox,
    _bench_python_pptx_dropin_add_group_auto_shape,
    _bench_python_pptx_dropin_add_group_chart,
    _bench_python_pptx_dropin_add_group_shape,
    _bench_python_pptx_dropin_add_group_textbox,
    _bench_python_pptx_dropin_add_nested_group_auto_shape,
    _bench_python_pptx_dropin_add_nested_group_chart,
    _bench_python_pptx_dropin_add_nested_group_shape,
    _bench_python_pptx_dropin_add_nested_group_textbox,
    _bench_python_pptx_dropin_group_existing_children,
    _bench_python_pptx_dropin_group_existing_shapes,
    _bench_wolfppt_facade_dropin_add_deeper_nested_group_auto_shape,
    _bench_wolfppt_facade_dropin_add_deeper_nested_group_chart,
    _bench_wolfppt_facade_dropin_add_deeper_nested_group_shape,
    _bench_wolfppt_facade_dropin_add_deeper_nested_group_textbox,
    _bench_wolfppt_facade_dropin_add_group_auto_shape,
    _bench_wolfppt_facade_dropin_add_group_chart,
    _bench_wolfppt_facade_dropin_add_group_shape,
    _bench_wolfppt_facade_dropin_add_group_textbox,
    _bench_wolfppt_facade_dropin_add_nested_group_auto_shape,
    _bench_wolfppt_facade_dropin_add_nested_group_chart,
    _bench_wolfppt_facade_dropin_add_nested_group_shape,
    _bench_wolfppt_facade_dropin_add_nested_group_textbox,
    _bench_wolfppt_facade_dropin_group_existing_children,
    _bench_wolfppt_facade_dropin_group_existing_shapes,
)
from .benchmark_dropin_shape_details import (
    _dropin_add_group_auto_shape_details,
    _dropin_add_group_chart_details,
    _dropin_add_group_shape_details,
    _dropin_add_group_textbox_details,
    _dropin_add_deeper_nested_group_auto_shape_details,
    _dropin_add_deeper_nested_group_chart_details,
    _dropin_add_deeper_nested_group_shape_details,
    _dropin_add_deeper_nested_group_textbox_details,
    _dropin_group_existing_children_details,
    _dropin_group_existing_shapes_details,
    _dropin_add_nested_group_auto_shape_details,
    _dropin_add_nested_group_chart_details,
    _dropin_add_nested_group_shape_details,
    _dropin_add_nested_group_textbox_details,
    _dropin_add_shape_details,
    _dropin_add_slide_details,
    _dropin_add_table_details,
    _dropin_add_textbox_details,
    _dropin_clone_layout_placeholders_details,
    _dropin_group_child_adjustment_edit_details,
    _dropin_group_child_geometry_edit_details,
    _dropin_picture_crop_edit_details,
    _dropin_shape_line_element_edit_details,
    _dropin_shape_adjustment_edit_details,
    _dropin_shape_geometry_edit_details,
)
from .benchmark_runtime import elapsed_ms as _elapsed_ms, stamp as _stamp
from .presentation import Presentation as WolfPresentation


def _bench_python_pptx_dropin_clone_layout_placeholders(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-clone-layout-placeholders-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_clone_layout_placeholders(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_clone_layout_placeholders_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_clone_layout_placeholders(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-clone-layout-placeholders-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_clone_layout_placeholders(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_clone_layout_placeholders_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_picture_crop_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-picture-crop-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_picture_crop_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_picture_crop_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_picture_crop_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-picture-crop-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_picture_crop_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_picture_crop_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )
