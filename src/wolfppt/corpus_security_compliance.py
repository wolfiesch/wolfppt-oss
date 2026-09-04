"""Security and compliance review workload fixture builder."""

from __future__ import annotations

from io import BytesIO


def add_security_compliance_review_pack(prs, slide) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches

    _add_textbox(
        slide,
        "Security and Compliance Review",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    _add_textbox(
        slide,
        "SOC 2 readiness, vendor-risk posture, incident drills, and control owners.",
        Inches(0.7),
        Inches(1.05),
        Inches(8.2),
        Inches(0.8),
    )
    slide.shapes.add_picture(
        BytesIO(_png_1x1()),
        Inches(8.85),
        Inches(0.35),
        Inches(0.4),
        Inches(0.4),
    )
    slide.notes_slide.notes_text_frame.text = (
        "Synthetic security review pack for real-corpus preservation evidence."
    )

    controls_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        controls_slide,
        "Control Readiness Register",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    controls = controls_slide.shapes.add_table(
        6,
        5,
        Inches(0.45),
        Inches(1.0),
        Inches(9.1),
        Inches(3.0),
    ).table
    for row_idx, row in enumerate(
        [
            ["Control", "Owner", "Evidence", "Status", "Next step"],
            ["Access review", "IT", "access_review.xlsx", "Green", "Archive"],
            ["Vendor SOC reports", "GRC", "vendor_soc.zip", "Yellow", "Chase"],
            ["Incident tabletop", "SecOps", "tabletop_notes.docx", "Green", "Retest"],
            ["Backup restore", "Infra", "restore_log.csv", "Yellow", "Run drill"],
            ["Change approvals", "Eng", "change_tickets.csv", "Green", "Sample"],
        ]
    ):
        for col_idx, value in enumerate(row):
            controls.cell(row_idx, col_idx).text = value
    controls_slide.notes_slide.notes_text_frame.text = (
        "Control table should preserve evidence filenames and owner status."
    )

    posture_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        posture_slide,
        "Risk Posture Trend",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    chart_data = CategoryChartData()
    chart_data.categories = ["Jan", "Feb", "Mar", "Apr"]
    chart_data.add_series("Open risks", (34, 29, 24, 18))
    chart_data.add_series("Remediated", (12, 16, 21, 27))
    posture_slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.7),
        Inches(1.0),
        Inches(8.4),
        Inches(4.0),
        chart_data,
    )
    posture_slide.notes_slide.notes_text_frame.text = (
        "Line chart workbook and security-risk notes must survive round-trip."
    )

    vendor_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        vendor_slide,
        "Vendor Risk Heatmap",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    vendor_table = vendor_slide.shapes.add_table(
        5,
        4,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(2.45),
    ).table
    for row_idx, row in enumerate(
        [
            ["Vendor", "Data", "Criticality", "Review"],
            ["PayrollCo", "PII", "High", "SOC report pending"],
            ["EmailOps", "Email", "High", "DPA refreshed"],
            ["DataRoom", "Confidential", "Medium", "Pen test due"],
            ["SupportDesk", "Tickets", "Medium", "Subprocessors reviewed"],
        ]
    ):
        for col_idx, value in enumerate(row):
            vendor_table.cell(row_idx, col_idx).text = value
    vendor_slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(1.0),
        Inches(3.75),
        Inches(8.4),
        Inches(3.75),
    )
    vendor_slide.notes_slide.notes_text_frame.text = (
        "Vendor-risk table and connector line should be preserved."
    )

    actions_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        actions_slide,
        "Remediation Actions",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    group = actions_slide.shapes.add_group_shape()
    for idx, (label, owner) in enumerate(
        (
            ("Evidence lockbox", "GRC"),
            ("Restore drill", "Infra"),
            ("Vendor chase", "Legal"),
            ("Access sign-off", "IT"),
        )
    ):
        tile = group.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.25 + (idx % 2) * 2.85),
            Inches(0.25 + (idx // 2) * 1.05),
            Inches(2.5),
            Inches(0.78),
        )
        tile.text = f"{label}\n{owner}"
    actions_slide.notes_slide.notes_text_frame.text = (
        "Grouped remediation action tiles should keep child text and transforms."
    )

    appendix_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        appendix_slide,
        "Appendix: Audit Evidence Index",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    appendix = appendix_slide.shapes.add_table(
        5,
        3,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(2.6),
    ).table
    for row_idx, row in enumerate(
        [
            ["Artifact", "Source", "Owner"],
            ["Risk register", "security_risks.xlsx", "GRC"],
            ["Vendor packet", "vendor_reviews.zip", "Legal"],
            ["Incident drill", "tabletop_minutes.docx", "SecOps"],
            ["Access sample", "user_access_sample.csv", "IT"],
        ]
    ):
        for col_idx, value in enumerate(row):
            appendix.cell(row_idx, col_idx).text = value


def _add_textbox(slide, text: str, left, top, width, height):
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.text = text
    return shape


def _png_1x1() -> bytes:
    return (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
        b"\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde\x00"
        b"\x00\x00\x0cIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02"
        b"\xfeA\xde\xfc\x8d\x00\x00\x00\x00IEND\xaeB`\x82"
    )
