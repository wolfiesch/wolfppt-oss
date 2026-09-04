"""Product-launch readiness workload fixture builder."""

from __future__ import annotations

from io import BytesIO


def add_product_launch_readiness_pack(prs, slide) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches

    _add_textbox(
        slide,
        "Product Launch Readiness Pack",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    _add_textbox(
        slide,
        "Milestones, launch risks, demand signals, and cross-functional owners.",
        Inches(0.7),
        Inches(1.05),
        Inches(8.2),
        Inches(0.8),
    )
    slide.shapes.add_picture(
        BytesIO(_png_1x1()),
        Inches(8.85),
        Inches(0.35),
        Inches(0.4),
        Inches(0.4),
    )
    slide.notes_slide.notes_text_frame.text = (
        "Synthetic product launch readiness pack for real-corpus preservation evidence."
    )

    milestones_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        milestones_slide,
        "Launch Milestone Tracker",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    milestones = milestones_slide.shapes.add_table(
        6,
        5,
        Inches(0.45),
        Inches(1.0),
        Inches(9.1),
        Inches(3.0),
    ).table
    for row_idx, row in enumerate(
        [
            ["Milestone", "Owner", "Status", "Due", "Blocker"],
            ["Beta feedback", "Product", "Green", "May 22", "None"],
            ["Pricing page", "Marketing", "Yellow", "May 24", "Copy review"],
            ["Enablement deck", "Sales", "Green", "May 28", "None"],
            ["Support runbook", "Support", "Yellow", "May 30", "Triage paths"],
            ["Launch checklist", "PMO", "Green", "Jun 03", "Final sign-off"],
        ]
    ):
        for col_idx, value in enumerate(row):
            milestones.cell(row_idx, col_idx).text = value
    milestones_slide.notes_slide.notes_text_frame.text = (
        "Milestone tracker should preserve launch owners, dates, and blocker text."
    )

    demand_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        demand_slide,
        "Demand Signal Trend",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    chart_data = CategoryChartData()
    chart_data.categories = ["Feb", "Mar", "Apr", "May"]
    chart_data.add_series("Waitlist", (1200, 1840, 2610, 3450))
    chart_data.add_series("Beta active", (210, 360, 620, 910))
    demand_slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.7),
        Inches(1.0),
        Inches(8.4),
        Inches(4.0),
        chart_data,
    )
    demand_slide.notes_slide.notes_text_frame.text = (
        "Demand chart workbook and launch-readiness notes must survive round-trip."
    )

    risks_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        risks_slide,
        "Launch Risk Register",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    risks = risks_slide.shapes.add_table(
        5,
        4,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(2.45),
    ).table
    for row_idx, row in enumerate(
        [
            ["Risk", "Impact", "Owner", "Mitigation"],
            ["Docs gap", "Medium", "Product Marketing", "Ship quickstart"],
            ["Support load", "High", "Support", "Create response macros"],
            ["Integration bug", "High", "Engineering", "Patch candidate"],
            ["Pricing confusion", "Medium", "Sales Ops", "FAQ and training"],
        ]
    ):
        for col_idx, value in enumerate(row):
            risks.cell(row_idx, col_idx).text = value
    risks_slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(1.0),
        Inches(3.75),
        Inches(8.4),
        Inches(3.75),
    )
    risks_slide.notes_slide.notes_text_frame.text = (
        "Risk-register table and connector line should be preserved."
    )

    owners_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        owners_slide,
        "Launch War Room Owners",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    group = owners_slide.shapes.add_group_shape()
    for idx, (label, owner) in enumerate(
        (
            ("Launch comms", "Marketing"),
            ("Beta triage", "Product"),
            ("Revenue desk", "Sales"),
            ("Support desk", "CX"),
        )
    ):
        tile = group.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.25 + (idx % 2) * 2.85),
            Inches(0.25 + (idx // 2) * 1.05),
            Inches(2.5),
            Inches(0.78),
        )
        tile.text = f"{label}\n{owner}"
    owners_slide.notes_slide.notes_text_frame.text = (
        "Grouped launch-owner tiles should keep child text and transforms."
    )

    appendix_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        appendix_slide,
        "Appendix: Launch Evidence Index",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    appendix = appendix_slide.shapes.add_table(
        5,
        3,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(2.6),
    ).table
    for row_idx, row in enumerate(
        [
            ["Artifact", "Source", "Owner"],
            ["Beta feedback", "beta_feedback.xlsx", "Product"],
            ["Launch checklist", "launch_checklist.csv", "PMO"],
            ["Support runbook", "support_runbook.docx", "Support"],
            ["Demand snapshot", "waitlist_export.csv", "Marketing"],
        ]
    ):
        for col_idx, value in enumerate(row):
            appendix.cell(row_idx, col_idx).text = value


def _add_textbox(slide, text: str, left, top, width, height):
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.text = text
    return shape


def _png_1x1() -> bytes:
    return (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
        b"\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde\x00"
        b"\x00\x00\x0cIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02"
        b"\xfeA\xde\xfc\x8d\x00\x00\x00\x00IEND\xaeB`\x82"
    )
