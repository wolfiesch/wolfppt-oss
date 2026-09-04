"""Nested-group drop-in benchmark edit actions."""

from __future__ import annotations

from typing import Any

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from .benchmark_dropin_common_actions import _shape_at
from .presentation import Presentation as WolfPresentation

GROUP_EXISTING_CHILD_TARGETS = {
    "shapes/grouped_shapes": (0, 0),
    "workloads/management_reporting_deck": (2, 1),
    "workloads/customer_success_review_pack": (4, 1),
}


def _dropin_group_shape(fixture_id: str, prs: Any, operation: str) -> Any:
    if fixture_id not in GROUP_EXISTING_CHILD_TARGETS:
        raise RuntimeError(
            f"drop-in {operation} benchmark does not support fixture {fixture_id}"
        )
    slide_index, group_index = GROUP_EXISTING_CHILD_TARGETS[fixture_id]
    if isinstance(prs, WolfPresentation):
        return _shape_at(prs, slide_index, group_index)
    return prs.slides[slide_index].shapes[group_index]


def _apply_python_pptx_dropin_group_existing_nested_children(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in GROUP_EXISTING_CHILD_TARGETS:
        raise RuntimeError(
            "drop-in group-existing-nested-children benchmark does not support "
            f"fixture {fixture_id}"
        )
    slide_index, group_index = GROUP_EXISTING_CHILD_TARGETS[fixture_id]
    nested = prs.slides[slide_index].shapes[group_index].shapes.add_group_shape()
    first = nested.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.1), Inches(0.1), Inches(1.0), Inches(0.5)
    )
    first.text = "First"
    second = nested.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(1.3), Inches(0.1), Inches(1.0), Inches(0.5)
    )
    second.text = "Second"
    nested.shapes.add_group_shape([second, first])


def _apply_wolfppt_dropin_group_existing_nested_children(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    nested = _dropin_group_shape(
        fixture_id, prs, "group-existing-nested-children"
    ).shapes.add_group_shape()
    first = nested.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.1), Inches(0.1), Inches(1.0), Inches(0.5)
    )
    first.text = "First"
    second = nested.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(1.3), Inches(0.1), Inches(1.0), Inches(0.5)
    )
    second.text = "Second"
    nested.shapes.add_group_shape([second, first])


def _apply_python_pptx_dropin_group_existing_deeper_nested_children(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in GROUP_EXISTING_CHILD_TARGETS:
        raise RuntimeError(
            "drop-in group-existing-deeper-nested-children benchmark does not "
            f"support fixture {fixture_id}"
        )
    slide_index, group_index = GROUP_EXISTING_CHILD_TARGETS[fixture_id]
    deeper = (
        prs.slides[slide_index]
        .shapes[group_index]
        .shapes.add_group_shape()
        .shapes.add_group_shape()
    )
    first = deeper.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.1), Inches(0.1), Inches(1.0), Inches(0.5)
    )
    first.text = "First"
    second = deeper.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(1.3), Inches(0.1), Inches(1.0), Inches(0.5)
    )
    second.text = "Second"
    deeper.shapes.add_group_shape([second, first])


def _apply_wolfppt_dropin_group_existing_deeper_nested_children(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    deeper = (
        _dropin_group_shape(fixture_id, prs, "group-existing-deeper-nested-children")
        .shapes.add_group_shape()
        .shapes.add_group_shape()
    )
    first = deeper.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.1), Inches(0.1), Inches(1.0), Inches(0.5)
    )
    first.text = "First"
    second = deeper.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(1.3), Inches(0.1), Inches(1.0), Inches(0.5)
    )
    second.text = "Second"
    deeper.shapes.add_group_shape([second, first])
