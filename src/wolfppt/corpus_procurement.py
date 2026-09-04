"""Procurement vendor-risk workload fixture builder."""

from __future__ import annotations

from io import BytesIO


def add_procurement_vendor_risk_pack(prs, slide) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches

    _add_textbox(
        slide,
        "Procurement Vendor Risk Pack",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    _add_textbox(
        slide,
        "Supplier exposure, renewal timing, contract concentration, and owner actions.",
        Inches(0.7),
        Inches(1.05),
        Inches(8.2),
        Inches(0.8),
    )
    slide.shapes.add_picture(
        BytesIO(_png_1x1()),
        Inches(8.85),
        Inches(0.35),
        Inches(0.4),
        Inches(0.4),
    )
    slide.notes_slide.notes_text_frame.text = (
        "Synthetic procurement vendor risk pack for real-corpus preservation evidence."
    )

    exposure_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        exposure_slide,
        "Supplier Exposure Register",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    exposure = exposure_slide.shapes.add_table(
        6,
        5,
        Inches(0.45),
        Inches(1.0),
        Inches(9.1),
        Inches(3.0),
    ).table
    for row_idx, row in enumerate(
        [
            ["Supplier", "Category", "Spend", "Risk", "Owner"],
            ["Northstar Cloud", "Infrastructure", "$4.8m", "High", "IT"],
            ["Apex Logistics", "Fulfillment", "$2.6m", "Medium", "Ops"],
            ["LedgerWorks", "Finance systems", "$1.4m", "Medium", "Finance"],
            ["Beacon Support", "BPO", "$3.1m", "High", "CX"],
            ["Total tracked", "Critical vendors", "$11.9m", "Watch", "Procurement"],
        ]
    ):
        for col_idx, value in enumerate(row):
            exposure.cell(row_idx, col_idx).text = value
    exposure_slide.notes_slide.notes_text_frame.text = (
        "Supplier exposure table should preserve vendor names, categories, and risk owners."
    )

    concentration_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        concentration_slide,
        "Spend Concentration Trend",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2", "Q3", "Q4"]
    chart_data.add_series("Top 5 vendors", (38, 41, 43, 46))
    chart_data.add_series("Top 10 vendors", (57, 60, 63, 66))
    concentration_slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.7),
        Inches(1.0),
        Inches(8.4),
        Inches(4.0),
        chart_data,
    )
    concentration_slide.notes_slide.notes_text_frame.text = (
        "Spend concentration chart workbook and notes must survive round-trip."
    )

    renewals_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        renewals_slide,
        "Renewal and Contract Risk",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    renewals = renewals_slide.shapes.add_table(
        5,
        4,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(2.45),
    ).table
    for row_idx, row in enumerate(
        [
            ["Contract", "Renewal", "Issue", "Mitigation"],
            ["Cloud commit", "Jun 30", "Overage exposure", "Usage true-up"],
            ["Warehouse 3PL", "Jul 15", "Capacity cap", "Secondary lane"],
            ["Support BPO", "Aug 01", "SLA credits", "Executive review"],
            ["Finance suite", "Sep 30", "Data residency", "Legal addendum"],
        ]
    ):
        for col_idx, value in enumerate(row):
            renewals.cell(row_idx, col_idx).text = value
    renewals_slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(1.0),
        Inches(3.75),
        Inches(8.4),
        Inches(3.75),
    )
    renewals_slide.notes_slide.notes_text_frame.text = (
        "Renewal table and connector should be preserved as contract-risk evidence."
    )

    owners_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        owners_slide,
        "Mitigation Owners",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    group = owners_slide.shapes.add_group_shape()
    for idx, (label, owner) in enumerate(
        (
            ("Cloud usage true-up", "IT"),
            ("Warehouse backup lane", "Operations"),
            ("BPO SLA review", "Customer Success"),
            ("Data residency addendum", "Legal"),
        )
    ):
        tile = group.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.25 + (idx % 2) * 2.85),
            Inches(0.25 + (idx // 2) * 1.05),
            Inches(2.5),
            Inches(0.78),
        )
        tile.text = f"{label}\n{owner}"
    owners_slide.notes_slide.notes_text_frame.text = (
        "Grouped procurement action tiles should keep child text and transforms."
    )

    appendix_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        appendix_slide,
        "Appendix: Procurement Evidence Index",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    appendix = appendix_slide.shapes.add_table(
        5,
        3,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(2.6),
    ).table
    for row_idx, row in enumerate(
        [
            ["Artifact", "Source", "Owner"],
            ["Spend cube", "vendor_spend_cube.xlsx", "Procurement"],
            ["Contract tracker", "renewal_tracker.csv", "Legal Ops"],
            ["Risk register", "supplier_risk.csv", "Vendor Management"],
            ["Usage export", "cloud_usage_export.csv", "IT"],
        ]
    ):
        for col_idx, value in enumerate(row):
            appendix.cell(row_idx, col_idx).text = value


def _add_textbox(slide, text: str, left, top, width, height):
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.text = text
    return shape


def _png_1x1() -> bytes:
    return (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
        b"\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde\x00"
        b"\x00\x00\x0cIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02"
        b"\xfeA\xde\xfc\x8d\x00\x00\x00\x00IEND\xaeB`\x82"
    )
