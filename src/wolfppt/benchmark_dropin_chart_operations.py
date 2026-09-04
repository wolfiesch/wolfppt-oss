"""Chart drop-in benchmark operation implementations."""

from __future__ import annotations

from pathlib import Path
from time import perf_counter
from typing import Any

from .benchmark_cases import CHART_TITLE_EXPECTED_TEXT
from .benchmark_dropin_actions import (
    _apply_python_pptx_dropin_chart_axis_property_edit,
    _apply_python_pptx_dropin_chart_axis_title_edit,
    _apply_python_pptx_dropin_chart_axis_title_remove,
    _apply_python_pptx_dropin_chart_title_paragraph_font_edit,
    _apply_python_pptx_dropin_chart_title_paragraph_format_edit,
    _apply_python_pptx_dropin_chart_title_text_frame_edit,
    _apply_python_pptx_dropin_chart_title_text_frame_flow_edit,
    _apply_python_pptx_dropin_chart_data_label_edit,
    _apply_python_pptx_dropin_chart_data_label_remove,
    _apply_python_pptx_dropin_chart_format_edit,
    _apply_python_pptx_dropin_chart_legend_edit,
    _apply_python_pptx_dropin_chart_plot_property_edit,
    _apply_python_pptx_dropin_chart_point_data_label_edit,
    _apply_python_pptx_dropin_chart_style_edit,
    _apply_python_pptx_dropin_chart_title_edit,
    _apply_python_pptx_dropin_chart_title_remove,
    _apply_wolfppt_dropin_chart_axis_property_edit,
    _apply_wolfppt_dropin_chart_axis_title_edit,
    _apply_wolfppt_dropin_chart_axis_title_remove,
    _apply_wolfppt_dropin_chart_title_paragraph_font_edit,
    _apply_wolfppt_dropin_chart_title_paragraph_format_edit,
    _apply_wolfppt_dropin_chart_title_text_frame_edit,
    _apply_wolfppt_dropin_chart_title_text_frame_flow_edit,
    _apply_wolfppt_dropin_chart_data_label_edit,
    _apply_wolfppt_dropin_chart_data_label_remove,
    _apply_wolfppt_dropin_chart_format_edit,
    _apply_wolfppt_dropin_chart_legend_edit,
    _apply_wolfppt_dropin_chart_plot_property_edit,
    _apply_wolfppt_dropin_chart_point_data_label_edit,
    _apply_wolfppt_dropin_chart_style_edit,
    _apply_wolfppt_dropin_chart_title_edit,
    _apply_wolfppt_dropin_chart_title_remove,
    _read_python_pptx_chart_metadata,
    _read_wolfppt_chart_metadata,
)
from .benchmark_dropin_chart_details import (
    _dropin_chart_axis_property_edit_details,
    _dropin_chart_axis_title_edit_details,
    _dropin_chart_axis_title_remove_details,
    _dropin_chart_title_paragraph_font_edit_details,
    _dropin_chart_title_paragraph_format_edit_details,
    _dropin_chart_title_text_frame_edit_details,
    _dropin_chart_title_text_frame_flow_edit_details,
    _dropin_chart_data_label_edit_details,
    _dropin_chart_data_label_remove_details,
    _dropin_chart_format_edit_details,
    _dropin_chart_legend_edit_details,
    _dropin_chart_plot_property_edit_details,
    _dropin_chart_point_data_label_edit_details,
    _dropin_chart_read_details,
    _dropin_chart_style_edit_details,
    _dropin_chart_title_edit_details,
    _dropin_chart_title_remove_details,
)
from .benchmark_dropin_chart_readers import _first_chart_shape
from .benchmark_runtime import elapsed_ms as _elapsed_ms, stamp as _stamp
from .presentation import Presentation as WolfPresentation


def _chart_title_remove_source(fixture_path: Path, tmp_path: Path) -> Path:
    from pptx import Presentation as PythonPptxPresentation

    source = tmp_path / (
        f"dropin-chart-title-remove-source-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    prs = PythonPptxPresentation(str(fixture_path))
    chart = _first_chart_shape(prs).chart
    chart.chart_title.text_frame.text = CHART_TITLE_EXPECTED_TEXT
    prs.save(str(source))
    return source


def _chart_axis_title_remove_source(fixture_path: Path, tmp_path: Path) -> Path:
    from pptx import Presentation as PythonPptxPresentation

    source = tmp_path / (
        f"dropin-chart-axis-title-remove-source-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    prs = PythonPptxPresentation(str(fixture_path))
    chart = _first_chart_shape(prs).chart
    chart.category_axis.axis_title.text_frame.text = "Quarter"
    chart.value_axis.axis_title.text_frame.text = "Revenue"
    prs.save(str(source))
    return source


def _bench_python_pptx_dropin_chart_read(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    metadata = _read_python_pptx_chart_metadata(fixture_id, prs)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_read_details(fixture_id, metadata)


def _bench_wolfppt_facade_dropin_chart_read(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    metadata = _read_wolfppt_chart_metadata(fixture_id, prs)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_read_details(fixture_id, metadata)


def _bench_python_pptx_dropin_chart_title_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-chart-title-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_chart_title_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_title_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_chart_title_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-chart-title-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_chart_title_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_title_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_chart_title_remove(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    source = _chart_title_remove_source(fixture_path, tmp_path)
    out = tmp_path / (
        f"python-pptx-dropin-chart-title-remove-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(source))
    _apply_python_pptx_dropin_chart_title_remove(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_title_remove_details(
        fixture_id,
        source,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_chart_title_remove(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    source = _chart_title_remove_source(fixture_path, tmp_path)
    out = tmp_path / (
        f"wolfppt-dropin-chart-title-remove-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(source)
    _apply_wolfppt_dropin_chart_title_remove(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_title_remove_details(
        fixture_id,
        source,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_chart_legend_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-chart-legend-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_chart_legend_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_legend_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_chart_legend_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-chart-legend-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_chart_legend_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_legend_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_chart_axis_title_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-chart-axis-title-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_chart_axis_title_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_axis_title_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_chart_axis_title_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-chart-axis-title-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_chart_axis_title_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_axis_title_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_chart_axis_title_remove(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    source = _chart_axis_title_remove_source(fixture_path, tmp_path)
    out = tmp_path / (
        f"python-pptx-dropin-chart-axis-title-remove-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(source))
    _apply_python_pptx_dropin_chart_axis_title_remove(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_axis_title_remove_details(
        fixture_id,
        source,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_chart_axis_title_remove(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    source = _chart_axis_title_remove_source(fixture_path, tmp_path)
    out = tmp_path / (
        f"wolfppt-dropin-chart-axis-title-remove-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(source)
    _apply_wolfppt_dropin_chart_axis_title_remove(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_axis_title_remove_details(
        fixture_id,
        source,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_chart_title_text_frame_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        "python-pptx-dropin-chart-title-text-frame-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_chart_title_text_frame_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_title_text_frame_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_chart_title_text_frame_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        "wolfppt-dropin-chart-title-text-frame-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_chart_title_text_frame_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_title_text_frame_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_chart_title_text_frame_flow_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        "python-pptx-dropin-chart-title-text-frame-flow-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_chart_title_text_frame_flow_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_title_text_frame_flow_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_chart_title_text_frame_flow_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        "wolfppt-dropin-chart-title-text-frame-flow-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_chart_title_text_frame_flow_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_title_text_frame_flow_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_chart_title_paragraph_format_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        "python-pptx-dropin-chart-title-paragraph-format-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_chart_title_paragraph_format_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_title_paragraph_format_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_chart_title_paragraph_format_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        "wolfppt-dropin-chart-title-paragraph-format-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_chart_title_paragraph_format_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_title_paragraph_format_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_chart_title_paragraph_font_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        "python-pptx-dropin-chart-title-paragraph-font-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_chart_title_paragraph_font_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_title_paragraph_font_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_chart_title_paragraph_font_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        "wolfppt-dropin-chart-title-paragraph-font-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_chart_title_paragraph_font_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_title_paragraph_font_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_chart_axis_property_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-chart-axis-property-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_chart_axis_property_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_axis_property_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_chart_axis_property_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-chart-axis-property-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_chart_axis_property_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_axis_property_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_chart_plot_property_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-chart-plot-property-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_chart_plot_property_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_plot_property_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_chart_plot_property_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-chart-plot-property-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_chart_plot_property_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_plot_property_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_chart_data_label_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-chart-data-label-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_chart_data_label_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_data_label_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_chart_data_label_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-chart-data-label-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_chart_data_label_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_data_label_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_chart_data_label_remove(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-chart-data-label-remove-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_chart_data_label_remove(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_data_label_remove_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_chart_data_label_remove(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-chart-data-label-remove-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_chart_data_label_remove(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_data_label_remove_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_chart_point_data_label_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-chart-point-data-label-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_chart_point_data_label_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_point_data_label_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_chart_point_data_label_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-chart-point-data-label-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_chart_point_data_label_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_point_data_label_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_chart_style_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-chart-style-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_chart_style_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_style_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_chart_style_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-chart-style-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_chart_style_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_style_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_chart_format_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-chart-format-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_chart_format_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_format_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_chart_format_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-chart-format-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_chart_format_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_chart_format_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )
