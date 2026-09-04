"""Post-merger integration workload fixture builder."""

from __future__ import annotations

from io import BytesIO

from .corpus_slide_helpers import add_textbox as _add_textbox
from .corpus_slide_helpers import png_1x1 as _png_1x1


def add_post_merger_integration_pack(prs, slide) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches

    _add_textbox(
        slide,
        "Post-Merger Integration Pack",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    _add_textbox(
        slide,
        "Synergy tracking, Day 1 readiness, systems cutover, and owner actions.",
        Inches(0.7),
        Inches(1.05),
        Inches(8.2),
        Inches(0.8),
    )
    slide.shapes.add_picture(
        BytesIO(_png_1x1()),
        Inches(8.85),
        Inches(0.35),
        Inches(0.4),
        Inches(0.4),
    )
    slide.notes_slide.notes_text_frame.text = (
        "Synthetic post-merger integration pack for real-corpus preservation evidence."
    )

    readiness_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        readiness_slide,
        "Day 1 Readiness Tracker",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    readiness = readiness_slide.shapes.add_table(
        6,
        5,
        Inches(0.45),
        Inches(1.0),
        Inches(9.1),
        Inches(3.0),
    ).table
    for row_idx, row in enumerate(
        [
            ["Workstream", "Owner", "Status", "Due", "Risk"],
            ["Finance close", "Controller", "Green", "Day 1", "None"],
            ["HR onboarding", "People Ops", "Yellow", "Day 7", "Policy mapping"],
            ["Customer comms", "CX", "Green", "Day 3", "None"],
            ["IT access", "Security", "Yellow", "Day 5", "SSO exceptions"],
            ["Legal entities", "Legal", "Green", "Day 10", "Filings"],
        ]
    ):
        for col_idx, value in enumerate(row):
            readiness.cell(row_idx, col_idx).text = value
    readiness_slide.notes_slide.notes_text_frame.text = (
        "Readiness table should preserve workstream owner and risk text."
    )

    synergy_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        synergy_slide,
        "Synergy Capture Trend",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    chart_data = CategoryChartData()
    chart_data.categories = ["Month 1", "Month 2", "Month 3", "Month 4"]
    chart_data.add_series("Run-rate target", (1.0, 2.5, 4.0, 6.0))
    chart_data.add_series("Captured", (0.7, 1.8, 3.1, 4.9))
    synergy_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.7),
        Inches(1.0),
        Inches(8.4),
        Inches(4.0),
        chart_data,
    )
    synergy_slide.notes_slide.notes_text_frame.text = (
        "Synergy chart workbook and notes should survive round-trip."
    )

    cutover_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        cutover_slide,
        "Systems Cutover Plan",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    cutover = cutover_slide.shapes.add_table(
        5,
        4,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(2.45),
    ).table
    for row_idx, row in enumerate(
        [
            ["System", "Target", "Dependency", "Fallback"],
            ["ERP", "Month 2", "Chart mapping", "Dual close"],
            ["CRM", "Month 1", "Account merge", "Read-only archive"],
            ["Data warehouse", "Month 3", "Identity map", "Batch replay"],
            ["Support desk", "Day 14", "Routing rules", "Shared queue"],
        ]
    ):
        for col_idx, value in enumerate(row):
            cutover.cell(row_idx, col_idx).text = value
    cutover_slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(1.0),
        Inches(3.75),
        Inches(8.4),
        Inches(3.75),
    )
    cutover_slide.notes_slide.notes_text_frame.text = (
        "Cutover table and connector should remain intact."
    )

    owners_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        owners_slide,
        "Integration Owner Actions",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    group = owners_slide.shapes.add_group_shape()
    for idx, (label, owner) in enumerate(
        (
            ("Synergy model refresh", "Finance"),
            ("SSO exception cleanup", "Security"),
            ("Customer notice QA", "CX"),
            ("ERP cutover rehearsal", "IT"),
        )
    ):
        tile = group.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.25 + (idx % 2) * 2.85),
            Inches(0.25 + (idx // 2) * 1.05),
            Inches(2.5),
            Inches(0.78),
        )
        tile.text = f"{label}\n{owner}"
    owners_slide.notes_slide.notes_text_frame.text = (
        "Grouped integration action tiles should keep child text and transforms."
    )

    appendix_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        appendix_slide,
        "Appendix: Integration Evidence Index",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    appendix = appendix_slide.shapes.add_table(
        5,
        3,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(2.6),
    ).table
    for row_idx, row in enumerate(
        [
            ["Artifact", "Source", "Owner"],
            ["Synergy model", "synergy_model.xlsx", "Finance"],
            ["Cutover tracker", "systems_cutover.csv", "IT"],
            ["Customer comms", "customer_notice.docx", "CX"],
            ["Access exceptions", "sso_exceptions.csv", "Security"],
        ]
    ):
        for col_idx, value in enumerate(row):
            appendix.cell(row_idx, col_idx).text = value
