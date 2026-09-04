"""python-pptx baseline adapter."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any


@dataclass(frozen=True)
class PythonPptxSummary:
    path: str
    slide_count: int
    shape_count: int
    texts: list[str]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def summarize(path: str | Path) -> PythonPptxSummary:
    """Summarize a presentation through python-pptx's public API."""

    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise RuntimeError("python-pptx is not installed; run with the baseline extra") from exc

    deck_path = Path(path)
    prs = Presentation(str(deck_path))
    texts: list[str] = []
    shape_count = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            shape_count += 1
            if getattr(shape, "has_text_frame", False):
                text = shape.text
                if text:
                    texts.append(text)
    return PythonPptxSummary(
        path=str(deck_path),
        slide_count=len(prs.slides),
        shape_count=shape_count,
        texts=texts,
    )


def roundtrip(input_path: str | Path, output_path: str | Path) -> Path:
    """Open and save a deck through python-pptx."""

    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise RuntimeError("python-pptx is not installed; run with the baseline extra") from exc

    source = Path(input_path)
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(source))
    prs.save(str(out))
    return out

