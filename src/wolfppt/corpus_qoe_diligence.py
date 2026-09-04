"""Quality of Earnings diligence workload fixture builder."""

from __future__ import annotations

from io import BytesIO


def add_qoe_diligence_pack(prs, slide) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches

    _add_textbox(
        slide,
        "Quality of Earnings Diligence Pack",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    _add_textbox(
        slide,
        "Adjusted EBITDA, revenue quality, net working capital, and diligence findings.",
        Inches(0.7),
        Inches(1.05),
        Inches(8.2),
        Inches(0.8),
    )
    slide.shapes.add_picture(
        BytesIO(_png_1x1()),
        Inches(8.85),
        Inches(0.35),
        Inches(0.4),
        Inches(0.4),
    )
    slide.notes_slide.notes_text_frame.text = (
        "Synthetic QoE diligence pack for buyer-side real-corpus preservation evidence."
    )

    ebitda_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        ebitda_slide,
        "Adjusted EBITDA Bridge",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    table = ebitda_slide.shapes.add_table(
        7,
        5,
        Inches(0.45),
        Inches(1.0),
        Inches(9.1),
        Inches(3.2),
    ).table
    rows = [
        ["Adjustment", "FY22", "FY23", "LTM", "Evidence"],
        ["Reported EBITDA", "$6.8m", "$8.4m", "$9.1m", "TB tie-out"],
        ["Owner comp", "$0.4m", "$0.5m", "$0.5m", "Payroll"],
        ["One-time legal", "$0.2m", "$0.1m", "$0.0m", "Invoices"],
        ["Revenue cutoff", "($0.1m)", "($0.4m)", "($0.3m)", "Sales orders"],
        ["Run-rate savings", "$0.0m", "$0.3m", "$0.6m", "Headcount plan"],
        ["Adjusted EBITDA", "$7.3m", "$8.9m", "$9.9m", "QoE model"],
    ]
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            table.cell(row_idx, col_idx).text = value
    ebitda_slide.notes_slide.notes_text_frame.text = (
        "EBITDA bridge should retain table topology, notes, and evidence labels."
    )

    chart_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        chart_slide,
        "Revenue Quality by Channel",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    chart_data = CategoryChartData()
    chart_data.categories = ["Direct", "Partner", "Enterprise", "Services"]
    chart_data.add_series("Recurring", (18.4, 9.6, 22.1, 4.2))
    chart_data.add_series("Non-recurring", (1.1, 0.8, 2.4, 3.7))
    chart_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_STACKED,
        Inches(0.7),
        Inches(1.0),
        Inches(8.4),
        Inches(4.0),
        chart_data,
    )
    chart_slide.notes_slide.notes_text_frame.text = (
        "Stacked revenue-quality chart and workbook must survive round-trip."
    )

    nwc_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        nwc_slide,
        "Net Working Capital Peg",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    nwc = nwc_slide.shapes.add_table(
        6,
        4,
        Inches(0.65),
        Inches(1.0),
        Inches(8.4),
        Inches(2.8),
    ).table
    for row_idx, row in enumerate(
        [
            ["Scenario", "Target", "Actual", "Purchase price impact"],
            ["Base peg", "$5.4m", "$5.1m", "($0.3m)"],
            ["AR normalization", "$5.8m", "$5.1m", "($0.7m)"],
            ["Inventory reserve", "$5.4m", "$4.8m", "($0.6m)"],
            ["Deferred revenue", "$5.0m", "$5.1m", "$0.1m"],
            ["Buyer proposal", "$5.6m", "$5.1m", "($0.5m)"],
        ]
    ):
        for col_idx, value in enumerate(row):
            nwc.cell(row_idx, col_idx).text = value
    nwc_slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(1.0),
        Inches(4.05),
        Inches(8.5),
        Inches(4.05),
    )
    nwc_slide.notes_slide.notes_text_frame.text = (
        "NWC scenario table and connector line should be preserved."
    )

    findings_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        findings_slide,
        "Diligence Findings and Open Items",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    group = findings_slide.shapes.add_group_shape()
    for idx, (label, owner) in enumerate(
        (
            ("Revenue cutoff", "Accounting"),
            ("Customer churn", "Commercial"),
            ("Payroll accrual", "People"),
            ("Inventory reserve", "Ops"),
        )
    ):
        tile = group.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.25 + (idx % 2) * 2.85),
            Inches(0.25 + (idx // 2) * 1.05),
            Inches(2.5),
            Inches(0.78),
        )
        tile.text = f"{label}\n{owner}"
    findings_slide.notes_slide.notes_text_frame.text = (
        "Grouped diligence finding tiles should keep child text and transforms."
    )

    appendix_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        appendix_slide,
        "Appendix: QoE Evidence Index",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    appendix = appendix_slide.shapes.add_table(
        6,
        3,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(3.0),
    ).table
    for row_idx, row in enumerate(
        [
            ["Artifact", "Source", "Owner"],
            ["Trial balance", "trial_balance.xlsx", "Controller"],
            ["Revenue sample", "sales_orders.csv", "Accounting"],
            ["Payroll detail", "payroll_register.xlsx", "People"],
            ["Inventory roll-forward", "inventory_ledger.xlsx", "Ops"],
            ["QoE model", "qoe_adjustments.xlsx", "Deal team"],
        ]
    ):
        for col_idx, value in enumerate(row):
            appendix.cell(row_idx, col_idx).text = value


def _add_textbox(slide, text: str, left, top, width, height):
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.text = text
    return shape


def _png_1x1() -> bytes:
    return (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
        b"\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde\x00"
        b"\x00\x00\x0cIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02"
        b"\xfeA\xde\xfc\x8d\x00\x00\x00\x00IEND\xaeB`\x82"
    )
