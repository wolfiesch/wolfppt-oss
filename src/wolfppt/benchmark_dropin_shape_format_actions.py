"""Shape formatting drop-in benchmark edit actions."""

from __future__ import annotations

from typing import Any

from .benchmark_cases import (
    SHAPE_HYPERLINK_EXPECTED_ADDRESS,
    SHAPE_NAME_EXPECTED,
    SHAPE_ROTATION_EXPECTED,
    SHAPE_SHADOW_EXPECTED_INHERIT,
)
from .benchmark_dropin_common_actions import _shape_at
from .presentation import Presentation as WolfPresentation

TOP_LEVEL_CREATION_FIXTURES = (
    "text_basic/title_body_bullets",
    "workloads/mixed_real_world_deck",
    "workloads/mixed_deal_review_deck",
)
TARGET_SLIDE_FIXTURES = (
    "slides/two_slide_text",
    "workloads/mixed_real_world_deck",
    "workloads/mixed_deal_review_deck",
)


def _dropin_shape_or_group_child(fixture_id: str, prs: Any, operation: str) -> Any:
    if fixture_id == "text_basic/title_body_bullets":
        if isinstance(prs, WolfPresentation):
            return _shape_at(prs, 0, 0)
        return prs.slides[0].shapes[0]
    if fixture_id == "shapes/grouped_shapes":
        if isinstance(prs, WolfPresentation):
            return _shape_at(prs, 0, 0).shapes[0]
        return prs.slides[0].shapes[0].shapes[0]
    raise RuntimeError(
        f"drop-in {operation} benchmark does not support fixture {fixture_id}"
    )


def _apply_python_pptx_dropin_shape_style_edit(fixture_id: str, prs: Any) -> None:
    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape style")
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE

    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    shape.line.color.rgb = RGBColor(0x65, 0x43, 0x21)
    shape.line.width = 28575
    shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def _apply_wolfppt_dropin_shape_style_edit(
    fixture_id: str, prs: WolfPresentation
) -> None:
    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape style")
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE

    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    shape.line.color.rgb = RGBColor(0x65, 0x43, 0x21)
    shape.line.width = 28575
    shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def _apply_python_pptx_dropin_shape_theme_color_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.enum.dml import MSO_THEME_COLOR

    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape theme color")
    shape.fill.solid()
    shape.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_2
    shape.fill.fore_color.brightness = 0.4
    shape.line.color.theme_color = MSO_THEME_COLOR.ACCENT_3
    shape.line.color.brightness = -0.25


def _apply_wolfppt_dropin_shape_theme_color_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.enum.dml import MSO_THEME_COLOR

    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape theme color")
    shape.fill.solid()
    shape.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_2
    shape.fill.fore_color.brightness = 0.4
    shape.line.color.theme_color = MSO_THEME_COLOR.ACCENT_3
    shape.line.color.brightness = -0.25


def _apply_python_pptx_dropin_shape_shadow_edit(fixture_id: str, prs: Any) -> None:
    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape shadow")
    shape.shadow.inherit = SHAPE_SHADOW_EXPECTED_INHERIT


def _apply_wolfppt_dropin_shape_shadow_edit(
    fixture_id: str, prs: WolfPresentation
) -> None:
    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape shadow")
    shape.shadow.inherit = SHAPE_SHADOW_EXPECTED_INHERIT


def _apply_python_pptx_dropin_shape_hyperlink_edit(fixture_id: str, prs: Any) -> None:
    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape hyperlink")
    shape.click_action.hyperlink.address = SHAPE_HYPERLINK_EXPECTED_ADDRESS


def _apply_wolfppt_dropin_shape_hyperlink_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape hyperlink")
    shape.click_action.hyperlink.address = SHAPE_HYPERLINK_EXPECTED_ADDRESS


def _apply_python_pptx_dropin_shape_target_slide_edit(
    fixture_id: str, prs: Any
) -> None:
    if fixture_id not in TARGET_SLIDE_FIXTURES:
        raise RuntimeError(
            f"drop-in shape target-slide benchmark does not support fixture {fixture_id}"
        )
    prs.slides[0].shapes[0].click_action.target_slide = prs.slides[1]


def _apply_wolfppt_dropin_shape_target_slide_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in TARGET_SLIDE_FIXTURES:
        raise RuntimeError(
            f"drop-in shape target-slide benchmark does not support fixture {fixture_id}"
        )
    _shape_at(prs, 0, 0).click_action.target_slide = prs.slides[1]


def _apply_python_pptx_dropin_shape_target_new_slide_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            "drop-in shape target-new-slide benchmark does not support fixture "
            f"{fixture_id}"
        )
    target = prs.slides.add_slide(prs.slide_layouts[1])
    prs.slides[0].shapes[0].click_action.target_slide = target


def _apply_wolfppt_dropin_shape_target_new_slide_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            "drop-in shape target-new-slide benchmark does not support fixture "
            f"{fixture_id}"
        )
    target = prs.slides.add_slide(prs.slide_layouts[1])
    _shape_at(prs, 0, 0).click_action.target_slide = target


def _apply_python_pptx_dropin_shape_fill_solid_edit(fixture_id: str, prs: Any) -> None:
    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape solid-fill")
    shape.fill.solid()


def _apply_wolfppt_dropin_shape_fill_solid_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape solid-fill")
    shape.fill.solid()


def _apply_python_pptx_dropin_shape_line_fill_background_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.dml.color import RGBColor

    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape line background-fill")
    shape.line.color.rgb = RGBColor(0x65, 0x43, 0x21)
    shape.line.width = 28575
    shape.line.fill.background()


def _apply_wolfppt_dropin_shape_line_fill_background_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.dml.color import RGBColor

    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape line background-fill")
    shape.line.color.rgb = RGBColor(0x65, 0x43, 0x21)
    shape.line.width = 28575
    shape.line.fill.background()


def _apply_python_pptx_dropin_shape_patterned_fill_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_PATTERN_TYPE

    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape patterned-fill")
    shape.fill.patterned()
    shape.fill.pattern = MSO_PATTERN_TYPE.DIVOT
    shape.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    shape.fill.back_color.rgb = RGBColor(0x65, 0x43, 0x21)
    shape.line.fill.patterned()
    shape.line.fill.pattern = MSO_PATTERN_TYPE.WAVE
    shape.line.fill.fore_color.rgb = RGBColor(0xAB, 0xCD, 0xEF)
    shape.line.fill.back_color.rgb = RGBColor(0xFE, 0xDC, 0xBA)


def _apply_wolfppt_dropin_shape_patterned_fill_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_PATTERN_TYPE

    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape patterned-fill")
    shape.fill.patterned()
    shape.fill.pattern = MSO_PATTERN_TYPE.DIVOT
    shape.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    shape.fill.back_color.rgb = RGBColor(0x65, 0x43, 0x21)
    shape.line.fill.patterned()
    shape.line.fill.pattern = MSO_PATTERN_TYPE.WAVE
    shape.line.fill.fore_color.rgb = RGBColor(0xAB, 0xCD, 0xEF)
    shape.line.fill.back_color.rgb = RGBColor(0xFE, 0xDC, 0xBA)


def _apply_python_pptx_dropin_shape_gradient_fill_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.dml.color import RGBColor

    fill = _dropin_shape_or_group_child(fixture_id, prs, "shape gradient-fill").fill
    fill.gradient()
    fill.gradient_angle = 45.0
    fill.gradient_stops[0].position = 0.25
    fill.gradient_stops[0].color.rgb = RGBColor(0x12, 0x34, 0x56)


def _apply_wolfppt_dropin_shape_gradient_fill_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.dml.color import RGBColor

    fill = _dropin_shape_or_group_child(fixture_id, prs, "shape gradient-fill").fill
    fill.gradient()
    fill.gradient_angle = 45.0
    fill.gradient_stops[0].position = 0.25
    fill.gradient_stops[0].color.rgb = RGBColor(0x12, 0x34, 0x56)


def _apply_python_pptx_dropin_shape_rotation_edit(fixture_id: str, prs: Any) -> None:
    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape rotation")
    shape.rotation = SHAPE_ROTATION_EXPECTED


def _apply_wolfppt_dropin_shape_rotation_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape rotation")
    shape.rotation = SHAPE_ROTATION_EXPECTED


def _apply_python_pptx_dropin_shape_name_edit(fixture_id: str, prs: Any) -> None:
    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape name")
    shape.name = SHAPE_NAME_EXPECTED


def _apply_wolfppt_dropin_shape_name_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    shape = _dropin_shape_or_group_child(fixture_id, prs, "shape name")
    shape.name = SHAPE_NAME_EXPECTED
