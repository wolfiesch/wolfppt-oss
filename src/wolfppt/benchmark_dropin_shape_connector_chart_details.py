"""Chart and connector shape-creation benchmark detail readers."""
# ruff: noqa: F401

from __future__ import annotations

from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from .benchmark_cases import (
    ADD_CHART_EXPECTED,
    ADD_CHART_EXPECTED_TRANSFORM,
    C_NS,
)
from .benchmark_validation import _validate_output
from .chart_adds import _CHART_TEMPLATE_NAME_TO_ID, _XY_CHART_TEMPLATE_NAME_TO_ID
from .benchmark_dropin_shape_connector_details import (
    GROUP_CONNECTOR_TARGETS,
    TOP_LEVEL_CONNECTOR_FIXTURES,
    _connector_connection_point,
    _dropin_add_connector_details,
    _dropin_add_group_connector_details,
    _dropin_add_nested_group_connector_details,
    _dropin_connector_connection_edit_details,
    _dropin_existing_connector_connection_edit_details,
    _dropin_group_connector_connection_edit_details,
)
from .extractor import extract_semantics
from .package_diff import diff_packages
from .presentation import Presentation as WolfPresentation

TOP_LEVEL_CHART_FIXTURES = (
    "text_basic/title_body_bullets",
    "workloads/mixed_real_world_deck",
    "workloads/mixed_deal_review_deck",
)
GROUP_CHART_TARGETS = {
    "shapes/grouped_shapes": (0, 0),
    "workloads/management_reporting_deck": (2, 1),
    "workloads/customer_success_review_pack": (4, 1),
}
STRICT_CHART_ADDED_PARTS = [
    "ppt/charts/_rels/chart1.xml.rels",
    "ppt/charts/chart1.xml",
    "ppt/embeddings/Microsoft_Excel_Sheet1.xlsx",
]
STRICT_CHART_CHANGED_PARTS = [
    "[Content_Types].xml",
    "ppt/slides/_rels/slide1.xml.rels",
    "ppt/slides/slide1.xml",
]


def _top_level_chart_package_delta_ok(
    fixture_id: str,
    added_parts: list[str],
    changed_parts: list[str],
    removed_parts: list[str],
    *,
    slide_index: int = 0,
) -> bool:
    changed_slide_part = f"ppt/slides/slide{slide_index + 1}.xml"
    changed_rels_part = f"ppt/slides/_rels/slide{slide_index + 1}.xml.rels"
    if (
        fixture_id in {"text_basic/title_body_bullets", "shapes/grouped_shapes"}
        and slide_index == 0
    ):
        return (
            added_parts == STRICT_CHART_ADDED_PARTS
            and changed_parts == STRICT_CHART_CHANGED_PARTS
            and removed_parts == []
        )
    chart_parts = [
        part
        for part in added_parts
        if part.startswith("ppt/charts/chart") and part.endswith(".xml")
    ]
    chart_rels = [
        part
        for part in added_parts
        if part.startswith("ppt/charts/_rels/chart") and part.endswith(".xml.rels")
    ]
    workbook_parts = [
        part
        for part in added_parts
        if part.startswith("ppt/embeddings/Microsoft_Excel_Sheet")
        and part.endswith(".xlsx")
    ]
    required_changed = {changed_rels_part, changed_slide_part}
    allowed_changed = required_changed | {"[Content_Types].xml"}
    return (
        len(chart_parts) == 1
        and len(chart_rels) == 1
        and len(workbook_parts) == 1
        and sorted(added_parts) == sorted(chart_parts + chart_rels + workbook_parts)
        and required_changed <= set(changed_parts)
        and set(changed_parts) <= allowed_changed
        and removed_parts == []
    )


def _dropin_add_chart_details(
    fixture_id: str,
    fixture_path: Path,
    output_path: Path,
    validate_openxml: bool,
    *,
    expected_chart_type: str = "COLUMN_CLUSTERED (51)",
    expected_chart_metadata: dict[str, Any] | None = None,
    nested_group: bool = False,
    slide_index: int = 0,
    group_index: int | None = None,
    expected_hierarchy_level_count: int | None = None,
) -> dict[str, Any]:
    if fixture_id not in {
        *TOP_LEVEL_CHART_FIXTURES,
        "shapes/grouped_shapes",
        *GROUP_CHART_TARGETS,
    }:
        raise RuntimeError(
            f"drop-in add-chart benchmark does not support fixture {fixture_id}"
        )
    changed_slide_part = f"ppt/slides/slide{slide_index + 1}.xml"
    changed_rels_part = f"ppt/slides/_rels/slide{slide_index + 1}.xml.rels"
    package_diff = diff_packages(fixture_path, output_path)
    validation = (
        _validate_output(output_path) if validate_openxml else {"enabled": False}
    )
    openxml_ok = validation.get("valid") is not False
    package_delta_ok = _top_level_chart_package_delta_ok(
        fixture_id,
        package_diff.added_parts,
        package_diff.changed_parts,
        package_diff.removed_parts,
        slide_index=slide_index,
    )
    slide_shapes = WolfPresentation(output_path).slides[slide_index].shapes
    parent_preserved = True
    parent_child_names: list[str] = []
    expected_parent_child_names: list[str] = []
    if group_index is not None:
        source_summary = extract_semantics(fixture_path)
        output_summary = extract_semantics(output_path)
        source_group = source_summary.slides[slide_index].shapes[group_index]
        output_group = output_summary.slides[slide_index].shapes[group_index]
        expected_parent_child_names = [child.name for child in source_group.children]
        parent_child_names = [
            child.name
            for child in output_group.children[: len(expected_parent_child_names)]
        ]
        parent_preserved = parent_child_names == expected_parent_child_names
    if group_index is not None:
        group_shapes = slide_shapes[group_index].shapes
        shape = (
            group_shapes[-1].shapes[-1]
            if nested_group
            else group_shapes[-1]
        )
    else:
        shape = slide_shapes[-1]
    chart = shape.chart
    transform = {
        "x": int(shape.left),
        "y": int(shape.top),
        "cx": int(shape.width),
        "cy": int(shape.height),
    }
    chart_metadata = {
        "chart_type": str(chart.chart_type),
        "categories": chart.categories,
        "series": [
            {"name": series.name, "values": list(series.values)}
            for series in chart.series
        ],
    }
    expected_chart_metadata = (
        {**ADD_CHART_EXPECTED, "chart_type": expected_chart_type}
        if expected_chart_metadata is None
        else dict(expected_chart_metadata)
    )
    if "bubble_sizes" in expected_chart_metadata:
        chart_metadata["bubble_sizes"] = [
            list(series.bubble_sizes) for series in chart.series
        ]
    hierarchy_xml = (
        _chart_hierarchy_xml_details(output_path, shape.chart_part.partname.lstrip("/"))
        if expected_hierarchy_level_count is not None
        else None
    )
    hierarchy_xml_pass = (
        True
        if expected_hierarchy_level_count is None
        else hierarchy_xml is not None
        and hierarchy_xml["has_multi_level_ref"] is True
        and hierarchy_xml["has_multi_level_cache"] is True
        and hierarchy_xml["level_count"] == expected_hierarchy_level_count
        and hierarchy_xml["point_count"] == str(
            len(expected_chart_metadata["categories"])
        )
    )
    semantic_pass = (
        shape.has_chart is True
        and transform == ADD_CHART_EXPECTED_TRANSFORM
        and chart_metadata == expected_chart_metadata
        and parent_preserved
        and hierarchy_xml_pass
    )
    ok = semantic_pass and package_delta_ok and openxml_ok
    return {
        "ok": ok,
        "semantic_pass": semantic_pass,
        "semantic_score": 1.0 if semantic_pass else 0.0,
        "chart_pass": semantic_pass,
        "chart_metadata": chart_metadata,
        "expected_chart_metadata": expected_chart_metadata,
        "hierarchy_xml_pass": hierarchy_xml_pass,
        "hierarchy_xml": hierarchy_xml,
        "chart_transform": transform,
        "parent_group_child_preserved": parent_preserved,
        "parent_child_names": parent_child_names,
        "expected_parent_child_names": expected_parent_child_names,
        "package_clean": package_diff.clean,
        "package_added_count": len(package_diff.added_parts),
        "package_changed_count": len(package_diff.changed_parts),
        "package_removed_count": len(package_diff.removed_parts),
        "added_parts": package_diff.added_parts,
        "changed_parts": package_diff.changed_parts,
        "changed_slide_part": changed_slide_part,
        "changed_rels_part": changed_rels_part,
        "openxml_valid": validation.get("valid"),
        "openxml_validation": validation,
    }


def _chart_hierarchy_xml_details(
    output_path: Path,
    changed_chart_part: str,
) -> dict[str, Any]:
    with ZipFile(output_path) as package:
        root = ET.fromstring(package.read(changed_chart_part))
    reference = root.find(f".//{{{C_NS}}}ser/{{{C_NS}}}cat/{{{C_NS}}}multiLvlStrRef")
    cache = None if reference is None else reference.find(f"{{{C_NS}}}multiLvlStrCache")
    point_count = None if cache is None else cache.find(f"{{{C_NS}}}ptCount")
    formula = None if reference is None else reference.find(f"{{{C_NS}}}f")
    return {
        "has_multi_level_ref": reference is not None,
        "has_multi_level_cache": cache is not None,
        "formula": None if formula is None else formula.text,
        "level_count": 0 if cache is None else len(cache.findall(f"{{{C_NS}}}lvl")),
        "point_count": None if point_count is None else point_count.get("val"),
    }


def _dropin_add_chart_template_family_details(
    fixture_id: str,
    fixture_path: Path,
    output_path: Path,
    validate_openxml: bool,
) -> dict[str, Any]:
    if fixture_id not in TOP_LEVEL_CHART_FIXTURES:
        raise RuntimeError(
            f"drop-in add-chart family benchmark does not support fixture {fixture_id}"
        )
    from pptx.enum.chart import XL_CHART_TYPE

    chart_type_names = tuple(sorted(_CHART_TEMPLATE_NAME_TO_ID))
    expected_chart_types = [
        str(getattr(XL_CHART_TYPE, chart_type_name))
        for chart_type_name in chart_type_names
    ]
    package_diff = diff_packages(fixture_path, output_path)
    validation = (
        _validate_output(output_path) if validate_openxml else {"enabled": False}
    )
    openxml_ok = validation.get("valid") is not False
    shapes = WolfPresentation(output_path).slides[0].shapes
    chart_shapes = [shape for shape in shapes if shape.has_chart][
        -len(chart_type_names) :
    ]
    chart_metadata = [
        {
            "chart_type": str(shape.chart.chart_type),
            "categories": shape.chart.categories,
            "series": [
                {"name": series.name, "values": list(series.values)}
                for series in shape.chart.series
            ],
        }
        for shape in chart_shapes
    ]
    expected_metadata = [
        {
            "chart_type": chart_type,
            "categories": ["Q1", "Q2", "Q3"],
            "series": [
                {"name": "Revenue", "values": [10.0, 14.0, 18.0]},
            ],
        }
        for chart_type in expected_chart_types
    ]
    chart_parts = [
        part
        for part in package_diff.added_parts
        if part.startswith("ppt/charts/chart") and part.endswith(".xml")
    ]
    chart_rels = [
        part
        for part in package_diff.added_parts
        if part.startswith("ppt/charts/_rels/chart") and part.endswith(".xml.rels")
    ]
    workbook_parts = [
        part
        for part in package_diff.added_parts
        if part.startswith("ppt/embeddings/Microsoft_Excel_Sheet")
        and part.endswith(".xlsx")
    ]
    package_delta_ok = (
        len(chart_parts) == len(chart_type_names)
        and len(chart_rels) == len(chart_type_names)
        and len(workbook_parts) == len(chart_type_names)
        and sorted(package_diff.added_parts)
        == sorted(chart_parts + chart_rels + workbook_parts)
        and package_diff.changed_parts
        == [
            "[Content_Types].xml",
            "ppt/slides/_rels/slide1.xml.rels",
            "ppt/slides/slide1.xml",
        ]
        and package_diff.removed_parts == []
    )
    semantic_pass = chart_metadata == expected_metadata
    ok = semantic_pass and package_delta_ok and openxml_ok
    return {
        "ok": ok,
        "semantic_pass": semantic_pass,
        "semantic_score": 1.0 if semantic_pass else 0.0,
        "chart_pass": semantic_pass,
        "chart_template_count": len(chart_type_names),
        "chart_types": [item["chart_type"] for item in chart_metadata],
        "expected_chart_types": expected_chart_types,
        "chart_metadata": chart_metadata,
        "expected_chart_metadata": expected_metadata,
        "package_clean": package_diff.clean,
        "package_added_count": len(package_diff.added_parts),
        "package_changed_count": len(package_diff.changed_parts),
        "package_removed_count": len(package_diff.removed_parts),
        "added_parts": package_diff.added_parts,
        "changed_parts": package_diff.changed_parts,
        "openxml_valid": validation.get("valid"),
        "openxml_validation": validation,
    }


def _dropin_add_xy_scatter_template_family_details(
    fixture_id: str,
    fixture_path: Path,
    output_path: Path,
    validate_openxml: bool,
) -> dict[str, Any]:
    if fixture_id not in TOP_LEVEL_CHART_FIXTURES:
        raise RuntimeError(
            f"drop-in add-xy-scatter family benchmark does not support fixture {fixture_id}"
        )
    from pptx.enum.chart import XL_CHART_TYPE

    chart_type_names = tuple(sorted(_XY_CHART_TEMPLATE_NAME_TO_ID))
    expected_chart_types = [
        str(getattr(XL_CHART_TYPE, chart_type_name))
        for chart_type_name in chart_type_names
    ]
    package_diff = diff_packages(fixture_path, output_path)
    validation = (
        _validate_output(output_path) if validate_openxml else {"enabled": False}
    )
    openxml_ok = validation.get("valid") is not False
    shapes = WolfPresentation(output_path).slides[0].shapes
    chart_shapes = [shape for shape in shapes if shape.has_chart][
        -len(chart_type_names) :
    ]
    chart_metadata = [
        {
            "chart_type": str(shape.chart.chart_type),
            "categories": shape.chart.categories,
            "series": [
                {"name": series.name, "values": list(series.values)}
                for series in shape.chart.series
            ],
        }
        for shape in chart_shapes
    ]
    expected_metadata = [
        {
            "chart_type": chart_type,
            "categories": [],
            "series": [
                {"name": "Revenue", "values": [10.0, 14.0]},
                {"name": "Cost", "values": [6.0, 7.0]},
            ],
        }
        for chart_type in expected_chart_types
    ]
    chart_parts = [
        part
        for part in package_diff.added_parts
        if part.startswith("ppt/charts/chart") and part.endswith(".xml")
    ]
    chart_rels = [
        part
        for part in package_diff.added_parts
        if part.startswith("ppt/charts/_rels/chart") and part.endswith(".xml.rels")
    ]
    workbook_parts = [
        part
        for part in package_diff.added_parts
        if part.startswith("ppt/embeddings/Microsoft_Excel_Sheet")
        and part.endswith(".xlsx")
    ]
    package_delta_ok = (
        len(chart_parts) == len(chart_type_names)
        and len(chart_rels) == len(chart_type_names)
        and len(workbook_parts) == len(chart_type_names)
        and sorted(package_diff.added_parts)
        == sorted(chart_parts + chart_rels + workbook_parts)
        and package_diff.changed_parts
        == [
            "[Content_Types].xml",
            "ppt/slides/_rels/slide1.xml.rels",
            "ppt/slides/slide1.xml",
        ]
        and package_diff.removed_parts == []
    )
    semantic_pass = chart_metadata == expected_metadata
    ok = semantic_pass and package_delta_ok and openxml_ok
    return {
        "ok": ok,
        "semantic_pass": semantic_pass,
        "semantic_score": 1.0 if semantic_pass else 0.0,
        "chart_pass": semantic_pass,
        "chart_template_count": len(chart_type_names),
        "chart_types": [item["chart_type"] for item in chart_metadata],
        "expected_chart_types": expected_chart_types,
        "chart_metadata": chart_metadata,
        "expected_chart_metadata": expected_metadata,
        "package_clean": package_diff.clean,
        "package_added_count": len(package_diff.added_parts),
        "package_changed_count": len(package_diff.changed_parts),
        "package_removed_count": len(package_diff.removed_parts),
        "added_parts": package_diff.added_parts,
        "changed_parts": package_diff.changed_parts,
        "openxml_valid": validation.get("valid"),
        "openxml_validation": validation,
    }


def _dropin_add_group_chart_details(
    fixture_id: str,
    fixture_path: Path,
    output_path: Path,
    validate_openxml: bool,
) -> dict[str, Any]:
    if fixture_id not in GROUP_CHART_TARGETS:
        raise RuntimeError(
            f"drop-in add-group-chart benchmark does not support fixture {fixture_id}"
        )
    slide_index, group_index = GROUP_CHART_TARGETS[fixture_id]
    return _dropin_add_chart_details(
        fixture_id,
        fixture_path,
        output_path,
        validate_openxml,
        slide_index=slide_index,
        group_index=group_index,
    )


def _dropin_add_nested_group_chart_details(
    fixture_id: str,
    fixture_path: Path,
    output_path: Path,
    validate_openxml: bool,
) -> dict[str, Any]:
    if fixture_id not in GROUP_CHART_TARGETS:
        raise RuntimeError(
            f"drop-in add-nested-group-chart benchmark does not support fixture {fixture_id}"
        )
    slide_index, group_index = GROUP_CHART_TARGETS[fixture_id]
    return _dropin_add_chart_details(
        fixture_id,
        fixture_path,
        output_path,
        validate_openxml,
        nested_group=True,
        slide_index=slide_index,
        group_index=group_index,
    )
