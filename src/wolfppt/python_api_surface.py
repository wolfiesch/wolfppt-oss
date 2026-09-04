"""Public-surface comparison against python-pptx.

This module keeps the drop-in claim honest by comparing representative
WolfPPT facade objects with the public names exposed by matching python-pptx
objects. It is intentionally about API shape, not behavior; behavior parity
belongs in the focused parity and benchmark lanes.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from .presentation import Presentation as WolfPresentation


REPRESENTATIVE_FIXTURES = {
    "text": Path("text_basic/title_body_bullets.pptx"),
    "chart": Path("charts/bar_chart.pptx"),
    "media": Path("media/png_picture.pptx"),
    "table": Path("tables/simple_table.pptx"),
    "group": Path("shapes/grouped_shapes.pptx"),
    "notes": Path("notes/speaker_notes.pptx"),
}


@dataclass(frozen=True)
class SurfacePair:
    label: str
    python_object: Any
    wolfppt_object: Any


@dataclass(frozen=True)
class MissingAttribute:
    name: str
    error: str


@dataclass(frozen=True)
class AttributeReadError:
    name: str
    error: str


READ_PROBE_SKIP_NAMES = frozenset(
    {
        # python-pptx may create note parts when these are read on decks that do
        # not already contain them; that side effect is covered by behavior tests.
        "notes_master",
        "notes_slide",
    }
)


def compare_public_surface(
    fixture_root: Path = Path("fixtures/pptx"),
    *,
    max_real_slides: int = 3,
    max_real_shapes: int = 5,
    min_fixtures: int | None = None,
    min_representative_pairs: int | None = None,
    min_real_pairs: int | None = None,
    run_id: str | None = None,
) -> dict[str, Any]:
    """Compare public object names on representative and real fixture samples."""

    _validate_non_negative("max_real_slides", max_real_slides)
    _validate_non_negative("max_real_shapes", max_real_shapes)
    _validate_positive_optional("min_fixtures", min_fixtures)
    _validate_positive_optional("min_representative_pairs", min_representative_pairs)
    _validate_positive_optional("min_real_pairs", min_real_pairs)

    python_pptx = _load_python_pptx()
    representative_pairs = _representative_surface_pairs(fixture_root, python_pptx)
    real_pairs = _real_fixture_surface_pairs(
        fixture_root,
        python_pptx,
        max_slides=max_real_slides,
        max_shapes=max_real_shapes,
    )
    representative_missing = _missing_by_label(representative_pairs)
    real_missing = _missing_by_label(real_pairs)
    representative_read_errors = _attribute_read_errors_by_label(representative_pairs)
    real_read_errors = _attribute_read_errors_by_label(real_pairs)
    all_missing = {
        **representative_missing,
        **{f"real:{label}": names for label, names in real_missing.items()},
    }
    all_read_errors = {
        **representative_read_errors,
        **{
            f"real:{label}": errors
            for label, errors in real_read_errors.items()
        },
    }
    fixture_count = len(_deck_paths(fixture_root))
    gates = {
        "fixtures": _gate_result(fixture_count, min_fixtures),
        "representative_pairs": _gate_result(
            len(representative_pairs),
            min_representative_pairs,
        ),
        "real_fixture_pairs": _gate_result(len(real_pairs), min_real_pairs),
    }
    gate_failures = [
        name
        for name, gate in gates.items()
        if gate["required"] is not None and not gate["passed"]
    ]
    read_error_count = sum(len(errors) for errors in all_read_errors.values())
    return {
        "status": "pass"
        if not all_missing and not all_read_errors and not gate_failures
        else "fail",
        "run_id": run_id or datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ"),
        "python_pptx_version": getattr(python_pptx, "__version__", "unknown"),
        "fixture_root": str(fixture_root),
        "fixture_count": fixture_count,
        "representative_pair_count": len(representative_pairs),
        "real_fixture_pair_count": len(real_pairs),
        "attribute_read_error_count": read_error_count,
        "max_real_slides_per_fixture": max_real_slides,
        "max_real_shapes_per_slide": max_real_shapes,
        "gates": gates,
        "gate_failures": gate_failures,
        "missing_by_label": {key: sorted(value) for key, value in all_missing.items()},
        "attribute_read_errors_by_label": {
            key: [
                {
                    "name": error.name,
                    "error": error.error,
                }
                for error in errors
            ]
            for key, errors in sorted(all_read_errors.items())
        },
    }


def render_public_surface_markdown(report: dict[str, Any]) -> str:
    """Render a compact Markdown report for humans."""

    lines = [
        "# Python API Surface Ratchet",
        "",
        "| Check | Result |",
        "|---|---:|",
        f"| status | {report['status']} |",
        f"| python-pptx version | {report['python_pptx_version']} |",
        f"| checked fixtures | {report['fixture_count']} |",
        f"| representative object pairs | {report['representative_pair_count']} |",
        f"| real-fixture object pairs | {report['real_fixture_pair_count']} |",
        f"| attribute read errors | {report['attribute_read_error_count']} |",
        f"| max real slides per fixture | {report['max_real_slides_per_fixture']} |",
        f"| max real shapes per slide | {report['max_real_shapes_per_slide']} |",
        "",
    ]
    gated_rows = [
        (name, gate)
        for name, gate in report["gates"].items()
        if gate["required"] is not None
    ]
    if gated_rows:
        lines.extend(["## Gates", "", "| Gate | Observed | Required | Status |", "|---|---:|---:|---|"])
        for name, gate in gated_rows:
            status = "pass" if gate["passed"] else "fail"
            lines.append(
                f"| {name.replace('_', ' ')} | {gate['observed']} | "
                f"{gate['required']} | {status} |"
            )
        lines.append("")
    missing = report["missing_by_label"]
    read_errors = report.get("attribute_read_errors_by_label", {})
    if not missing and not read_errors:
        lines.append(
            "No missing public python-pptx names or read errors were found on "
            "the checked object surface."
        )
    else:
        if missing:
            lines.extend(
                ["## Missing Public Names", "", "| Object | Missing names |", "|---|---|"]
            )
            for label, names in sorted(missing.items()):
                lines.append(f"| `{label}` | `{', '.join(names)}` |")
        if read_errors:
            lines.extend(
                [
                    "",
                    "## Attribute Read Errors",
                    "",
                    "| Object | Attribute | Error |",
                    "|---|---|---|",
                ]
            )
            for label, errors in sorted(read_errors.items()):
                for error in errors:
                    lines.append(
                        f"| `{label}` | `{error['name']}` | `{error['error']}` |"
                    )
    return "\n".join(lines) + "\n"


def write_public_surface_report(
    report: dict[str, Any],
    output_dir: Path = Path("results/api-surface/latest"),
) -> None:
    """Write latest and archived public-surface evidence reports."""

    for target_dir in _surface_report_dirs(report["run_id"], output_dir):
        target_dir.mkdir(parents=True, exist_ok=True)
        (target_dir / "run.json").write_text(
            json.dumps(report, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        (target_dir / "report.md").write_text(
            render_public_surface_markdown(report),
            encoding="utf-8",
        )


def _load_python_pptx() -> Any:
    try:
        import pptx
    except ImportError as exc:  # pragma: no cover - exercised without dev extras.
        raise RuntimeError(
            "python-pptx is required; install dev or baseline extras first"
        ) from exc
    return pptx


def _representative_surface_pairs(
    fixture_root: Path, python_pptx: Any
) -> list[SurfacePair]:
    sources = {
        key: fixture_root / relative
        for key, relative in REPRESENTATIVE_FIXTURES.items()
    }
    python_text = python_pptx.Presentation(str(sources["text"]))
    wolf_text = WolfPresentation(sources["text"])
    python_chart = python_pptx.Presentation(str(sources["chart"]))
    wolf_chart = WolfPresentation(sources["chart"])
    python_media = python_pptx.Presentation(str(sources["media"]))
    wolf_media = WolfPresentation(sources["media"])
    python_table = python_pptx.Presentation(str(sources["table"]))
    wolf_table = WolfPresentation(sources["table"])
    python_group = python_pptx.Presentation(str(sources["group"]))
    wolf_group = WolfPresentation(sources["group"])
    python_notes = python_pptx.Presentation(str(sources["notes"]))
    wolf_notes = WolfPresentation(sources["notes"])

    python_text_shape = python_text.slides[0].shapes[0]
    wolf_text_shape = wolf_text.slides[0].shapes[0]
    python_chart_shape = python_chart.slides[0].shapes[0]
    wolf_chart_shape = wolf_chart.slides[0].shapes[0]
    python_chart_obj = python_chart_shape.chart
    wolf_chart_obj = wolf_chart_shape.chart
    python_chart_obj.category_axis.axis_title.text_frame.text = "Category Surface"
    wolf_chart_obj.category_axis.axis_title.text_frame.text = "Category Surface"
    python_chart_obj.value_axis.axis_title.text_frame.text = "Value Surface"
    wolf_chart_obj.value_axis.axis_title.text_frame.text = "Value Surface"
    python_chart_title_paragraph = python_chart_obj.chart_title.text_frame.paragraphs[0]
    wolf_chart_title_paragraph = wolf_chart_obj.chart_title.text_frame.paragraphs[0]
    python_chart_title_run = python_chart_title_paragraph.add_run()
    python_chart_title_run.text = "surface"
    wolf_chart_title_run = wolf_chart_title_paragraph.add_run()
    wolf_chart_title_run.text = "surface"
    python_chart_plot = python_chart_obj.plots[0]
    wolf_chart_plot = wolf_chart_obj.plots[0]
    python_chart_plot.has_data_labels = True
    wolf_chart_plot.has_data_labels = True
    python_chart_series = python_chart_obj.series[0]
    wolf_chart_series = wolf_chart_obj.series[0]
    python_chart_point = python_chart_series.points[0]
    wolf_chart_point = wolf_chart_series.points[0]
    python_chart_point_label = python_chart_point.data_label
    wolf_chart_point_label = wolf_chart_point.data_label
    python_chart_point_label_paragraph = (
        python_chart_point_label.text_frame.paragraphs[0]
    )
    wolf_chart_point_label_paragraph = wolf_chart_point_label.text_frame.paragraphs[0]
    python_chart_point_label_run = python_chart_point_label_paragraph.add_run()
    python_chart_point_label_run.text = "surface"
    wolf_chart_point_label_run = wolf_chart_point_label_paragraph.add_run()
    wolf_chart_point_label_run.text = "surface"
    python_picture_shape = python_media.slides[0].shapes[0]
    wolf_picture_shape = wolf_media.slides[0].shapes[0]
    python_table_shape = python_table.slides[0].shapes[0]
    wolf_table_shape = wolf_table.slides[0].shapes[0]
    python_cell = python_table_shape.table.cell(0, 0)
    wolf_cell = wolf_table_shape.table.cell(0, 0)
    python_group_shape = python_group.slides[0].shapes[0]
    wolf_group_shape = wolf_group.slides[0].shapes[0]
    python_notes_slide = python_notes.slides[0].notes_slide
    wolf_notes_slide = wolf_notes.slides[0].notes_slide

    return [
        SurfacePair("presentation", python_text, wolf_text),
        SurfacePair("slides", python_text.slides, wolf_text.slides),
        SurfacePair("slide", python_text.slides[0], wolf_text.slides[0]),
        SurfacePair(
            "slide_layouts", python_text.slide_layouts, wolf_text.slide_layouts
        ),
        SurfacePair(
            "slide_layout", python_text.slide_layouts[0], wolf_text.slide_layouts[0]
        ),
        SurfacePair(
            "layout_shapes",
            python_text.slide_layouts[0].shapes,
            wolf_text.slide_layouts[0].shapes,
        ),
        SurfacePair(
            "layout_placeholders",
            python_text.slide_layouts[0].placeholders,
            wolf_text.slide_layouts[0].placeholders,
        ),
        SurfacePair(
            "master_shapes",
            python_text.slide_master.shapes,
            wolf_text.slide_master.shapes,
        ),
        SurfacePair(
            "master_placeholders",
            python_text.slide_master.placeholders,
            wolf_text.slide_master.placeholders,
        ),
        SurfacePair(
            "placeholder_format",
            python_text.slide_layouts[0].placeholders[0].placeholder_format,
            wolf_text.slide_layouts[0].placeholders[0].placeholder_format,
        ),
        SurfacePair(
            "slide_shapes", python_text.slides[0].shapes, wolf_text.slides[0].shapes
        ),
        SurfacePair("shape", python_text_shape, wolf_text_shape),
        SurfacePair("chart_shape", python_chart_shape, wolf_chart_shape),
        SurfacePair(
            "text_frame", python_text_shape.text_frame, wolf_text_shape.text_frame
        ),
        SurfacePair(
            "paragraph",
            python_text_shape.text_frame.paragraphs[0],
            wolf_text_shape.text_frame.paragraphs[0],
        ),
        SurfacePair(
            "run",
            python_text_shape.text_frame.paragraphs[0].runs[0],
            wolf_text_shape.text_frame.paragraphs[0].runs[0],
        ),
        SurfacePair(
            "font",
            python_text_shape.text_frame.paragraphs[0].runs[0].font,
            wolf_text_shape.text_frame.paragraphs[0].runs[0].font,
        ),
        SurfacePair("fill", python_text_shape.fill, wolf_text_shape.fill),
        SurfacePair("line", python_text_shape.line, wolf_text_shape.line),
        SurfacePair(
            "click_action", python_text_shape.click_action, wolf_text_shape.click_action
        ),
        SurfacePair("chart", python_chart_obj, wolf_chart_obj),
        SurfacePair(
            "chart_title",
            python_chart_obj.chart_title,
            wolf_chart_obj.chart_title,
        ),
        SurfacePair(
            "chart_title_text_frame",
            python_chart_obj.chart_title.text_frame,
            wolf_chart_obj.chart_title.text_frame,
        ),
        SurfacePair(
            "chart_title_paragraph_font",
            python_chart_title_paragraph.font,
            wolf_chart_title_paragraph.font,
        ),
        SurfacePair(
            "chart_title_run_font",
            python_chart_title_run.font,
            wolf_chart_title_run.font,
        ),
        SurfacePair("chart_plot", python_chart_plot, wolf_chart_plot),
        SurfacePair(
            "chart_plot_series",
            python_chart_plot.series,
            wolf_chart_plot.series,
        ),
        SurfacePair("chart_series", python_chart_series, wolf_chart_series),
        SurfacePair(
            "chart_points",
            python_chart_series.points,
            wolf_chart_series.points,
        ),
        SurfacePair("chart_point", python_chart_point, wolf_chart_point),
        SurfacePair(
            "chart_point_marker",
            python_chart_point.marker,
            wolf_chart_point.marker,
        ),
        SurfacePair(
            "chart_data_labels",
            python_chart_plot.data_labels,
            wolf_chart_plot.data_labels,
        ),
        SurfacePair(
            "chart_data_labels_font",
            python_chart_plot.data_labels.font,
            wolf_chart_plot.data_labels.font,
        ),
        SurfacePair(
            "chart_series_data_labels",
            python_chart_series.data_labels,
            wolf_chart_series.data_labels,
        ),
        SurfacePair(
            "chart_series_data_labels_font",
            python_chart_series.data_labels.font,
            wolf_chart_series.data_labels.font,
        ),
        SurfacePair(
            "chart_point_data_label",
            python_chart_point_label,
            wolf_chart_point_label,
        ),
        SurfacePair(
            "chart_point_data_label_text_frame",
            python_chart_point_label.text_frame,
            wolf_chart_point_label.text_frame,
        ),
        SurfacePair(
            "chart_point_data_label_paragraphs",
            python_chart_point_label.text_frame.paragraphs,
            wolf_chart_point_label.text_frame.paragraphs,
        ),
        SurfacePair(
            "chart_point_data_label_paragraph",
            python_chart_point_label_paragraph,
            wolf_chart_point_label_paragraph,
        ),
        SurfacePair(
            "chart_point_data_label_paragraph_font",
            python_chart_point_label_paragraph.font,
            wolf_chart_point_label_paragraph.font,
        ),
        SurfacePair(
            "chart_point_data_label_runs",
            python_chart_point_label_paragraph.runs,
            wolf_chart_point_label_paragraph.runs,
        ),
        SurfacePair(
            "chart_point_data_label_run",
            python_chart_point_label_run,
            wolf_chart_point_label_run,
        ),
        SurfacePair(
            "chart_point_data_label_run_font",
            python_chart_point_label_run.font,
            wolf_chart_point_label_run.font,
        ),
        SurfacePair(
            "chart_point_data_label_run_hyperlink",
            python_chart_point_label_run.hyperlink,
            wolf_chart_point_label_run.hyperlink,
        ),
        SurfacePair(
            "category_axis",
            python_chart_obj.category_axis,
            wolf_chart_obj.category_axis,
        ),
        SurfacePair(
            "category_axis_title",
            python_chart_obj.category_axis.axis_title,
            wolf_chart_obj.category_axis.axis_title,
        ),
        SurfacePair(
            "category_axis_title_text_frame",
            python_chart_obj.category_axis.axis_title.text_frame,
            wolf_chart_obj.category_axis.axis_title.text_frame,
        ),
        SurfacePair(
            "value_axis",
            python_chart_obj.value_axis,
            wolf_chart_obj.value_axis,
        ),
        SurfacePair(
            "value_axis_title",
            python_chart_obj.value_axis.axis_title,
            wolf_chart_obj.value_axis.axis_title,
        ),
        SurfacePair(
            "value_axis_title_text_frame",
            python_chart_obj.value_axis.axis_title.text_frame,
            wolf_chart_obj.value_axis.axis_title.text_frame,
        ),
        SurfacePair("picture_shape", python_picture_shape, wolf_picture_shape),
        SurfacePair(
            "picture_image",
            python_picture_shape.image,
            wolf_picture_shape.image,
        ),
        SurfacePair("table", python_table_shape.table, wolf_table_shape.table),
        SurfacePair(
            "table_rows", python_table_shape.table.rows, wolf_table_shape.table.rows
        ),
        SurfacePair(
            "table_columns",
            python_table_shape.table.columns,
            wolf_table_shape.table.columns,
        ),
        SurfacePair(
            "table_row",
            python_table_shape.table.rows[0],
            wolf_table_shape.table.rows[0],
        ),
        SurfacePair(
            "table_column",
            python_table_shape.table.columns[0],
            wolf_table_shape.table.columns[0],
        ),
        SurfacePair("table_cell", python_cell, wolf_cell),
        SurfacePair("table_cell_fill", python_cell.fill, wolf_cell.fill),
        SurfacePair(
            "table_cell_text_frame", python_cell.text_frame, wolf_cell.text_frame
        ),
        SurfacePair("group_shape", python_group_shape, wolf_group_shape),
        SurfacePair("group_shapes", python_group_shape.shapes, wolf_group_shape.shapes),
        SurfacePair("notes_slide", python_notes_slide, wolf_notes_slide),
        SurfacePair("notes_shapes", python_notes_slide.shapes, wolf_notes_slide.shapes),
    ]


def _real_fixture_surface_pairs(
    fixture_root: Path,
    python_pptx: Any,
    *,
    max_slides: int,
    max_shapes: int,
) -> list[SurfacePair]:
    pairs: list[SurfacePair] = []
    for source in _deck_paths(fixture_root):
        python_prs = python_pptx.Presentation(str(source))
        wolf_prs = WolfPresentation(source)
        for slide_index in range(
            min(len(python_prs.slides), len(wolf_prs.slides), max_slides)
        ):
            python_shapes = python_prs.slides[slide_index].shapes
            wolf_shapes = wolf_prs.slides[slide_index].shapes
            for shape_index in range(
                min(len(python_shapes), len(wolf_shapes), max_shapes)
            ):
                python_shape = python_shapes[shape_index]
                wolf_shape = wolf_shapes[shape_index]
                label = f"{source}:{slide_index}:{shape_index}"
                pairs.append(SurfacePair(f"{label}:shape", python_shape, wolf_shape))
                _append_optional_surface_pair(
                    pairs, label, "fill", python_shape, wolf_shape
                )
                _append_optional_surface_pair(
                    pairs, label, "line", python_shape, wolf_shape
                )
                _append_optional_surface_pair(
                    pairs, label, "shadow", python_shape, wolf_shape
                )
                _append_text_surface_pairs(pairs, label, python_shape, wolf_shape)
                _append_table_surface_pairs(pairs, label, python_shape, wolf_shape)
                if python_shape.has_chart and wolf_shape.has_chart:
                    pairs.append(
                        SurfacePair(
                            f"{label}:chart", python_shape.chart, wolf_shape.chart
                        )
                    )
    return pairs


def _append_text_surface_pairs(
    pairs: list[SurfacePair],
    label: str,
    python_shape: Any,
    wolf_shape: Any,
) -> None:
    if not (python_shape.has_text_frame and wolf_shape.has_text_frame):
        return
    python_paragraph = python_shape.text_frame.paragraphs[0]
    wolf_paragraph = wolf_shape.text_frame.paragraphs[0]
    pairs.extend(
        [
            SurfacePair(
                f"{label}:text_frame", python_shape.text_frame, wolf_shape.text_frame
            ),
            SurfacePair(f"{label}:paragraph", python_paragraph, wolf_paragraph),
        ]
    )
    if python_paragraph.runs and wolf_paragraph.runs:
        pairs.extend(
            [
                SurfacePair(
                    f"{label}:run", python_paragraph.runs[0], wolf_paragraph.runs[0]
                ),
                SurfacePair(
                    f"{label}:font",
                    python_paragraph.runs[0].font,
                    wolf_paragraph.runs[0].font,
                ),
            ]
        )


def _append_table_surface_pairs(
    pairs: list[SurfacePair],
    label: str,
    python_shape: Any,
    wolf_shape: Any,
) -> None:
    if not (python_shape.has_table and wolf_shape.has_table):
        return
    python_cell = python_shape.table.cell(0, 0)
    wolf_cell = wolf_shape.table.cell(0, 0)
    pairs.extend(
        [
            SurfacePair(f"{label}:table", python_shape.table, wolf_shape.table),
            SurfacePair(f"{label}:table_cell", python_cell, wolf_cell),
            SurfacePair(f"{label}:table_cell_fill", python_cell.fill, wolf_cell.fill),
            SurfacePair(
                f"{label}:table_cell_text_frame",
                python_cell.text_frame,
                wolf_cell.text_frame,
            ),
        ]
    )


def _append_optional_surface_pair(
    pairs: list[SurfacePair],
    label: str,
    name: str,
    python_object: Any,
    wolfppt_object: Any,
) -> None:
    try:
        python_value = getattr(python_object, name)
    except (AttributeError, NotImplementedError):
        return
    try:
        wolfppt_value = getattr(wolfppt_object, name)
    except Exception as exc:  # noqa: BLE001 - report the missing proxy instead of aborting.
        wolfppt_value = MissingAttribute(name=name, error=str(exc))
    pairs.append(SurfacePair(f"{label}:{name}", python_value, wolfppt_value))


def _missing_by_label(pairs: list[SurfacePair]) -> dict[str, set[str]]:
    missing: dict[str, set[str]] = {}
    for pair in pairs:
        names = _public_names(pair.python_object) - _public_names(pair.wolfppt_object)
        if isinstance(pair.wolfppt_object, MissingAttribute):
            names.add(pair.wolfppt_object.name)
        if names:
            missing[pair.label] = names
    return missing


def _attribute_read_errors_by_label(
    pairs: list[SurfacePair],
) -> dict[str, list[AttributeReadError]]:
    failures: dict[str, list[AttributeReadError]] = {}
    for pair in pairs:
        common_names = _public_names(pair.python_object) & _public_names(
            pair.wolfppt_object
        )
        for name in sorted(common_names - READ_PROBE_SKIP_NAMES):
            try:
                getattr(pair.python_object, name)
            except Exception:
                continue
            try:
                getattr(pair.wolfppt_object, name)
            except Exception as exc:  # noqa: BLE001 - report facade read mismatches.
                failures.setdefault(pair.label, []).append(
                    AttributeReadError(
                        name=name,
                        error=f"{type(exc).__name__}: {exc}",
                    )
                )
    return failures


def _public_names(value: Any) -> set[str]:
    if isinstance(value, MissingAttribute):
        return set()
    return {name for name in dir(value) if not name.startswith("_")}


def _gate_result(observed: int, required: int | None) -> dict[str, int | bool | None]:
    return {
        "observed": observed,
        "required": required,
        "passed": required is None or observed >= required,
    }


def _validate_non_negative(name: str, value: int) -> None:
    if value < 0:
        raise ValueError(f"{name} must be non-negative")


def _validate_positive_optional(name: str, value: int | None) -> None:
    if value is not None and value < 1:
        raise ValueError(f"{name} must be at least 1")


def _surface_report_dirs(run_id: str, output_dir: Path) -> tuple[Path, ...]:
    if output_dir.name != "latest":
        return (output_dir,)
    return (output_dir, output_dir.parent / "runs" / run_id)


def _deck_paths(fixture_root: Path) -> list[Path]:
    return sorted(
        path
        for path in fixture_root.rglob("*")
        if path.suffix.lower() in {".pptx", ".pptm"}
    )
