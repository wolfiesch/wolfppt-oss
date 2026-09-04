"""Chart drop-in benchmark edit actions and readers."""
# ruff: noqa: F401

from __future__ import annotations

from typing import Any

from .benchmark_cases import (
    TEXT_FRAME_AUTO_SIZE_EXPECTED_NAME,
    TEXT_FRAME_MARGIN_EXPECTED,
    TEXT_FRAME_VERTICAL_ANCHOR_EXPECTED_NAME,
    TEXT_FRAME_WORD_WRAP_EXPECTED,
)
from .presentation import Presentation as WolfPresentation

from .benchmark_dropin_chart_readers import (
    CHART_EDIT_FIXTURES,
    _first_chart_shape,
)
from .benchmark_dropin_chart_title_actions import (
    _apply_python_pptx_dropin_chart_axis_title_edit,
    _apply_python_pptx_dropin_chart_axis_title_remove,
    _apply_python_pptx_dropin_chart_legend_edit,
    _apply_python_pptx_dropin_chart_title_edit,
    _apply_python_pptx_dropin_chart_title_paragraph_font_edit,
    _apply_python_pptx_dropin_chart_title_paragraph_format_edit,
    _apply_python_pptx_dropin_chart_title_remove,
    _apply_python_pptx_dropin_chart_title_text_frame_edit,
    _apply_python_pptx_dropin_chart_title_text_frame_flow_edit,
    _apply_wolfppt_dropin_chart_axis_title_edit,
    _apply_wolfppt_dropin_chart_axis_title_remove,
    _apply_wolfppt_dropin_chart_legend_edit,
    _apply_wolfppt_dropin_chart_title_edit,
    _apply_wolfppt_dropin_chart_title_paragraph_font_edit,
    _apply_wolfppt_dropin_chart_title_paragraph_format_edit,
    _apply_wolfppt_dropin_chart_title_remove,
    _apply_wolfppt_dropin_chart_title_text_frame_edit,
    _apply_wolfppt_dropin_chart_title_text_frame_flow_edit,
)


def _apply_python_pptx_dropin_chart_axis_property_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_AXIS_CROSSES, XL_TICK_LABEL_POSITION, XL_TICK_MARK
    from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_PATTERN_TYPE
    from pptx.enum.lang import MSO_LANGUAGE_ID
    from pptx.util import Pt

    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart axis property benchmark does not support fixture {fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.category_axis.visible = False
    chart.category_axis.major_tick_mark = XL_TICK_MARK.OUTSIDE
    chart.category_axis.minor_tick_mark = XL_TICK_MARK.INSIDE
    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    chart.category_axis.tick_labels.number_format = "mmm-yy"
    chart.category_axis.tick_labels.number_format_is_linked = True
    chart.category_axis.tick_labels.font.bold = True
    chart.category_axis.tick_labels.font.italic = True
    chart.category_axis.tick_labels.font.underline = True
    chart.category_axis.tick_labels.font.size = Pt(15)
    chart.category_axis.tick_labels.font.fill.gradient()
    chart.category_axis.tick_labels.font.fill.gradient_angle = 45.0
    chart.category_axis.tick_labels.font.fill.gradient_stops[0].position = 0.25
    chart.category_axis.tick_labels.font.fill.gradient_stops[0].color.rgb = RGBColor(
        0x12,
        0x34,
        0x56,
    )
    chart.category_axis.tick_labels.font.name = "Aptos"
    chart.category_axis.tick_labels.font.language_id = MSO_LANGUAGE_ID.GERMAN
    chart.category_axis.tick_labels.offset = 250
    chart.category_axis.reverse_order = True
    chart.category_axis.format.line.fill.solid()
    chart.category_axis.format.line.color.rgb = RGBColor(0x11, 0x22, 0x33)
    chart.category_axis.format.line.width = Pt(1.25)
    chart.category_axis.format.line.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT
    chart.value_axis.visible = False
    chart.value_axis.tick_labels.number_format = "$#,##0"
    chart.value_axis.tick_labels.number_format_is_linked = False
    chart.value_axis.tick_labels.font.bold = True
    chart.value_axis.tick_labels.font.italic = True
    chart.value_axis.tick_labels.font.underline = True
    chart.value_axis.tick_labels.font.size = Pt(15)
    chart.value_axis.tick_labels.font.fill.patterned()
    chart.value_axis.tick_labels.font.fill.pattern = MSO_PATTERN_TYPE.DIVOT
    chart.value_axis.tick_labels.font.fill.fore_color.rgb = RGBColor(0x65, 0x43, 0x21)
    chart.value_axis.tick_labels.font.fill.back_color.rgb = RGBColor(0xFE, 0xDC, 0xBA)
    chart.value_axis.tick_labels.font.name = "Arial"
    chart.value_axis.tick_labels.font.language_id = MSO_LANGUAGE_ID.FRENCH
    chart.value_axis.minimum_scale = 1.0
    chart.value_axis.maximum_scale = 20.0
    chart.value_axis.major_unit = 5.0
    chart.value_axis.minor_unit = 1.0
    chart.value_axis.format.line.fill.solid()
    chart.value_axis.format.line.color.rgb = RGBColor(0xAA, 0xBB, 0xCC)
    chart.value_axis.format.line.width = Pt(2.25)
    chart.value_axis.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.fill.solid()
    chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor(
        0x12,
        0x34,
        0x56,
    )
    chart.value_axis.major_gridlines.format.line.width = Pt(1.5)
    chart.value_axis.major_gridlines.format.line.dash_style = (
        MSO_LINE_DASH_STYLE.DASH_DOT
    )
    chart.value_axis.has_minor_gridlines = True
    chart.value_axis.crosses = XL_AXIS_CROSSES.MAXIMUM


def _apply_python_pptx_dropin_chart_plot_property_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart plot property benchmark does not support fixture {fixture_id}"
        )
    plot = _first_chart_shape(prs).chart.plots[0]
    plot.vary_by_categories = False
    plot.gap_width = 219
    plot.overlap = -25


def _apply_wolfppt_dropin_chart_plot_property_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart plot property benchmark does not support fixture {fixture_id}"
        )
    plot = _first_chart_shape(prs).chart.plots[0]
    plot.vary_by_categories = False
    plot.gap_width = 219
    plot.overlap = -25


def _apply_python_pptx_dropin_chart_data_label_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_LABEL_POSITION
    from pptx.enum.dml import MSO_PATTERN_TYPE
    from pptx.enum.lang import MSO_LANGUAGE_ID
    from pptx.util import Pt

    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart data-label benchmark does not support fixture {fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.show_value = True
    plot.data_labels.show_category_name = True
    plot.data_labels.show_series_name = False
    plot.data_labels.show_percentage = False
    plot.data_labels.show_legend_key = True
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.number_format = "$#,##0"
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.bold = True
    plot.data_labels.font.italic = True
    plot.data_labels.font.underline = True
    plot.data_labels.font.size = Pt(14)
    plot.data_labels.font.fill.gradient()
    plot.data_labels.font.fill.gradient_angle = 45.0
    plot.data_labels.font.fill.gradient_stops[0].position = 0.25
    plot.data_labels.font.fill.gradient_stops[0].color.rgb = RGBColor(
        0x12,
        0x34,
        0x56,
    )
    plot.data_labels.font.name = "Aptos"
    plot.data_labels.font.language_id = MSO_LANGUAGE_ID.FRENCH
    labels = chart.series[0].data_labels
    labels.show_value = True
    labels.show_category_name = True
    labels.show_series_name = False
    labels.show_percentage = False
    labels.show_legend_key = True
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.number_format = "$#,##0"
    labels.number_format_is_linked = False
    labels.font.bold = True
    labels.font.italic = True
    labels.font.underline = True
    labels.font.size = Pt(14)
    labels.font.fill.patterned()
    labels.font.fill.pattern = MSO_PATTERN_TYPE.DIVOT
    labels.font.fill.fore_color.rgb = RGBColor(0x0C, 0x22, 0x38)
    labels.font.fill.back_color.rgb = RGBColor(0xAB, 0xCD, 0xEF)
    labels.font.name = "Aptos"
    labels.font.language_id = MSO_LANGUAGE_ID.GERMAN


def _apply_wolfppt_dropin_chart_data_label_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_LABEL_POSITION
    from pptx.enum.dml import MSO_PATTERN_TYPE
    from pptx.enum.lang import MSO_LANGUAGE_ID
    from pptx.util import Pt

    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart data-label benchmark does not support fixture {fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.show_value = True
    plot.data_labels.show_category_name = True
    plot.data_labels.show_series_name = False
    plot.data_labels.show_percentage = False
    plot.data_labels.show_legend_key = True
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.number_format = "$#,##0"
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.bold = True
    plot.data_labels.font.italic = True
    plot.data_labels.font.underline = True
    plot.data_labels.font.size = Pt(14)
    plot.data_labels.font.fill.gradient()
    plot.data_labels.font.fill.gradient_angle = 45.0
    plot.data_labels.font.fill.gradient_stops[0].position = 0.25
    plot.data_labels.font.fill.gradient_stops[0].color.rgb = RGBColor(
        0x12,
        0x34,
        0x56,
    )
    plot.data_labels.font.name = "Aptos"
    plot.data_labels.font.language_id = MSO_LANGUAGE_ID.FRENCH
    labels = chart.series[0].data_labels
    labels.show_value = True
    labels.show_category_name = True
    labels.show_series_name = False
    labels.show_percentage = False
    labels.show_legend_key = True
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.number_format = "$#,##0"
    labels.number_format_is_linked = False
    labels.font.bold = True
    labels.font.italic = True
    labels.font.underline = True
    labels.font.size = Pt(14)
    labels.font.fill.patterned()
    labels.font.fill.pattern = MSO_PATTERN_TYPE.DIVOT
    labels.font.fill.fore_color.rgb = RGBColor(0x0C, 0x22, 0x38)
    labels.font.fill.back_color.rgb = RGBColor(0xAB, 0xCD, 0xEF)
    labels.font.name = "Aptos"
    labels.font.language_id = MSO_LANGUAGE_ID.GERMAN


def _apply_python_pptx_dropin_chart_data_label_remove(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart data-label removal benchmark does not support fixture {fixture_id}"
        )
    plot = _first_chart_shape(prs).chart.plots[0]
    plot.has_data_labels = True
    plot.has_data_labels = False


def _apply_wolfppt_dropin_chart_data_label_remove(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart data-label removal benchmark does not support fixture {fixture_id}"
        )
    plot = _first_chart_shape(prs).chart.plots[0]
    plot.has_data_labels = True
    plot.has_data_labels = False


def _apply_python_pptx_dropin_chart_point_data_label_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_LABEL_POSITION
    from pptx.enum.dml import MSO_PATTERN_TYPE
    from pptx.enum.lang import MSO_LANGUAGE_ID
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Pt

    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart point data-label benchmark does not support fixture {fixture_id}"
        )
    label = _first_chart_shape(prs).chart.series[0].points[0].data_label
    label.position = XL_LABEL_POSITION.OUTSIDE_END
    label.text_frame.margin_left = TEXT_FRAME_MARGIN_EXPECTED["margin_left"]
    label.text_frame.margin_right = TEXT_FRAME_MARGIN_EXPECTED["margin_right"]
    label.text_frame.margin_top = TEXT_FRAME_MARGIN_EXPECTED["margin_top"]
    label.text_frame.margin_bottom = TEXT_FRAME_MARGIN_EXPECTED["margin_bottom"]
    label.text_frame.word_wrap = TEXT_FRAME_WORD_WRAP_EXPECTED
    label.text_frame.vertical_anchor = MSO_ANCHOR[
        TEXT_FRAME_VERTICAL_ANCHOR_EXPECTED_NAME
    ]
    label.text_frame.auto_size = MSO_AUTO_SIZE[TEXT_FRAME_AUTO_SIZE_EXPECTED_NAME]
    run = label.text_frame.paragraphs[0].add_run()
    run.text = "Q1 "
    run.font.bold = True
    run.font.italic = True
    run.font.underline = True
    run.font.size = Pt(11)
    run.font.fill.gradient()
    run.font.fill.gradient_angle = 45.0
    run.font.fill.gradient_stops[0].position = 0.25
    run.font.fill.gradient_stops[0].color.rgb = RGBColor(0x12, 0x34, 0x56)
    run.font.name = "Aptos"
    run.font.language_id = MSO_LANGUAGE_ID.ITALIAN
    paragraph = label.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.level = 1
    paragraph.space_before = Pt(12)
    paragraph.space_after = Pt(6)
    paragraph.line_spacing = 1.25
    paragraph_font = label.text_frame.paragraphs[0].font
    paragraph_font.bold = True
    paragraph_font.italic = True
    paragraph_font.underline = True
    paragraph_font.size = Pt(13)
    paragraph_font.fill.gradient()
    paragraph_font.fill.gradient_angle = 45.0
    paragraph_font.fill.gradient_stops[0].position = 0.25
    paragraph_font.fill.gradient_stops[0].color.rgb = RGBColor(0x0C, 0x22, 0x38)
    paragraph_font.name = "Aptos"
    paragraph_font.language_id = MSO_LANGUAGE_ID.GERMAN
    label.text_frame.paragraphs[0].add_line_break()
    label.text_frame.paragraphs[0].add_run().text = "beat"
    label.text_frame.add_paragraph().add_run().text = "Watch Q2"
    label.text_frame.paragraphs[1].clear()
    label.text_frame.paragraphs[1].add_run().text = "Watch Q2"
    label.font.bold = True
    label.font.italic = True
    label.font.underline = True
    label.font.size = Pt(14)
    label.font.fill.patterned()
    label.font.fill.pattern = MSO_PATTERN_TYPE.DIVOT
    label.font.fill.fore_color.rgb = RGBColor(0xBA, 0xDC, 0x0D)
    label.font.fill.back_color.rgb = RGBColor(0x12, 0x34, 0x56)
    label.font.name = "Aptos"
    label.font.language_id = MSO_LANGUAGE_ID.SPANISH


def _apply_wolfppt_dropin_chart_point_data_label_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_LABEL_POSITION
    from pptx.enum.dml import MSO_PATTERN_TYPE
    from pptx.enum.lang import MSO_LANGUAGE_ID
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Pt

    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart point data-label benchmark does not support fixture {fixture_id}"
        )
    label = _first_chart_shape(prs).chart.series[0].points[0].data_label
    label.position = XL_LABEL_POSITION.OUTSIDE_END
    label.text_frame.margin_left = TEXT_FRAME_MARGIN_EXPECTED["margin_left"]
    label.text_frame.margin_right = TEXT_FRAME_MARGIN_EXPECTED["margin_right"]
    label.text_frame.margin_top = TEXT_FRAME_MARGIN_EXPECTED["margin_top"]
    label.text_frame.margin_bottom = TEXT_FRAME_MARGIN_EXPECTED["margin_bottom"]
    label.text_frame.word_wrap = TEXT_FRAME_WORD_WRAP_EXPECTED
    label.text_frame.vertical_anchor = MSO_ANCHOR[
        TEXT_FRAME_VERTICAL_ANCHOR_EXPECTED_NAME
    ]
    label.text_frame.auto_size = MSO_AUTO_SIZE[TEXT_FRAME_AUTO_SIZE_EXPECTED_NAME]
    run = label.text_frame.paragraphs[0].add_run()
    run.text = "Q1 "
    run.font.bold = True
    run.font.italic = True
    run.font.underline = True
    run.font.size = Pt(11)
    run.font.fill.gradient()
    run.font.fill.gradient_angle = 45.0
    run.font.fill.gradient_stops[0].position = 0.25
    run.font.fill.gradient_stops[0].color.rgb = RGBColor(0x12, 0x34, 0x56)
    run.font.name = "Aptos"
    run.font.language_id = MSO_LANGUAGE_ID.ITALIAN
    paragraph = label.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.level = 1
    paragraph.space_before = Pt(12)
    paragraph.space_after = Pt(6)
    paragraph.line_spacing = 1.25
    paragraph_font = label.text_frame.paragraphs[0].font
    paragraph_font.bold = True
    paragraph_font.italic = True
    paragraph_font.underline = True
    paragraph_font.size = Pt(13)
    paragraph_font.fill.gradient()
    paragraph_font.fill.gradient_angle = 45.0
    paragraph_font.fill.gradient_stops[0].position = 0.25
    paragraph_font.fill.gradient_stops[0].color.rgb = RGBColor(0x0C, 0x22, 0x38)
    paragraph_font.name = "Aptos"
    paragraph_font.language_id = MSO_LANGUAGE_ID.GERMAN
    label.text_frame.paragraphs[0].add_line_break()
    label.text_frame.paragraphs[0].add_run().text = "beat"
    label.text_frame.add_paragraph().add_run().text = "Watch Q2"
    label.text_frame.paragraphs[1].clear()
    label.text_frame.paragraphs[1].add_run().text = "Watch Q2"
    label.font.bold = True
    label.font.italic = True
    label.font.underline = True
    label.font.size = Pt(14)
    label.font.fill.patterned()
    label.font.fill.pattern = MSO_PATTERN_TYPE.DIVOT
    label.font.fill.fore_color.rgb = RGBColor(0xBA, 0xDC, 0x0D)
    label.font.fill.back_color.rgb = RGBColor(0x12, 0x34, 0x56)
    label.font.name = "Aptos"
    label.font.language_id = MSO_LANGUAGE_ID.SPANISH


def _apply_python_pptx_dropin_chart_style_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart style benchmark does not support fixture {fixture_id}"
        )
    _first_chart_shape(prs).chart.chart_style = 10


def _apply_wolfppt_dropin_chart_style_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart style benchmark does not support fixture {fixture_id}"
        )
    _first_chart_shape(prs).chart.chart_style = 10


def _apply_wolfppt_dropin_chart_axis_property_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_AXIS_CROSSES, XL_TICK_LABEL_POSITION, XL_TICK_MARK
    from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_PATTERN_TYPE
    from pptx.enum.lang import MSO_LANGUAGE_ID
    from pptx.util import Pt

    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart axis property benchmark does not support fixture {fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.category_axis.visible = False
    chart.category_axis.major_tick_mark = XL_TICK_MARK.OUTSIDE
    chart.category_axis.minor_tick_mark = XL_TICK_MARK.INSIDE
    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    chart.category_axis.tick_labels.number_format = "mmm-yy"
    chart.category_axis.tick_labels.number_format_is_linked = True
    chart.category_axis.tick_labels.font.bold = True
    chart.category_axis.tick_labels.font.italic = True
    chart.category_axis.tick_labels.font.underline = True
    chart.category_axis.tick_labels.font.size = Pt(15)
    chart.category_axis.tick_labels.font.fill.gradient()
    chart.category_axis.tick_labels.font.fill.gradient_angle = 45.0
    chart.category_axis.tick_labels.font.fill.gradient_stops[0].position = 0.25
    chart.category_axis.tick_labels.font.fill.gradient_stops[0].color.rgb = RGBColor(
        0x12,
        0x34,
        0x56,
    )
    chart.category_axis.tick_labels.font.name = "Aptos"
    chart.category_axis.tick_labels.font.language_id = MSO_LANGUAGE_ID.GERMAN
    chart.category_axis.tick_labels.offset = 250
    chart.category_axis.reverse_order = True
    chart.category_axis.format.line.fill.solid()
    chart.category_axis.format.line.color.rgb = RGBColor(0x11, 0x22, 0x33)
    chart.category_axis.format.line.width = Pt(1.25)
    chart.category_axis.format.line.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT
    chart.value_axis.visible = False
    chart.value_axis.tick_labels.number_format = "$#,##0"
    chart.value_axis.tick_labels.number_format_is_linked = False
    chart.value_axis.tick_labels.font.bold = True
    chart.value_axis.tick_labels.font.italic = True
    chart.value_axis.tick_labels.font.underline = True
    chart.value_axis.tick_labels.font.size = Pt(15)
    chart.value_axis.tick_labels.font.fill.patterned()
    chart.value_axis.tick_labels.font.fill.pattern = MSO_PATTERN_TYPE.DIVOT
    chart.value_axis.tick_labels.font.fill.fore_color.rgb = RGBColor(0x65, 0x43, 0x21)
    chart.value_axis.tick_labels.font.fill.back_color.rgb = RGBColor(0xFE, 0xDC, 0xBA)
    chart.value_axis.tick_labels.font.name = "Arial"
    chart.value_axis.tick_labels.font.language_id = MSO_LANGUAGE_ID.FRENCH
    chart.value_axis.minimum_scale = 1.0
    chart.value_axis.maximum_scale = 20.0
    chart.value_axis.major_unit = 5.0
    chart.value_axis.minor_unit = 1.0
    chart.value_axis.format.line.fill.solid()
    chart.value_axis.format.line.color.rgb = RGBColor(0xAA, 0xBB, 0xCC)
    chart.value_axis.format.line.width = Pt(2.25)
    chart.value_axis.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.fill.solid()
    chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor(
        0x12,
        0x34,
        0x56,
    )
    chart.value_axis.major_gridlines.format.line.width = Pt(1.5)
    chart.value_axis.major_gridlines.format.line.dash_style = (
        MSO_LINE_DASH_STYLE.DASH_DOT
    )
    chart.value_axis.has_minor_gridlines = True
    chart.value_axis.crosses = XL_AXIS_CROSSES.MAXIMUM


def _apply_python_pptx_dropin_chart_format_edit(fixture_id: str, prs: Any) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_MARKER_STYLE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_PATTERN_TYPE
    from pptx.enum.lang import MSO_LANGUAGE_ID
    from pptx.util import Pt

    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart format benchmark does not support fixture {fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.font.bold = True
    chart.font.italic = True
    chart.font.underline = True
    chart.font.size = Pt(13)
    chart.font.fill.gradient()
    chart.font.fill.gradient_angle = 45.0
    chart.font.fill.gradient_stops[0].position = 0.25
    chart.font.fill.gradient_stops[0].color.rgb = RGBColor(0xCA, 0xFE, 0x01)
    chart.font.name = "Aptos"
    chart.font.language_id = MSO_LANGUAGE_ID.FRENCH
    series = chart.series[0]
    point = series.points[0]
    background_point = series.points[1]
    marker = point.marker
    background_marker = background_point.marker
    series.format.fill.patterned()
    series.format.fill.pattern = MSO_PATTERN_TYPE.DIVOT
    series.format.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    series.format.fill.back_color.rgb = RGBColor(0xAB, 0xCD, 0xEF)
    series.format.line.fill.patterned()
    series.format.line.fill.pattern = MSO_PATTERN_TYPE.DIVOT
    series.format.line.fill.fore_color.rgb = RGBColor(0x11, 0x22, 0x33)
    series.format.line.fill.back_color.rgb = RGBColor(0x44, 0x55, 0x66)
    series.format.line.width = Pt(2.5)
    series.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    point.format.fill.gradient()
    point.format.fill.gradient_angle = 45.0
    point.format.fill.gradient_stops[0].position = 0.25
    point.format.fill.gradient_stops[0].color.rgb = RGBColor(0x76, 0x54, 0x32)
    point.format.line.fill.background()
    point.format.line.width = Pt(4)
    point.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH_DOT
    background_point.format.fill.background()
    background_point.format.line.fill.patterned()
    background_point.format.line.fill.pattern = (
        MSO_PATTERN_TYPE.LIGHT_DOWNWARD_DIAGONAL
    )
    background_point.format.line.fill.fore_color.rgb = RGBColor(0x77, 0x88, 0x99)
    background_point.format.line.fill.back_color.rgb = RGBColor(0xAA, 0xBB, 0xCC)
    marker.format.fill.patterned()
    marker.format.fill.pattern = MSO_PATTERN_TYPE.LIGHT_DOWNWARD_DIAGONAL
    marker.format.fill.fore_color.rgb = RGBColor(0x01, 0x02, 0x03)
    marker.format.fill.back_color.rgb = RGBColor(0xF0, 0xE1, 0xD2)
    marker.format.line.fill.patterned()
    marker.format.line.fill.pattern = MSO_PATTERN_TYPE.PERCENT_10
    marker.format.line.fill.fore_color.rgb = RGBColor(0x04, 0x05, 0x06)
    marker.format.line.fill.back_color.rgb = RGBColor(0x07, 0x08, 0x09)
    marker.format.line.width = Pt(2)
    marker.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    marker.style = XL_MARKER_STYLE.DIAMOND
    marker.size = 9
    background_marker.format.fill.background()


def _apply_wolfppt_dropin_chart_format_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_MARKER_STYLE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_PATTERN_TYPE
    from pptx.enum.lang import MSO_LANGUAGE_ID
    from pptx.util import Pt

    if fixture_id not in CHART_EDIT_FIXTURES:
        raise RuntimeError(
            f"drop-in chart format benchmark does not support fixture {fixture_id}"
        )
    chart = _first_chart_shape(prs).chart
    chart.font.bold = True
    chart.font.italic = True
    chart.font.underline = True
    chart.font.size = Pt(13)
    chart.font.fill.gradient()
    chart.font.fill.gradient_angle = 45.0
    chart.font.fill.gradient_stops[0].position = 0.25
    chart.font.fill.gradient_stops[0].color.rgb = RGBColor(0xCA, 0xFE, 0x01)
    chart.font.name = "Aptos"
    chart.font.language_id = MSO_LANGUAGE_ID.FRENCH
    series = chart.series[0]
    point = series.points[0]
    background_point = series.points[1]
    marker = point.marker
    background_marker = background_point.marker
    series.format.fill.patterned()
    series.format.fill.pattern = MSO_PATTERN_TYPE.DIVOT
    series.format.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    series.format.fill.back_color.rgb = RGBColor(0xAB, 0xCD, 0xEF)
    series.format.line.fill.patterned()
    series.format.line.fill.pattern = MSO_PATTERN_TYPE.DIVOT
    series.format.line.fill.fore_color.rgb = RGBColor(0x11, 0x22, 0x33)
    series.format.line.fill.back_color.rgb = RGBColor(0x44, 0x55, 0x66)
    series.format.line.width = Pt(2.5)
    series.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    point.format.fill.gradient()
    point.format.fill.gradient_angle = 45.0
    point.format.fill.gradient_stops[0].position = 0.25
    point.format.fill.gradient_stops[0].color.rgb = RGBColor(0x76, 0x54, 0x32)
    point.format.line.fill.background()
    point.format.line.width = Pt(4)
    point.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH_DOT
    background_point.format.fill.background()
    background_point.format.line.fill.patterned()
    background_point.format.line.fill.pattern = (
        MSO_PATTERN_TYPE.LIGHT_DOWNWARD_DIAGONAL
    )
    background_point.format.line.fill.fore_color.rgb = RGBColor(0x77, 0x88, 0x99)
    background_point.format.line.fill.back_color.rgb = RGBColor(0xAA, 0xBB, 0xCC)
    marker.format.fill.patterned()
    marker.format.fill.pattern = MSO_PATTERN_TYPE.LIGHT_DOWNWARD_DIAGONAL
    marker.format.fill.fore_color.rgb = RGBColor(0x01, 0x02, 0x03)
    marker.format.fill.back_color.rgb = RGBColor(0xF0, 0xE1, 0xD2)
    marker.format.line.fill.patterned()
    marker.format.line.fill.pattern = MSO_PATTERN_TYPE.PERCENT_10
    marker.format.line.fill.fore_color.rgb = RGBColor(0x04, 0x05, 0x06)
    marker.format.line.fill.back_color.rgb = RGBColor(0x07, 0x08, 0x09)
    marker.format.line.width = Pt(2)
    marker.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    marker.style = XL_MARKER_STYLE.DIAMOND
    marker.size = 9
    background_marker.format.fill.background()
