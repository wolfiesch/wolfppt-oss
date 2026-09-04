"""Shape drop-in benchmark edit actions."""
# ruff: noqa: F401

from __future__ import annotations

from typing import Any

from .benchmark_cases import (
    ADD_CHART_EXPECTED,
    ADD_CHART_EXPECTED_TRANSFORM,
    ADD_SHAPE_EXPECTED_TEXT,
    ADD_SHAPE_EXPECTED_TRANSFORM,
    ADD_TEXTBOX_EXPECTED_TEXT,
    ADD_TEXTBOX_EXPECTED_TRANSFORM,
    GROUP_CHILD_GEOMETRY_EXPECTED,
    PICTURE_CROP_EXPECTED,
    SHAPE_ADJUSTMENT_EXPECTED,
    SHAPE_GEOMETRY_EXPECTED,
)
from .benchmark_dropin_media_actions import (
    _apply_python_pptx_dropin_add_group_ole_object,
    _apply_python_pptx_dropin_add_movie,
    _apply_python_pptx_dropin_add_movie_file_like,
    _apply_python_pptx_dropin_add_nested_group_ole_object,
    _apply_python_pptx_dropin_add_ole_object,
    _apply_python_pptx_dropin_add_ole_object_file_like,
    _apply_python_pptx_dropin_add_picture,
    _apply_python_pptx_dropin_add_picture_auto_size,
    _apply_python_pptx_dropin_add_picture_file_like,
    _apply_wolfppt_dropin_add_group_ole_object,
    _apply_wolfppt_dropin_add_movie,
    _apply_wolfppt_dropin_add_movie_file_like,
    _apply_wolfppt_dropin_add_nested_group_ole_object,
    _apply_wolfppt_dropin_add_ole_object,
    _apply_wolfppt_dropin_add_ole_object_file_like,
    _apply_wolfppt_dropin_add_picture,
    _apply_wolfppt_dropin_add_picture_auto_size,
    _apply_wolfppt_dropin_add_picture_file_like,
)
from .benchmark_dropin_shape_chart_actions import (
    ADD_CHART_TEMPLATE_FAMILY_TYPES,
    ADD_XY_CHART_TEMPLATE_FAMILY_TYPES,
    _apply_python_pptx_dropin_add_bar_chart,
    _apply_python_pptx_dropin_add_bubble_3d_chart,
    _apply_python_pptx_dropin_add_bubble_chart,
    _apply_python_pptx_dropin_add_chart,
    _apply_python_pptx_dropin_add_chart_template_family,
    _apply_python_pptx_dropin_add_chart_type,
    _apply_python_pptx_dropin_add_chart_types,
    _apply_python_pptx_dropin_add_hierarchical_chart,
    _apply_python_pptx_dropin_add_line_chart,
    _apply_python_pptx_dropin_add_pie_chart,
    _apply_python_pptx_dropin_add_xy_scatter_chart,
    _apply_python_pptx_dropin_add_xy_scatter_template_family,
    _apply_wolfppt_dropin_add_bar_chart,
    _apply_wolfppt_dropin_add_bubble_3d_chart,
    _apply_wolfppt_dropin_add_bubble_chart,
    _apply_wolfppt_dropin_add_chart,
    _apply_wolfppt_dropin_add_chart_template_family,
    _apply_wolfppt_dropin_add_chart_type,
    _apply_wolfppt_dropin_add_chart_types,
    _apply_wolfppt_dropin_add_hierarchical_chart,
    _apply_wolfppt_dropin_add_line_chart,
    _apply_wolfppt_dropin_add_pie_chart,
    _apply_wolfppt_dropin_add_xy_scatter_chart,
    _apply_wolfppt_dropin_add_xy_scatter_template_family,
)
from .benchmark_dropin_shape_connector_actions import (
    _apply_python_pptx_dropin_add_connector,
    _apply_python_pptx_dropin_add_deeper_nested_group_connector,
    _apply_python_pptx_dropin_add_group_connector,
    _apply_python_pptx_dropin_add_nested_group_connector,
    _apply_python_pptx_dropin_connector_connection_edit,
    _apply_python_pptx_dropin_connector_line_style_edit,
    _apply_python_pptx_dropin_existing_connector_connection_edit,
    _apply_python_pptx_dropin_group_connector_connection_edit,
    _apply_wolfppt_dropin_add_connector,
    _apply_wolfppt_dropin_add_deeper_nested_group_connector,
    _apply_wolfppt_dropin_add_group_connector,
    _apply_wolfppt_dropin_add_nested_group_connector,
    _apply_wolfppt_dropin_connector_connection_edit,
    _apply_wolfppt_dropin_connector_line_style_edit,
    _apply_wolfppt_dropin_existing_connector_connection_edit,
    _apply_wolfppt_dropin_group_connector_connection_edit,
)
from .benchmark_dropin_shape_creation_actions import (
    TOP_LEVEL_CREATION_FIXTURES,
    _apply_python_pptx_dropin_add_shape,
    _apply_python_pptx_dropin_add_slide,
    _apply_python_pptx_dropin_add_table,
    _apply_python_pptx_dropin_add_textbox,
    _apply_python_pptx_dropin_add_title_slide,
    _apply_wolfppt_dropin_add_shape,
    _apply_wolfppt_dropin_add_slide,
    _apply_wolfppt_dropin_add_table,
    _apply_wolfppt_dropin_add_textbox,
    _apply_wolfppt_dropin_add_title_slide,
)
from .benchmark_dropin_shape_freeform_actions import (
    _apply_python_pptx_dropin_build_deeper_nested_group_freeform,
    _apply_python_pptx_dropin_build_freeform,
    _apply_python_pptx_dropin_build_group_freeform,
    _apply_python_pptx_dropin_build_nested_group_freeform,
    _apply_wolfppt_dropin_build_deeper_nested_group_freeform,
    _apply_wolfppt_dropin_build_freeform,
    _apply_wolfppt_dropin_build_group_freeform,
    _apply_wolfppt_dropin_build_nested_group_freeform,
)
from .benchmark_dropin_shape_format_actions import (
    _apply_python_pptx_dropin_shape_fill_solid_edit,
    _apply_python_pptx_dropin_shape_gradient_fill_edit,
    _apply_python_pptx_dropin_shape_hyperlink_edit,
    _apply_python_pptx_dropin_shape_line_fill_background_edit,
    _apply_python_pptx_dropin_shape_name_edit,
    _apply_python_pptx_dropin_shape_patterned_fill_edit,
    _apply_python_pptx_dropin_shape_rotation_edit,
    _apply_python_pptx_dropin_shape_shadow_edit,
    _apply_python_pptx_dropin_shape_style_edit,
    _apply_python_pptx_dropin_shape_target_new_slide_edit,
    _apply_python_pptx_dropin_shape_target_slide_edit,
    _apply_python_pptx_dropin_shape_theme_color_edit,
    _apply_wolfppt_dropin_shape_fill_solid_edit,
    _apply_wolfppt_dropin_shape_gradient_fill_edit,
    _apply_wolfppt_dropin_shape_hyperlink_edit,
    _apply_wolfppt_dropin_shape_line_fill_background_edit,
    _apply_wolfppt_dropin_shape_name_edit,
    _apply_wolfppt_dropin_shape_patterned_fill_edit,
    _apply_wolfppt_dropin_shape_rotation_edit,
    _apply_wolfppt_dropin_shape_shadow_edit,
    _apply_wolfppt_dropin_shape_style_edit,
    _apply_wolfppt_dropin_shape_target_new_slide_edit,
    _apply_wolfppt_dropin_shape_target_slide_edit,
    _apply_wolfppt_dropin_shape_theme_color_edit,
)
from .benchmark_dropin_common_actions import _shape_at, _shape_collection_at
from .presentation import Presentation as WolfPresentation

PICTURE_CROP_FIXTURES = ("media/png_picture", "workloads/mixed_real_world_deck")
GROUP_SHAPE_TARGETS = {
    "shapes/grouped_shapes": (0, 0),
    "workloads/management_reporting_deck": (2, 1),
    "workloads/customer_success_review_pack": (4, 1),
}


def _dropin_picture_shape(fixture_id: str, prs: Any) -> Any:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if fixture_id not in PICTURE_CROP_FIXTURES:
        raise RuntimeError(
            f"drop-in picture crop benchmark does not support fixture {fixture_id}"
        )
    if isinstance(prs, WolfPresentation):
        for slide_index, slide in enumerate(prs.slides):
            payload = getattr(slide, "_payload", None)
            if not isinstance(payload, dict):
                continue
            for hint in payload.get("shape_shell_hints") or []:
                if not isinstance(hint, dict):
                    continue
                if hint.get("has_picture") or hint.get("kind") == "picture":
                    return _shape_at(prs, slide_index, int(hint["_shape_index"]))
    for slide in prs.slides:
        for shape in slide.shapes:
            if (
                getattr(shape, "has_picture", False)
                or getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE
            ):
                return shape
    raise RuntimeError(f"fixture {fixture_id} does not contain a picture shape")


def _dropin_group_shape(fixture_id: str, prs: Any, operation: str) -> Any:
    if fixture_id not in GROUP_SHAPE_TARGETS:
        raise RuntimeError(
            f"drop-in {operation} benchmark does not support fixture {fixture_id}"
        )
    slide_index, group_index = GROUP_SHAPE_TARGETS[fixture_id]
    if isinstance(prs, WolfPresentation):
        return _shape_at(prs, slide_index, group_index)
    return prs.slides[slide_index].shapes[group_index]


def _apply_python_pptx_dropin_add_group_textbox(fixture_id: str, prs: Any) -> None:
    if fixture_id not in GROUP_SHAPE_TARGETS:
        raise RuntimeError(
            f"drop-in add-group-textbox benchmark does not support fixture {fixture_id}"
        )
    slide_index, group_index = GROUP_SHAPE_TARGETS[fixture_id]
    shape = (
        prs.slides[slide_index]
        .shapes[group_index]
        .shapes.add_textbox(
            ADD_TEXTBOX_EXPECTED_TRANSFORM["x"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["y"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["cx"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["cy"],
        )
    )
    shape.text = ADD_TEXTBOX_EXPECTED_TEXT


def _apply_wolfppt_dropin_add_group_textbox(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    shape = (
        _dropin_group_shape(fixture_id, prs, "add-group-textbox")
        .shapes.add_textbox(
            ADD_TEXTBOX_EXPECTED_TRANSFORM["x"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["y"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["cx"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["cy"],
        )
    )
    shape.text = ADD_TEXTBOX_EXPECTED_TEXT


def _apply_python_pptx_dropin_add_group_chart(fixture_id: str, prs: Any) -> None:
    if fixture_id not in GROUP_SHAPE_TARGETS:
        raise RuntimeError(
            f"drop-in add-group-chart benchmark does not support fixture {fixture_id}"
        )
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    chart_data = CategoryChartData()
    chart_data.categories = ADD_CHART_EXPECTED["categories"]
    for series in ADD_CHART_EXPECTED["series"]:
        chart_data.add_series(series["name"], series["values"])
    slide_index, group_index = GROUP_SHAPE_TARGETS[fixture_id]
    prs.slides[slide_index].shapes[group_index].shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        ADD_CHART_EXPECTED_TRANSFORM["x"],
        ADD_CHART_EXPECTED_TRANSFORM["y"],
        ADD_CHART_EXPECTED_TRANSFORM["cx"],
        ADD_CHART_EXPECTED_TRANSFORM["cy"],
        chart_data,
    )


def _apply_wolfppt_dropin_add_group_chart(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    chart_data = CategoryChartData()
    chart_data.categories = ADD_CHART_EXPECTED["categories"]
    for series in ADD_CHART_EXPECTED["series"]:
        chart_data.add_series(series["name"], series["values"])
    _dropin_group_shape(fixture_id, prs, "add-group-chart").shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        ADD_CHART_EXPECTED_TRANSFORM["x"],
        ADD_CHART_EXPECTED_TRANSFORM["y"],
        ADD_CHART_EXPECTED_TRANSFORM["cx"],
        ADD_CHART_EXPECTED_TRANSFORM["cy"],
        chart_data,
    )


def _apply_python_pptx_dropin_add_nested_group_chart(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in GROUP_SHAPE_TARGETS:
        raise RuntimeError(
            f"drop-in add-nested-group-chart benchmark does not support fixture {fixture_id}"
        )
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    chart_data = CategoryChartData()
    chart_data.categories = ADD_CHART_EXPECTED["categories"]
    for series in ADD_CHART_EXPECTED["series"]:
        chart_data.add_series(series["name"], series["values"])
    slide_index, group_index = GROUP_SHAPE_TARGETS[fixture_id]
    nested = prs.slides[slide_index].shapes[group_index].shapes.add_group_shape()
    nested.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        ADD_CHART_EXPECTED_TRANSFORM["x"],
        ADD_CHART_EXPECTED_TRANSFORM["y"],
        ADD_CHART_EXPECTED_TRANSFORM["cx"],
        ADD_CHART_EXPECTED_TRANSFORM["cy"],
        chart_data,
    )


def _apply_wolfppt_dropin_add_nested_group_chart(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    chart_data = CategoryChartData()
    chart_data.categories = ADD_CHART_EXPECTED["categories"]
    for series in ADD_CHART_EXPECTED["series"]:
        chart_data.add_series(series["name"], series["values"])
    nested = _dropin_group_shape(
        fixture_id, prs, "add-nested-group-chart"
    ).shapes.add_group_shape()
    nested.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        ADD_CHART_EXPECTED_TRANSFORM["x"],
        ADD_CHART_EXPECTED_TRANSFORM["y"],
        ADD_CHART_EXPECTED_TRANSFORM["cx"],
        ADD_CHART_EXPECTED_TRANSFORM["cy"],
        chart_data,
    )


def _apply_python_pptx_dropin_add_deeper_nested_group_chart(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in GROUP_SHAPE_TARGETS:
        raise RuntimeError(
            "drop-in add-deeper-nested-group-chart benchmark does not support "
            f"fixture {fixture_id}"
        )
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    chart_data = CategoryChartData()
    chart_data.categories = ADD_CHART_EXPECTED["categories"]
    for series in ADD_CHART_EXPECTED["series"]:
        chart_data.add_series(series["name"], series["values"])
    slide_index, group_index = GROUP_SHAPE_TARGETS[fixture_id]
    (
        prs.slides[slide_index]
        .shapes[group_index]
        .shapes.add_group_shape()
        .shapes.add_group_shape()
        .shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            ADD_CHART_EXPECTED_TRANSFORM["x"],
            ADD_CHART_EXPECTED_TRANSFORM["y"],
            ADD_CHART_EXPECTED_TRANSFORM["cx"],
            ADD_CHART_EXPECTED_TRANSFORM["cy"],
            chart_data,
        )
    )


def _apply_wolfppt_dropin_add_deeper_nested_group_chart(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    chart_data = CategoryChartData()
    chart_data.categories = ADD_CHART_EXPECTED["categories"]
    for series in ADD_CHART_EXPECTED["series"]:
        chart_data.add_series(series["name"], series["values"])
    (
        _dropin_group_shape(fixture_id, prs, "add-deeper-nested-group-chart")
        .shapes.add_group_shape()
        .shapes.add_group_shape()
        .shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            ADD_CHART_EXPECTED_TRANSFORM["x"],
            ADD_CHART_EXPECTED_TRANSFORM["y"],
            ADD_CHART_EXPECTED_TRANSFORM["cx"],
            ADD_CHART_EXPECTED_TRANSFORM["cy"],
            chart_data,
        )
    )


def _apply_python_pptx_dropin_shape_adjustment_edit(fixture_id: str, prs: Any) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in shape adjustment benchmark does not support fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_SHAPE

    shape = prs.slides[0].shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        ADD_SHAPE_EXPECTED_TRANSFORM["x"],
        ADD_SHAPE_EXPECTED_TRANSFORM["y"],
        ADD_SHAPE_EXPECTED_TRANSFORM["cx"],
        ADD_SHAPE_EXPECTED_TRANSFORM["cy"],
    )
    shape.adjustments[0] = SHAPE_ADJUSTMENT_EXPECTED


def _apply_wolfppt_dropin_shape_adjustment_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in shape adjustment benchmark does not support fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_SHAPE

    shape = _shape_collection_at(prs, 0).add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        ADD_SHAPE_EXPECTED_TRANSFORM["x"],
        ADD_SHAPE_EXPECTED_TRANSFORM["y"],
        ADD_SHAPE_EXPECTED_TRANSFORM["cx"],
        ADD_SHAPE_EXPECTED_TRANSFORM["cy"],
    )
    shape.adjustments[0] = SHAPE_ADJUSTMENT_EXPECTED


def _apply_python_pptx_dropin_shape_geometry_edit(fixture_id: str, prs: Any) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in shape geometry benchmark does not support fixture {fixture_id}"
        )
    shape = prs.slides[0].shapes[0]
    shape.left = SHAPE_GEOMETRY_EXPECTED["x"]
    shape.top = SHAPE_GEOMETRY_EXPECTED["y"]
    shape.width = SHAPE_GEOMETRY_EXPECTED["cx"]
    shape.height = SHAPE_GEOMETRY_EXPECTED["cy"]


def _apply_wolfppt_dropin_shape_geometry_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in shape geometry benchmark does not support fixture {fixture_id}"
        )
    shape = _shape_at(prs, 0, 0)
    shape.left = SHAPE_GEOMETRY_EXPECTED["x"]
    shape.top = SHAPE_GEOMETRY_EXPECTED["y"]
    shape.width = SHAPE_GEOMETRY_EXPECTED["cx"]
    shape.height = SHAPE_GEOMETRY_EXPECTED["cy"]


def _apply_python_pptx_dropin_group_child_geometry_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in GROUP_SHAPE_TARGETS:
        raise RuntimeError(
            f"drop-in group child geometry benchmark does not support fixture {fixture_id}"
        )
    slide_index, group_index = GROUP_SHAPE_TARGETS[fixture_id]
    child = prs.slides[slide_index].shapes[group_index].shapes[0]
    child.left = GROUP_CHILD_GEOMETRY_EXPECTED["x"]
    child.top = GROUP_CHILD_GEOMETRY_EXPECTED["y"]
    child.width = GROUP_CHILD_GEOMETRY_EXPECTED["cx"]
    child.height = GROUP_CHILD_GEOMETRY_EXPECTED["cy"]


def _apply_wolfppt_dropin_group_child_geometry_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    child = _dropin_group_shape(fixture_id, prs, "group child geometry").shapes[0]
    child.left = GROUP_CHILD_GEOMETRY_EXPECTED["x"]
    child.top = GROUP_CHILD_GEOMETRY_EXPECTED["y"]
    child.width = GROUP_CHILD_GEOMETRY_EXPECTED["cx"]
    child.height = GROUP_CHILD_GEOMETRY_EXPECTED["cy"]


def _apply_python_pptx_dropin_group_child_adjustment_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in GROUP_SHAPE_TARGETS:
        raise RuntimeError(
            f"drop-in group child adjustment benchmark does not support fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_SHAPE

    slide_index, group_index = GROUP_SHAPE_TARGETS[fixture_id]
    child = prs.slides[slide_index].shapes[group_index].shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        ADD_SHAPE_EXPECTED_TRANSFORM["x"],
        ADD_SHAPE_EXPECTED_TRANSFORM["y"],
        ADD_SHAPE_EXPECTED_TRANSFORM["cx"],
        ADD_SHAPE_EXPECTED_TRANSFORM["cy"],
    )
    child.adjustments[0] = SHAPE_ADJUSTMENT_EXPECTED


def _apply_wolfppt_dropin_group_child_adjustment_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    child = _dropin_group_shape(
        fixture_id, prs, "group child adjustment"
    ).shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        ADD_SHAPE_EXPECTED_TRANSFORM["x"],
        ADD_SHAPE_EXPECTED_TRANSFORM["y"],
        ADD_SHAPE_EXPECTED_TRANSFORM["cx"],
        ADD_SHAPE_EXPECTED_TRANSFORM["cy"],
    )
    child.adjustments[0] = SHAPE_ADJUSTMENT_EXPECTED


def _apply_python_pptx_dropin_add_group_auto_shape(fixture_id: str, prs: Any) -> None:
    if fixture_id not in GROUP_SHAPE_TARGETS:
        raise RuntimeError(
            f"drop-in add-group-auto-shape benchmark does not support fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_SHAPE

    slide_index, group_index = GROUP_SHAPE_TARGETS[fixture_id]
    shape = (
        prs.slides[slide_index]
        .shapes[group_index]
        .shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            ADD_SHAPE_EXPECTED_TRANSFORM["x"],
            ADD_SHAPE_EXPECTED_TRANSFORM["y"],
            ADD_SHAPE_EXPECTED_TRANSFORM["cx"],
            ADD_SHAPE_EXPECTED_TRANSFORM["cy"],
        )
    )
    shape.text = ADD_SHAPE_EXPECTED_TEXT


def _apply_wolfppt_dropin_add_group_auto_shape(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    shape = (
        _dropin_group_shape(fixture_id, prs, "add-group-auto-shape")
        .shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            ADD_SHAPE_EXPECTED_TRANSFORM["x"],
            ADD_SHAPE_EXPECTED_TRANSFORM["y"],
            ADD_SHAPE_EXPECTED_TRANSFORM["cx"],
            ADD_SHAPE_EXPECTED_TRANSFORM["cy"],
        )
    )
    shape.text = ADD_SHAPE_EXPECTED_TEXT


def _apply_python_pptx_dropin_shape_line_element_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in shape line element benchmark does not support fixture {fixture_id}"
        )
    prs.slides[0].shapes[0].get_or_add_ln()


def _apply_wolfppt_dropin_shape_line_element_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in shape line element benchmark does not support fixture {fixture_id}"
        )
    _shape_at(prs, 0, 0).get_or_add_ln()


def _apply_python_pptx_dropin_add_group_shape(fixture_id: str, prs: Any) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-group-shape benchmark does not support fixture {fixture_id}"
        )
    prs.slides[0].shapes.add_group_shape()


def _apply_wolfppt_dropin_add_group_shape(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-group-shape benchmark does not support fixture {fixture_id}"
        )
    _shape_collection_at(prs, 0).add_group_shape()


def _apply_python_pptx_dropin_group_existing_shapes(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            "drop-in group-existing-shapes benchmark does not support fixture "
            f"{fixture_id}"
        )
    slide = prs.slides[0]
    slide.shapes.add_group_shape([slide.shapes[1], slide.shapes[0]])


def _apply_wolfppt_dropin_group_existing_shapes(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            "drop-in group-existing-shapes benchmark does not support fixture "
            f"{fixture_id}"
        )
    shapes = _shape_collection_at(prs, 0)
    shapes.add_group_shape([shapes[1], shapes[0]])


def _apply_python_pptx_dropin_add_nested_group_shape(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in GROUP_SHAPE_TARGETS:
        raise RuntimeError(
            f"drop-in add-nested-group-shape benchmark does not support fixture {fixture_id}"
        )
    slide_index, group_index = GROUP_SHAPE_TARGETS[fixture_id]
    prs.slides[slide_index].shapes[group_index].shapes.add_group_shape()


def _apply_wolfppt_dropin_add_nested_group_shape(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _dropin_group_shape(
        fixture_id, prs, "add-nested-group-shape"
    ).shapes.add_group_shape()


def _apply_python_pptx_dropin_add_deeper_nested_group_shape(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in GROUP_SHAPE_TARGETS:
        raise RuntimeError(
            "drop-in add-deeper-nested-group-shape benchmark does not support "
            f"fixture {fixture_id}"
        )
    slide_index, group_index = GROUP_SHAPE_TARGETS[fixture_id]
    (
        prs.slides[slide_index]
        .shapes[group_index]
        .shapes.add_group_shape()
        .shapes.add_group_shape()
    )


def _apply_wolfppt_dropin_add_deeper_nested_group_shape(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    (
        _dropin_group_shape(fixture_id, prs, "add-deeper-nested-group-shape")
        .shapes.add_group_shape()
        .shapes.add_group_shape()
    )


def _apply_python_pptx_dropin_group_existing_children(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id != "workloads/customer_success_review_pack":
        raise RuntimeError(
            "drop-in group-existing-children benchmark does not support fixture "
            f"{fixture_id}"
        )
    group = prs.slides[4].shapes[1]
    group.shapes.add_group_shape([group.shapes[1], group.shapes[0]])


def _apply_wolfppt_dropin_group_existing_children(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id != "workloads/customer_success_review_pack":
        raise RuntimeError(
            "drop-in group-existing-children benchmark does not support fixture "
            f"{fixture_id}"
        )
    group = _dropin_group_shape(fixture_id, prs, "group-existing-children")
    group.shapes.add_group_shape([group.shapes[1], group.shapes[0]])


def _apply_python_pptx_dropin_add_nested_group_auto_shape(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in GROUP_SHAPE_TARGETS:
        raise RuntimeError(
            f"drop-in add-nested-group-auto-shape benchmark does not support fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_SHAPE

    slide_index, group_index = GROUP_SHAPE_TARGETS[fixture_id]
    shape = (
        prs.slides[slide_index]
        .shapes[group_index]
        .shapes.add_group_shape()
        .shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            ADD_SHAPE_EXPECTED_TRANSFORM["x"],
            ADD_SHAPE_EXPECTED_TRANSFORM["y"],
            ADD_SHAPE_EXPECTED_TRANSFORM["cx"],
            ADD_SHAPE_EXPECTED_TRANSFORM["cy"],
        )
    )
    shape.text = ADD_SHAPE_EXPECTED_TEXT


def _apply_wolfppt_dropin_add_nested_group_auto_shape(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    shape = (
        _dropin_group_shape(fixture_id, prs, "add-nested-group-auto-shape")
        .shapes.add_group_shape()
        .shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            ADD_SHAPE_EXPECTED_TRANSFORM["x"],
            ADD_SHAPE_EXPECTED_TRANSFORM["y"],
            ADD_SHAPE_EXPECTED_TRANSFORM["cx"],
            ADD_SHAPE_EXPECTED_TRANSFORM["cy"],
        )
    )
    shape.text = ADD_SHAPE_EXPECTED_TEXT


def _apply_python_pptx_dropin_add_nested_group_textbox(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in GROUP_SHAPE_TARGETS:
        raise RuntimeError(
            f"drop-in add-nested-group-textbox benchmark does not support fixture {fixture_id}"
        )
    slide_index, group_index = GROUP_SHAPE_TARGETS[fixture_id]
    shape = (
        prs.slides[slide_index]
        .shapes[group_index]
        .shapes.add_group_shape()
        .shapes.add_textbox(
            ADD_TEXTBOX_EXPECTED_TRANSFORM["x"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["y"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["cx"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["cy"],
        )
    )
    shape.text = ADD_TEXTBOX_EXPECTED_TEXT


def _apply_wolfppt_dropin_add_nested_group_textbox(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    shape = (
        _dropin_group_shape(fixture_id, prs, "add-nested-group-textbox")
        .shapes.add_group_shape()
        .shapes.add_textbox(
            ADD_TEXTBOX_EXPECTED_TRANSFORM["x"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["y"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["cx"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["cy"],
        )
    )
    shape.text = ADD_TEXTBOX_EXPECTED_TEXT


def _apply_python_pptx_dropin_add_deeper_nested_group_textbox(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in GROUP_SHAPE_TARGETS:
        raise RuntimeError(
            "drop-in add-deeper-nested-group-textbox benchmark does not support "
            f"fixture {fixture_id}"
        )
    slide_index, group_index = GROUP_SHAPE_TARGETS[fixture_id]
    shape = (
        prs.slides[slide_index]
        .shapes[group_index]
        .shapes.add_group_shape()
        .shapes.add_group_shape()
        .shapes.add_textbox(
            ADD_TEXTBOX_EXPECTED_TRANSFORM["x"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["y"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["cx"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["cy"],
        )
    )
    shape.text = ADD_TEXTBOX_EXPECTED_TEXT


def _apply_wolfppt_dropin_add_deeper_nested_group_textbox(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    shape = (
        _dropin_group_shape(fixture_id, prs, "add-deeper-nested-group-textbox")
        .shapes.add_group_shape()
        .shapes.add_group_shape()
        .shapes.add_textbox(
            ADD_TEXTBOX_EXPECTED_TRANSFORM["x"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["y"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["cx"],
            ADD_TEXTBOX_EXPECTED_TRANSFORM["cy"],
        )
    )
    shape.text = ADD_TEXTBOX_EXPECTED_TEXT


def _apply_python_pptx_dropin_add_deeper_nested_group_auto_shape(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in GROUP_SHAPE_TARGETS:
        raise RuntimeError(
            "drop-in add-deeper-nested-group-auto-shape benchmark does not support "
            f"fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_SHAPE

    slide_index, group_index = GROUP_SHAPE_TARGETS[fixture_id]
    shape = (
        prs.slides[slide_index]
        .shapes[group_index]
        .shapes.add_group_shape()
        .shapes.add_group_shape()
        .shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            ADD_SHAPE_EXPECTED_TRANSFORM["x"],
            ADD_SHAPE_EXPECTED_TRANSFORM["y"],
            ADD_SHAPE_EXPECTED_TRANSFORM["cx"],
            ADD_SHAPE_EXPECTED_TRANSFORM["cy"],
        )
    )
    shape.text = ADD_SHAPE_EXPECTED_TEXT


def _apply_wolfppt_dropin_add_deeper_nested_group_auto_shape(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    shape = (
        _dropin_group_shape(fixture_id, prs, "add-deeper-nested-group-auto-shape")
        .shapes.add_group_shape()
        .shapes.add_group_shape()
        .shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            ADD_SHAPE_EXPECTED_TRANSFORM["x"],
            ADD_SHAPE_EXPECTED_TRANSFORM["y"],
            ADD_SHAPE_EXPECTED_TRANSFORM["cx"],
            ADD_SHAPE_EXPECTED_TRANSFORM["cy"],
        )
    )
    shape.text = ADD_SHAPE_EXPECTED_TEXT


def _apply_python_pptx_dropin_clone_layout_placeholders(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id != "text_basic/title_body_bullets":
        raise RuntimeError(
            f"drop-in clone-layout-placeholders benchmark does not support fixture {fixture_id}"
        )
    prs.slides[0].shapes.clone_layout_placeholders(prs.slide_layouts[1])


def _apply_wolfppt_dropin_clone_layout_placeholders(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id != "text_basic/title_body_bullets":
        raise RuntimeError(
            f"drop-in clone-layout-placeholders benchmark does not support fixture {fixture_id}"
        )
    _shape_collection_at(prs, 0).clone_layout_placeholders(prs.slide_layouts[1])


def _apply_python_pptx_dropin_picture_crop_edit(fixture_id: str, prs: Any) -> None:
    picture = _dropin_picture_shape(fixture_id, prs)
    picture.crop_left = PICTURE_CROP_EXPECTED["left"]
    picture.crop_right = PICTURE_CROP_EXPECTED["right"]
    picture.crop_top = PICTURE_CROP_EXPECTED["top"]
    picture.crop_bottom = PICTURE_CROP_EXPECTED["bottom"]

def _apply_wolfppt_dropin_picture_crop_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    picture = _dropin_picture_shape(fixture_id, prs)
    picture.crop_left = PICTURE_CROP_EXPECTED["left"]
    picture.crop_right = PICTURE_CROP_EXPECTED["right"]
    picture.crop_top = PICTURE_CROP_EXPECTED["top"]
    picture.crop_bottom = PICTURE_CROP_EXPECTED["bottom"]
