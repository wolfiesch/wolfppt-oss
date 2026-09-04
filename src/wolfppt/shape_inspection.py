"""Shape payload/XML inspection helpers for the Python facade."""

from __future__ import annotations

import zipfile
from collections.abc import Sequence
from typing import Any
from xml.etree import ElementTree as ET

from .shape_action_inspection import (
    _shape_click_action_name,
    _shape_click_action_value,
    _shape_hyperlink_address,
    _shape_hyperlink_click,
    _shape_target_slide,
)
from .shape_picture_inspection import (
    MediaFormat,
    PictureImage,
    _picture_crop,
    _picture_crop_value,
    _picture_image,
    _replace_picture_image,
    _set_picture_crop_value,
    _shape_media_format,
    _shape_media_type,
    _shape_picture_crop,
    _shape_poster_frame,
)
from .shape_payloads import _shape_transform
from .shape_preset_maps import (
    _AUTO_SHAPE_ADJUSTMENT_DEFAULTS,
    _AUTO_SHAPE_PRESET_BY_ID,
    _AUTO_SHAPE_PRESET_BY_NAME,
    _AUTO_SHAPE_TYPE_BY_PRESET,
    _CONNECTOR_PRESET_BY_ID,
    _CONNECTOR_PRESET_BY_NAME,
)
from .shape_xml import _shape_xml_element

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

__all__ = (
    "MediaFormat",
    "PictureImage",
    "_picture_crop",
    "_picture_crop_value",
    "_picture_image",
    "_replace_picture_image",
    "_set_picture_crop_value",
    "_shape_click_action_name",
    "_shape_click_action_value",
    "_shape_hyperlink_address",
    "_shape_hyperlink_click",
    "_shape_media_format",
    "_shape_media_type",
    "_shape_picture_crop",
    "_shape_poster_frame",
    "_shape_target_slide",
)


def _set_shape_transform_value(shape: Any, key: str, value: int) -> None:
    transform = _shape_transform(shape)
    if transform.get(key) == value:
        return
    transform[key] = value
    shape._payload["transform"] = transform
    shape._payload["effective_transform"] = transform
    if shape._payload.get("_deep_nested_group_child"):
        shape._slide._presentation._queue_deeper_group_child_geometry(
            shape._slide._index,
            int(shape._payload["_group_index"]),
            int(shape._payload["_nested_group_child_index"]),
            int(shape._payload["_deeper_group_child_index"]),
            int(shape._payload["_group_child_index"]),
            transform["x"],
            transform["y"],
            transform["cx"],
            transform["cy"],
        )
        return
    if shape._payload.get("_nested_group_child"):
        shape._slide._presentation._queue_nested_group_child_geometry(
            shape._slide._index,
            int(shape._payload["_group_index"]),
            int(shape._payload["_nested_group_child_index"]),
            int(shape._payload["_group_child_index"]),
            transform["x"],
            transform["y"],
            transform["cx"],
            transform["cy"],
        )
        return
    if shape._payload.get("_group_child"):
        shape._slide._presentation._queue_group_child_geometry(
            shape._slide._index,
            int(shape._payload["_group_index"]),
            int(shape._payload["_group_child_index"]),
            transform["x"],
            transform["y"],
            transform["cx"],
            transform["cy"],
        )
        return
    shape._slide._presentation._queue_shape_geometry(
        shape._slide._index,
        shape._index,
        transform["x"],
        transform["y"],
        transform["cx"],
        transform["cy"],
    )


def _placeholder_type_xml(shape: Any) -> str | None:
    raw = shape._payload.get("placeholder_type")
    if raw in (None, ""):
        return None
    return str(raw)


def _placeholder_raw_value(shape: Any, key: str) -> str | None:
    payload = getattr(shape, "_payload", {})
    raw = payload.get(key) if isinstance(payload, dict) else None
    return _string_or_none(raw)


def _payload_string(payload: dict[str, Any], key: str) -> str | None:
    return _string_or_none(payload.get(key))


def _string_or_none(raw: Any) -> str | None:
    if raw in (None, ""):
        return None
    return str(raw)


def _cloneable_layout_placeholders(slide_layout: Any) -> list[Any]:
    if hasattr(slide_layout, "iter_cloneable_placeholders"):
        return list(slide_layout.iter_cloneable_placeholders())
    placeholders = getattr(slide_layout, "placeholders", ())
    return [
        placeholder
        for placeholder in placeholders
        if _placeholder_type_xml(placeholder) not in {"dt", "ftr", "sldNum"}
    ]


def _next_shape_id(shapes: Sequence[Any]) -> int:
    shape_ids = [1]
    for shape in shapes:
        shape_id = getattr(shape, "shape_id", None)
        if isinstance(shape_id, int):
            shape_ids.append(shape_id)
        elif isinstance(shape_id, str) and shape_id.isdigit():
            shape_ids.append(int(shape_id))
    return max(shape_ids) + 1


def _next_placeholder_name(
    shapes: Sequence[Any],
    shape_id: int,
    placeholder_type: str | None,
    placeholder_orient: str | None,
) -> str:
    basename = _placeholder_basename_from_xml(placeholder_type)
    if placeholder_orient == "vert":
        basename = f"Vertical {basename}"
    existing_names = {getattr(shape, "name", "") for shape in shapes}
    suffix = shape_id - 1
    while True:
        name = f"{basename} {suffix}"
        if name not in existing_names:
            return name
        suffix += 1


def _placeholder_basename_from_xml(placeholder_type: str | None) -> str:
    return {
        "body": "Text Placeholder",
        "chart": "Chart Placeholder",
        "clipArt": "ClipArt Placeholder",
        "ctrTitle": "Title",
        "dgm": "SmartArt Placeholder",
        "dt": "Date Placeholder",
        "ftr": "Footer Placeholder",
        "media": "Media Placeholder",
        "pic": "Picture Placeholder",
        "sldImg": "Image Placeholder",
        "sldNum": "Slide Number Placeholder",
        "subTitle": "Subtitle",
        "tbl": "Table Placeholder",
        "title": "Title",
    }.get(placeholder_type, "Content Placeholder")


def _placeholder_type_value(raw: str | None) -> Any:
    try:
        from pptx.enum.shapes import PP_PLACEHOLDER
    except ImportError:
        return raw or "obj"

    if raw is None:
        return PP_PLACEHOLDER.OBJECT
    for member in PP_PLACEHOLDER:
        if getattr(member, "xml_value", None) == raw:
            return member
    return raw


def _placeholder_basename(ph_type: Any) -> str:
    from pptx.enum.shapes import PP_PLACEHOLDER

    return {
        PP_PLACEHOLDER.BITMAP: "ClipArt Placeholder",
        PP_PLACEHOLDER.BODY: "Text Placeholder",
        PP_PLACEHOLDER.CENTER_TITLE: "Title",
        PP_PLACEHOLDER.CHART: "Chart Placeholder",
        PP_PLACEHOLDER.DATE: "Date Placeholder",
        PP_PLACEHOLDER.FOOTER: "Footer Placeholder",
        PP_PLACEHOLDER.HEADER: "Header Placeholder",
        PP_PLACEHOLDER.MEDIA_CLIP: "Media Placeholder",
        PP_PLACEHOLDER.OBJECT: "Content Placeholder",
        PP_PLACEHOLDER.ORG_CHART: "SmartArt Placeholder",
        PP_PLACEHOLDER.PICTURE: "Picture Placeholder",
        PP_PLACEHOLDER.SLIDE_NUMBER: "Slide Number Placeholder",
        PP_PLACEHOLDER.SUBTITLE: "Subtitle",
        PP_PLACEHOLDER.TABLE: "Table Placeholder",
        PP_PLACEHOLDER.TITLE: "Title",
    }[ph_type]


def _notes_placeholder_basename(ph_type: Any) -> str:
    from pptx.enum.shapes import PP_PLACEHOLDER

    return {
        PP_PLACEHOLDER.BODY: "Notes Placeholder",
        PP_PLACEHOLDER.DATE: "Date Placeholder",
        PP_PLACEHOLDER.FOOTER: "Footer Placeholder",
        PP_PLACEHOLDER.HEADER: "Header Placeholder",
        PP_PLACEHOLDER.SLIDE_IMAGE: "Slide Image Placeholder",
        PP_PLACEHOLDER.SLIDE_NUMBER: "Slide Number Placeholder",
    }[ph_type]


def _shape_collection_table_count(shapes: Sequence[Any]) -> int:
    return sum(len(shape._payload.get("tables", [])) for shape in shapes)


def _shape_type_value(shape: Any) -> Any:
    kind = str(shape._payload.get("kind") or "")
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        if shape.is_placeholder:
            return "PLACEHOLDER"
        if kind == "connector":
            return "LINE"
        if shape.has_chart:
            return "CHART"
        if shape.has_table:
            return "TABLE"
        if kind == "ole_object":
            return "EMBEDDED_OLE_OBJECT"
        if kind == "movie":
            return "MEDIA"
        if kind == "picture":
            return "PICTURE"
        if kind == "group":
            return "GROUP"
        if kind == "freeform" or _shape_is_freeform(shape):
            return "FREEFORM"
        if kind == "shape":
            if _shape_is_text_box(shape):
                return "TEXT_BOX"
            return "AUTO_SHAPE"
        return kind or None
    if shape.is_placeholder:
        return MSO_SHAPE_TYPE.PLACEHOLDER
    if kind == "connector":
        return MSO_SHAPE_TYPE.LINE
    if shape.has_chart:
        return MSO_SHAPE_TYPE.CHART
    if shape.has_table:
        return MSO_SHAPE_TYPE.TABLE
    if kind == "ole_object":
        return MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT
    if kind == "movie":
        return MSO_SHAPE_TYPE.MEDIA
    if kind == "picture":
        return MSO_SHAPE_TYPE.PICTURE
    if kind == "group":
        return MSO_SHAPE_TYPE.GROUP
    if kind == "freeform" or _shape_is_freeform(shape):
        return MSO_SHAPE_TYPE.FREEFORM
    if kind == "shape":
        if _shape_is_text_box(shape):
            return MSO_SHAPE_TYPE.TEXT_BOX
        return MSO_SHAPE_TYPE.AUTO_SHAPE
    return None


def _shape_auto_shape_type_value(shape: Any) -> Any:
    preset_geometry = _shape_auto_shape_preset_geometry(shape)
    try:
        from pptx.enum.shapes import MSO_SHAPE
    except ImportError:
        return preset_geometry
    return MSO_SHAPE(_AUTO_SHAPE_TYPE_BY_PRESET[preset_geometry])


def _shape_adjustment_guides(shape: Any) -> list[tuple[str, int]]:
    queued = shape._payload.get("adjustment_guides")
    if isinstance(queued, list):
        return [
            (str(name), int(value))
            for name, value in queued
            if isinstance(name, str)
        ]

    if _shape_is_text_box(shape):
        shape._payload["adjustment_guides"] = []
        return []
    try:
        preset_geometry = _shape_auto_shape_preset_geometry(shape)
    except ValueError:
        guides = list(_shape_adjustment_actuals(shape).items())
        shape._payload["adjustment_guides"] = list(guides)
        return guides
    guides = list(_AUTO_SHAPE_ADJUSTMENT_DEFAULTS.get(preset_geometry, ()))
    if shape._payload.get("id") is None:
        shape._payload["adjustment_guides"] = list(guides)
        return guides
    actuals = _shape_adjustment_actuals(shape)
    guides = [(name, actuals.get(name, default)) for name, default in guides]
    shape._payload["adjustment_guides"] = list(guides)
    return guides


def _shape_adjustment_actuals(shape: Any) -> dict[str, int]:
    try:
        with zipfile.ZipFile(shape._slide._presentation.path) as package:
            root = ET.fromstring(package.read(shape._slide.partname))
        shape_element = _shape_xml_element(root, shape._payload)
    except (
        AttributeError,
        FileNotFoundError,
        IndexError,
        KeyError,
        ET.ParseError,
        zipfile.BadZipFile,
    ):
        return {}
    actuals: dict[str, int] = {}
    for guide in shape_element.findall(
        f"./{{{P_NS}}}spPr/{{{A_NS}}}prstGeom/{{{A_NS}}}avLst/{{{A_NS}}}gd"
    ):
        name = guide.attrib.get("name")
        formula = guide.attrib.get("fmla", "")
        if not name or not formula.startswith("val "):
            continue
        try:
            actuals[name] = int(formula[4:])
        except ValueError:
            continue
    return actuals


def _shape_line_xml_element(shape: Any) -> ET.Element | None:
    try:
        with zipfile.ZipFile(shape._slide._presentation.path) as package:
            root = ET.fromstring(package.read(shape._slide.partname))
        shape_element = _shape_xml_element(root, shape._payload)
    except (
        AttributeError,
        FileNotFoundError,
        IndexError,
        KeyError,
        ET.ParseError,
        zipfile.BadZipFile,
    ):
        shape_element = None

    if shape_element is not None:
        shape_properties = shape_element.find(f"{{{P_NS}}}spPr")
        if shape_properties is not None:
            line = shape_properties.find(f"{{{A_NS}}}ln")
            if line is not None:
                shape._payload["has_line_element"] = True
                return line
    if shape._payload.get("has_line_element"):
        return ET.Element(f"{{{A_NS}}}ln")
    return None


def _shape_auto_shape_preset_geometry(shape: Any) -> str:
    if shape._payload.get("kind") == "picture":
        return "rect"
    if shape._payload.get("kind") != "shape":
        raise ValueError("shape is not an auto shape")

    preset = shape._payload.get("preset_geometry")
    if isinstance(preset, str) and preset in _AUTO_SHAPE_TYPE_BY_PRESET:
        return preset
    if _shape_is_text_box(shape):
        raise ValueError("shape is not an auto shape")

    try:
        with zipfile.ZipFile(shape._slide._presentation.path) as package:
            root = ET.fromstring(package.read(shape._slide.partname))
        shape_element = _shape_xml_element(root, shape._payload)
    except (
        AttributeError,
        FileNotFoundError,
        IndexError,
        KeyError,
        ET.ParseError,
        zipfile.BadZipFile,
    ) as exc:
        raise ValueError("shape is not an auto shape") from exc

    non_visual_shape_properties = shape_element.find(f".//{{{P_NS}}}cNvSpPr")
    if (
        non_visual_shape_properties is not None
        and non_visual_shape_properties.attrib.get("txBox") == "1"
    ):
        raise ValueError("shape is not an auto shape")

    preset_geometry = shape_element.find(f"./{{{P_NS}}}spPr/{{{A_NS}}}prstGeom")
    if preset_geometry is None:
        raise ValueError("shape is not an auto shape")
    preset_value = preset_geometry.attrib.get("prst")
    if preset_value not in _AUTO_SHAPE_TYPE_BY_PRESET:
        raise ValueError("shape is not an auto shape")
    shape._payload["preset_geometry"] = preset_value
    return preset_value


def _shape_is_text_box(shape: Any) -> bool:
    raw = shape._payload.get("is_text_box")
    if isinstance(raw, bool):
        return raw
    if shape._payload.get("kind") != "shape":
        return False
    try:
        with zipfile.ZipFile(shape._slide._presentation.path) as package:
            root = ET.fromstring(package.read(shape._slide.partname))
        shape_element = _shape_xml_element(root, shape._payload)
    except (
        AttributeError,
        FileNotFoundError,
        IndexError,
        KeyError,
        ET.ParseError,
        zipfile.BadZipFile,
    ):
        return False

    non_visual_shape_properties = shape_element.find(f".//{{{P_NS}}}cNvSpPr")
    is_text_box = (
        non_visual_shape_properties is not None
        and non_visual_shape_properties.attrib.get("txBox") == "1"
    )
    shape._payload["is_text_box"] = is_text_box
    return is_text_box


def _shape_is_freeform(shape: Any) -> bool:
    raw = shape._payload.get("is_freeform")
    if isinstance(raw, bool):
        return raw
    if shape._payload.get("kind") == "freeform":
        shape._payload["is_freeform"] = True
        return True
    if shape._payload.get("kind") != "shape":
        return False
    try:
        with zipfile.ZipFile(shape._slide._presentation.path) as package:
            root = ET.fromstring(package.read(shape._slide.partname))
        shape_element = _shape_xml_element(root, shape._payload)
    except (
        AttributeError,
        FileNotFoundError,
        IndexError,
        KeyError,
        ET.ParseError,
        zipfile.BadZipFile,
    ):
        return False
    shape_properties = shape_element.find(f"./{{{P_NS}}}spPr")
    is_freeform = (
        shape_properties is not None
        and shape_properties.find(f"{{{A_NS}}}custGeom") is not None
    )
    shape._payload["is_freeform"] = is_freeform
    return is_freeform


def _connector_preset_geometry(connector_type: Any) -> str:
    xml_value = getattr(connector_type, "xml_value", None)
    if isinstance(xml_value, str) and xml_value:
        if xml_value in {"line", "bentConnector3", "curvedConnector3"}:
            return xml_value

    try:
        connector_id = int(connector_type)
    except (TypeError, ValueError):
        connector_id = None
    if connector_id in _CONNECTOR_PRESET_BY_ID:
        return _CONNECTOR_PRESET_BY_ID[connector_id]

    connector_name = getattr(connector_type, "name", None)
    if connector_name is None:
        connector_name = str(connector_type)
    connector_name = connector_name.rsplit(".", 1)[-1].strip().upper()
    if connector_name in _CONNECTOR_PRESET_BY_NAME:
        return _CONNECTOR_PRESET_BY_NAME[connector_name]
    if connector_name in {"LINE", "BENTCONNECTOR3", "CURVEDCONNECTOR3"}:
        return {
            "LINE": "line",
            "BENTCONNECTOR3": "bentConnector3",
            "CURVEDCONNECTOR3": "curvedConnector3",
        }[connector_name]

    raise NotImplementedError(
        "add_connector currently supports MSO_CONNECTOR.STRAIGHT, ELBOW, and CURVE"
    )


def _auto_shape_preset_geometry(auto_shape_type: Any) -> str:
    try:
        shape_id = int(auto_shape_type)
    except (TypeError, ValueError):
        shape_id = None
    if shape_id in _AUTO_SHAPE_PRESET_BY_ID:
        return _AUTO_SHAPE_PRESET_BY_ID[shape_id]

    shape_name = getattr(auto_shape_type, "name", None)
    if shape_name is None:
        shape_name = str(auto_shape_type)
    shape_name = shape_name.rsplit(".", 1)[-1].strip().upper()
    if shape_name in _AUTO_SHAPE_PRESET_BY_NAME:
        return _AUTO_SHAPE_PRESET_BY_NAME[shape_name]

    raise NotImplementedError(
        "add_shape currently supports common MSO_SHAPE presets including "
        "RECTANGLE, OVAL, ROUNDED_RECTANGLE, ISOSCELES_TRIANGLE, RIGHT_TRIANGLE, "
        "DIAMOND, PARALLELOGRAM, TRAPEZOID, HEXAGON, flowchart shapes, "
        "common arrows, callouts, stars, math symbols, braces/brackets, "
        "diagram symbols, and other frequent Office preset geometries"
    )
