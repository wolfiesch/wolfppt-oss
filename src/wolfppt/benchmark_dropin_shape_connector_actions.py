"""Connector-focused shape drop-in benchmark actions."""

from __future__ import annotations

from typing import Any

from .benchmark_cases import (
    CONNECTOR_EXPECTED_TRANSFORM,
    CONNECTOR_LINE_STYLE_EXPECTED,
)
from .benchmark_dropin_common_actions import _shape_at, _shape_collection_at
from .presentation import Presentation as WolfPresentation

TOP_LEVEL_CONNECTOR_FIXTURES = (
    "text_basic/title_body_bullets",
    "workloads/mixed_real_world_deck",
    "workloads/mixed_deal_review_deck",
)
GROUP_CONNECTOR_TARGETS = {
    "shapes/grouped_shapes": (0, 0),
    "workloads/management_reporting_deck": (2, 1),
    "workloads/customer_success_review_pack": (4, 1),
}


def _dropin_group_shape(fixture_id: str, prs: Any, operation: str) -> Any:
    if fixture_id not in GROUP_CONNECTOR_TARGETS:
        raise RuntimeError(
            f"drop-in {operation} benchmark does not support fixture {fixture_id}"
        )
    slide_index, group_index = GROUP_CONNECTOR_TARGETS[fixture_id]
    if isinstance(prs, WolfPresentation):
        return _shape_at(prs, slide_index, group_index)
    return prs.slides[slide_index].shapes[group_index]


def _apply_python_pptx_dropin_add_deeper_nested_group_connector(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in GROUP_CONNECTOR_TARGETS:
        raise RuntimeError(
            "drop-in add-deeper-nested-group-connector benchmark does not support "
            f"fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_CONNECTOR

    slide_index, group_index = GROUP_CONNECTOR_TARGETS[fixture_id]
    (
        prs.slides[slide_index]
        .shapes[group_index]
        .shapes.add_group_shape()
        .shapes.add_group_shape()
        .shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            CONNECTOR_EXPECTED_TRANSFORM["x"],
            CONNECTOR_EXPECTED_TRANSFORM["y"],
            CONNECTOR_EXPECTED_TRANSFORM["x"] + CONNECTOR_EXPECTED_TRANSFORM["cx"],
            CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
        )
    )


def _apply_wolfppt_dropin_add_deeper_nested_group_connector(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.enum.shapes import MSO_CONNECTOR

    (
        _dropin_group_shape(fixture_id, prs, "add-deeper-nested-group-connector")
        .shapes.add_group_shape()
        .shapes.add_group_shape()
        .shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            CONNECTOR_EXPECTED_TRANSFORM["x"],
            CONNECTOR_EXPECTED_TRANSFORM["y"],
            CONNECTOR_EXPECTED_TRANSFORM["x"] + CONNECTOR_EXPECTED_TRANSFORM["cx"],
            CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
        )
    )


def _apply_python_pptx_dropin_add_group_connector(fixture_id: str, prs: Any) -> None:
    if fixture_id not in GROUP_CONNECTOR_TARGETS:
        raise RuntimeError(
            f"drop-in add-group-connector benchmark does not support fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_CONNECTOR

    slide_index, group_index = GROUP_CONNECTOR_TARGETS[fixture_id]
    prs.slides[slide_index].shapes[group_index].shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        CONNECTOR_EXPECTED_TRANSFORM["x"],
        CONNECTOR_EXPECTED_TRANSFORM["y"],
        CONNECTOR_EXPECTED_TRANSFORM["x"] + CONNECTOR_EXPECTED_TRANSFORM["cx"],
        CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
    )


def _apply_wolfppt_dropin_add_group_connector(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.enum.shapes import MSO_CONNECTOR

    _dropin_group_shape(fixture_id, prs, "add-group-connector").shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        CONNECTOR_EXPECTED_TRANSFORM["x"],
        CONNECTOR_EXPECTED_TRANSFORM["y"],
        CONNECTOR_EXPECTED_TRANSFORM["x"] + CONNECTOR_EXPECTED_TRANSFORM["cx"],
        CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
    )


def _apply_python_pptx_dropin_add_nested_group_connector(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in GROUP_CONNECTOR_TARGETS:
        raise RuntimeError(
            f"drop-in add-nested-group-connector benchmark does not support fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_CONNECTOR

    slide_index, group_index = GROUP_CONNECTOR_TARGETS[fixture_id]
    nested = prs.slides[slide_index].shapes[group_index].shapes.add_group_shape()
    nested.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        CONNECTOR_EXPECTED_TRANSFORM["x"],
        CONNECTOR_EXPECTED_TRANSFORM["y"],
        CONNECTOR_EXPECTED_TRANSFORM["x"] + CONNECTOR_EXPECTED_TRANSFORM["cx"],
        CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
    )


def _apply_wolfppt_dropin_add_nested_group_connector(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.enum.shapes import MSO_CONNECTOR

    nested = _dropin_group_shape(
        fixture_id, prs, "add-nested-group-connector"
    ).shapes.add_group_shape()
    nested.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        CONNECTOR_EXPECTED_TRANSFORM["x"],
        CONNECTOR_EXPECTED_TRANSFORM["y"],
        CONNECTOR_EXPECTED_TRANSFORM["x"] + CONNECTOR_EXPECTED_TRANSFORM["cx"],
        CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
    )


def _apply_python_pptx_dropin_add_connector(fixture_id: str, prs: Any) -> None:
    if fixture_id not in TOP_LEVEL_CONNECTOR_FIXTURES:
        raise RuntimeError(
            f"drop-in add-connector benchmark does not support fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_CONNECTOR

    prs.slides[0].shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        CONNECTOR_EXPECTED_TRANSFORM["x"],
        CONNECTOR_EXPECTED_TRANSFORM["y"],
        CONNECTOR_EXPECTED_TRANSFORM["x"] + CONNECTOR_EXPECTED_TRANSFORM["cx"],
        CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
    )


def _apply_wolfppt_dropin_add_connector(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in TOP_LEVEL_CONNECTOR_FIXTURES:
        raise RuntimeError(
            f"drop-in add-connector benchmark does not support fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_CONNECTOR

    _shape_collection_at(prs, 0).add_connector(
        MSO_CONNECTOR.STRAIGHT,
        CONNECTOR_EXPECTED_TRANSFORM["x"],
        CONNECTOR_EXPECTED_TRANSFORM["y"],
        CONNECTOR_EXPECTED_TRANSFORM["x"] + CONNECTOR_EXPECTED_TRANSFORM["cx"],
        CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
    )


def _apply_python_pptx_dropin_connector_line_style_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in TOP_LEVEL_CONNECTOR_FIXTURES:
        raise RuntimeError(
            f"drop-in connector line-style benchmark does not support fixture {fixture_id}"
        )
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE

    connector = prs.slides[0].shapes[-1]
    connector.line.color.rgb = RGBColor(0x65, 0x43, 0x21)
    connector.line.width = CONNECTOR_LINE_STYLE_EXPECTED["line_width"]
    connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def _apply_python_pptx_dropin_connector_connection_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in TOP_LEVEL_CONNECTOR_FIXTURES:
        raise RuntimeError(
            f"drop-in connector connection benchmark does not support fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE

    slide = prs.slides[0]
    begin_target = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        CONNECTOR_EXPECTED_TRANSFORM["x"],
        CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
        1371600,
        685800,
    )
    end_target = slide.shapes.add_shape(
        MSO_SHAPE.DIAMOND,
        CONNECTOR_EXPECTED_TRANSFORM["x"] + CONNECTOR_EXPECTED_TRANSFORM["cx"],
        CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
        1143000,
        685800,
    )
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        CONNECTOR_EXPECTED_TRANSFORM["x"],
        CONNECTOR_EXPECTED_TRANSFORM["y"],
        CONNECTOR_EXPECTED_TRANSFORM["x"] + CONNECTOR_EXPECTED_TRANSFORM["cx"],
        CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
    )
    connector.begin_connect(begin_target, 3)
    connector.end_connect(end_target, 1)


def _apply_python_pptx_dropin_group_connector_connection_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in GROUP_CONNECTOR_TARGETS:
        raise RuntimeError(
            "drop-in grouped connector connection benchmark does not "
            f"support fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE

    slide_index, group_index = GROUP_CONNECTOR_TARGETS[fixture_id]
    group = prs.slides[slide_index].shapes[group_index]
    begin_target = group.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        914400,
        1828800,
        1371600,
        685800,
    )
    end_target = group.shapes.add_shape(
        MSO_SHAPE.DIAMOND,
        3657600,
        1828800,
        1143000,
        685800,
    )
    connector = group.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        914400,
        914400,
        3657600,
        1828800,
    )
    connector.begin_connect(begin_target, 3)
    connector.end_connect(end_target, 1)


def _apply_python_pptx_dropin_existing_connector_connection_edit(
    fixture_id: str,
    prs: Any,
) -> None:
    if fixture_id not in TOP_LEVEL_CONNECTOR_FIXTURES:
        raise RuntimeError(
            "drop-in existing-shape connector connection benchmark does not "
            f"support fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_CONNECTOR

    slide = prs.slides[0]
    begin_target = slide.shapes[0]
    end_target = slide.shapes[1]
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        CONNECTOR_EXPECTED_TRANSFORM["x"],
        CONNECTOR_EXPECTED_TRANSFORM["y"],
        CONNECTOR_EXPECTED_TRANSFORM["x"] + CONNECTOR_EXPECTED_TRANSFORM["cx"],
        CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
    )
    connector.begin_connect(begin_target, 2)
    connector.end_connect(end_target, 0)


def _apply_wolfppt_dropin_connector_line_style_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in TOP_LEVEL_CONNECTOR_FIXTURES:
        raise RuntimeError(
            f"drop-in connector line-style benchmark does not support fixture {fixture_id}"
        )
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE

    connector = _shape_collection_at(prs, 0)[-1]
    connector.line.color.rgb = RGBColor(0x65, 0x43, 0x21)
    connector.line.width = CONNECTOR_LINE_STYLE_EXPECTED["line_width"]
    connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def _apply_wolfppt_dropin_connector_connection_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in TOP_LEVEL_CONNECTOR_FIXTURES:
        raise RuntimeError(
            f"drop-in connector connection benchmark does not support fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE

    shapes = _shape_collection_at(prs, 0)
    begin_target = shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        CONNECTOR_EXPECTED_TRANSFORM["x"],
        CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
        1371600,
        685800,
    )
    end_target = shapes.add_shape(
        MSO_SHAPE.DIAMOND,
        CONNECTOR_EXPECTED_TRANSFORM["x"] + CONNECTOR_EXPECTED_TRANSFORM["cx"],
        CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
        1143000,
        685800,
    )
    connector = shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        CONNECTOR_EXPECTED_TRANSFORM["x"],
        CONNECTOR_EXPECTED_TRANSFORM["y"],
        CONNECTOR_EXPECTED_TRANSFORM["x"] + CONNECTOR_EXPECTED_TRANSFORM["cx"],
        CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
    )
    connector.begin_connect(begin_target, 3)
    connector.end_connect(end_target, 1)


def _apply_wolfppt_dropin_group_connector_connection_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE

    group = _dropin_group_shape(fixture_id, prs, "grouped connector connection")
    begin_target = group.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        914400,
        1828800,
        1371600,
        685800,
    )
    end_target = group.shapes.add_shape(
        MSO_SHAPE.DIAMOND,
        3657600,
        1828800,
        1143000,
        685800,
    )
    connector = group.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        914400,
        914400,
        3657600,
        1828800,
    )
    connector.begin_connect(begin_target, 3)
    connector.end_connect(end_target, 1)


def _apply_wolfppt_dropin_existing_connector_connection_edit(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    if fixture_id not in TOP_LEVEL_CONNECTOR_FIXTURES:
        raise RuntimeError(
            "drop-in existing-shape connector connection benchmark does not "
            f"support fixture {fixture_id}"
        )
    from pptx.enum.shapes import MSO_CONNECTOR

    shapes = _shape_collection_at(prs, 0)
    begin_target = shapes[0]
    end_target = shapes[1]
    connector = shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        CONNECTOR_EXPECTED_TRANSFORM["x"],
        CONNECTOR_EXPECTED_TRANSFORM["y"],
        CONNECTOR_EXPECTED_TRANSFORM["x"] + CONNECTOR_EXPECTED_TRANSFORM["cx"],
        CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
    )
    connector.begin_connect(begin_target, 2)
    connector.end_connect(end_target, 0)
