"""Customer-success review workload fixture builder."""

from __future__ import annotations

from io import BytesIO


def add_customer_success_review_pack(prs, slide) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches

    _add_textbox(
        slide,
        "Customer Success Quarterly Review",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    _add_textbox(
        slide,
        "Renewal forecast, adoption trends, account risks, and executive actions.",
        Inches(0.7),
        Inches(1.05),
        Inches(8.2),
        Inches(0.8),
    )
    slide.shapes.add_picture(
        BytesIO(_png_1x1()),
        Inches(8.85),
        Inches(0.35),
        Inches(0.4),
        Inches(0.4),
    )
    slide.notes_slide.notes_text_frame.text = (
        "Synthetic customer success review pack for real-corpus preservation evidence."
    )

    renewals_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        renewals_slide,
        "Renewal Forecast and At-Risk ARR",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    renewals = renewals_slide.shapes.add_table(
        6,
        5,
        Inches(0.45),
        Inches(1.0),
        Inches(9.1),
        Inches(3.0),
    ).table
    for row_idx, row in enumerate(
        [
            ["Segment", "Renewing ARR", "At risk", "Champion", "Next step"],
            ["Enterprise", "$18.6m", "$1.4m", "VP Ops", "Exec sponsor call"],
            ["Mid-market", "$9.8m", "$0.9m", "Director CS", "Usage review"],
            ["Strategic", "$12.2m", "$2.1m", "CRO", "Pricing exception"],
            ["SMB", "$4.6m", "$0.4m", "CSM pool", "Automation campaign"],
            ["Total", "$45.2m", "$4.8m", "CS Leadership", "Weekly risk standup"],
        ]
    ):
        for col_idx, value in enumerate(row):
            renewals.cell(row_idx, col_idx).text = value
    renewals_slide.notes_slide.notes_text_frame.text = (
        "Renewal table should preserve exact ARR, owner, and next-step text."
    )

    adoption_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        adoption_slide,
        "Adoption and Health Score Trends",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    chart_data = CategoryChartData()
    chart_data.categories = ["Jan", "Feb", "Mar", "Apr"]
    chart_data.add_series("Active seats", (8200, 8750, 9100, 9640))
    chart_data.add_series("Healthy accounts", (71, 74, 77, 81))
    adoption_slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.7),
        Inches(1.0),
        Inches(8.4),
        Inches(4.0),
        chart_data,
    )
    adoption_slide.notes_slide.notes_text_frame.text = (
        "Line chart workbook and customer-health notes must survive round-trip."
    )

    risks_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        risks_slide,
        "Account Risk Register",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    risk_table = risks_slide.shapes.add_table(
        5,
        4,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(2.45),
    ).table
    for row_idx, row in enumerate(
        [
            ["Account", "Risk", "ARR", "Mitigation"],
            ["Northwind", "Sponsor left", "$1.2m", "Map new buyer"],
            ["Fabrikam", "Low usage", "$0.8m", "Admin enablement"],
            ["Adventure", "Pricing pushback", "$1.6m", "Bundle proposal"],
            ["Contoso", "Implementation delay", "$1.1m", "Exec escalation"],
        ]
    ):
        for col_idx, value in enumerate(row):
            risk_table.cell(row_idx, col_idx).text = value
    risks_slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(1.0),
        Inches(3.75),
        Inches(8.4),
        Inches(3.75),
    )
    risks_slide.notes_slide.notes_text_frame.text = (
        "Risk-register table and connector line should be preserved."
    )

    actions_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        actions_slide,
        "Executive Actions",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    group = actions_slide.shapes.add_group_shape()
    for idx, (label, owner) in enumerate(
        (
            ("Renewal war room", "CS Ops"),
            ("Usage recovery", "Product"),
            ("Champion mapping", "Sales"),
            ("QBR content refresh", "Marketing"),
        )
    ):
        tile = group.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.25 + (idx % 2) * 2.85),
            Inches(0.25 + (idx // 2) * 1.05),
            Inches(2.5),
            Inches(0.78),
        )
        tile.text = f"{label}\n{owner}"
    actions_slide.notes_slide.notes_text_frame.text = (
        "Grouped customer-success action tiles should keep child text and transforms."
    )

    appendix_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        appendix_slide,
        "Appendix: CS Evidence Index",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    appendix = appendix_slide.shapes.add_table(
        5,
        3,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(2.6),
    ).table
    for row_idx, row in enumerate(
        [
            ["Artifact", "Source", "Owner"],
            ["Renewal forecast", "renewal_forecast.xlsx", "CS Ops"],
            ["Usage export", "product_usage.csv", "Data"],
            ["Risk register", "account_risks.xlsx", "CS Leadership"],
            ["QBR plan", "qbr_playbook.docx", "Marketing"],
        ]
    ):
        for col_idx, value in enumerate(row):
            appendix.cell(row_idx, col_idx).text = value


def _add_textbox(slide, text: str, left, top, width, height):
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.text = text
    return shape


def _png_1x1() -> bytes:
    return (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
        b"\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde\x00"
        b"\x00\x00\x0cIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02"
        b"\xfeA\xde\xfc\x8d\x00\x00\x00\x00IEND\xaeB`\x82"
    )
