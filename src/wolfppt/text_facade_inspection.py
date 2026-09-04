"""XML and payload readers for the text facade."""

# ruff: noqa: F401

from __future__ import annotations

import zipfile
from typing import TYPE_CHECKING, Any
from xml.etree import ElementTree as ET

from .facade_values import (
    centipoints_value as _centipoints_value,
    coerce_emu as _coerce_emu,
    emu_to_centipoints as _emu_to_centipoints,
    rotation_degrees_from_units as _rotation_degrees_from_units,
    text_frame_auto_size_value as _text_frame_auto_size_value,
    text_frame_vertical_anchor_value as _text_frame_vertical_anchor_value,
)
from .shape_payloads import (
    _shape_paragraph_font_value,
    _shape_paragraph_run_bold,
    _shape_paragraph_run_font_name,
    _shape_paragraph_run_font_size,
    _shape_paragraph_run_italic,
    _shape_paragraph_run_matrix,
    _shape_paragraph_spacing_values,
    _shape_transform,
)
from .shape_xml import (
    _find_shape_transform_xml_element,
    _find_text_body_properties_element,
    _shape_xml_element,
)
from .text_facade_queue import (
    _is_group_child_text_shape,
    _is_part_scoped_text_shape,
    _queue_paragraph_spacing_state,
)
from .text_facade_xml_inspection import (
    _shape_run_count_payload,
    _shape_text_paragraph_alignment,
    _shape_text_paragraph_default_run_properties,
    _shape_text_paragraph_font_bool,
    _shape_text_paragraph_font_fill_type,
    _shape_text_paragraph_font_gradient_payload,
    _shape_text_paragraph_font_language,
    _shape_text_paragraph_font_name,
    _shape_text_paragraph_font_pattern,
    _shape_text_paragraph_font_pattern_rgb,
    _shape_text_paragraph_font_rgb,
    _shape_text_paragraph_font_size,
    _shape_text_paragraph_font_theme_color,
    _shape_text_paragraph_font_underline,
    _shape_text_paragraph_level,
    _shape_text_paragraph_spacing,
    _shape_text_run_fill_type,
    _shape_text_run_gradient_payload,
    _shape_text_run_hyperlink_address,
    _shape_text_run_language_id,
    _shape_text_run_pattern,
    _shape_text_run_pattern_rgb,
    _shape_text_run_rgb,
    _shape_text_run_theme_color,
)
from .xml_helpers import xml_local_name as _xml_local_name

if TYPE_CHECKING:
    from .shape_core_facade import Shape

_TEXT_FRAME_MARGIN_ATTRS = {
    "margin_left": "lIns",
    "margin_right": "rIns",
    "margin_top": "tIns",
    "margin_bottom": "bIns",
}
_TEXT_FRAME_MARGIN_DEFAULTS = {
    "lIns": 91440,
    "rIns": 91440,
    "tIns": 45720,
    "bIns": 45720,
}
_TEXT_FRAME_AUTO_SIZE_TAGS = {"noAutofit", "spAutoFit", "normAutofit"}

def _shape_paragraph_font_bold(shape: Shape, paragraph_index: int) -> bool | None:
    return _shape_paragraph_font_value(
        shape,
        paragraph_index,
        "paragraph_font_bold",
        lambda: _shape_text_paragraph_font_bool(shape, paragraph_index, "b"),
    )


def _shape_paragraph_font_italic(shape: Shape, paragraph_index: int) -> bool | None:
    return _shape_paragraph_font_value(
        shape,
        paragraph_index,
        "paragraph_font_italic",
        lambda: _shape_text_paragraph_font_bool(shape, paragraph_index, "i"),
    )


def _shape_paragraph_font_underline(shape: Shape, paragraph_index: int) -> bool | None:
    return _shape_paragraph_font_value(
        shape,
        paragraph_index,
        "paragraph_font_underline",
        lambda: _shape_text_paragraph_font_underline(shape, paragraph_index),
    )


def _shape_paragraph_font_size(shape: Shape, paragraph_index: int) -> int | None:
    return _shape_paragraph_font_value(
        shape,
        paragraph_index,
        "paragraph_font_size",
        lambda: _shape_text_paragraph_font_size(shape, paragraph_index),
    )


def _shape_paragraph_font_name(shape: Shape, paragraph_index: int) -> str | None:
    return _shape_paragraph_font_value(
        shape,
        paragraph_index,
        "paragraph_font_name",
        lambda: _shape_text_paragraph_font_name(shape, paragraph_index),
    )


def _shape_paragraph_font_rgb(shape: Shape, paragraph_index: int) -> str | None:
    return _shape_paragraph_font_value(
        shape,
        paragraph_index,
        "paragraph_font_rgb",
        lambda: _shape_text_paragraph_font_rgb(shape, paragraph_index),
    )


def _shape_paragraph_font_fill_type(shape: Shape, paragraph_index: int) -> str | None:
    return _shape_paragraph_font_value(
        shape,
        paragraph_index,
        "paragraph_font_fill_type",
        lambda: _shape_text_paragraph_font_fill_type(shape, paragraph_index),
    )


def _shape_paragraph_font_language(shape: Shape, paragraph_index: int) -> str | None:
    return _shape_paragraph_font_value(
        shape,
        paragraph_index,
        "paragraph_font_language",
        lambda: _shape_text_paragraph_font_language(shape, paragraph_index),
    )


def _shape_rotation(shape: Shape) -> float:
    try:
        with zipfile.ZipFile(shape._slide._presentation.path) as package:
            root = ET.fromstring(package.read(shape._slide.partname))
        shape_element = _shape_xml_element(root, shape._payload)
        transform = _find_shape_transform_xml_element(shape_element)
    except (
        AttributeError,
        FileNotFoundError,
        KeyError,
        ET.ParseError,
        zipfile.BadZipFile,
    ):
        return 0.0
    if transform is None:
        return 0.0
    raw = transform.attrib.get("rot")
    if raw in (None, ""):
        return 0.0
    try:
        return _rotation_degrees_from_units(int(raw))
    except ValueError:
        return 0.0


def _text_frame_margin(shape: Shape, attr: str) -> int:
    margins = shape._payload.get("text_frame_margins")
    if not isinstance(margins, dict):
        margins = _shape_text_frame_margins(shape)
        shape._payload["text_frame_margins"] = margins
    return int(margins.get(attr, _TEXT_FRAME_MARGIN_DEFAULTS[attr]))


def _set_text_frame_margin(shape: Shape, attr: str, value: Any) -> None:
    margin = _coerce_emu(value, _text_frame_margin_name(attr))
    if margin == _text_frame_margin(shape, attr):
        return
    margins = shape._payload.get("text_frame_margins")
    if not isinstance(margins, dict):
        margins = _shape_text_frame_margins(shape)
    margins[attr] = margin
    shape._payload["text_frame_margins"] = margins
    if _is_part_scoped_text_shape(shape):
        shape._slide._presentation._queue_part_text_frame_margin(
            shape._slide.partname,
            shape._index,
            attr,
            margin,
        )
        return
    if _is_group_child_text_shape(shape):
        shape._slide._presentation._queue_group_child_text_frame_margin(
            shape._slide._index,
            int(shape._payload["_group_index"]),
            int(shape._payload["_group_child_index"]),
            attr,
            margin,
        )
        return
    shape._slide._presentation._queue_text_frame_margin(
        shape._slide._index,
        shape._index,
        attr,
        margin,
    )


def _text_frame_margin_name(attr: str) -> str:
    for name, xml_attr in _TEXT_FRAME_MARGIN_ATTRS.items():
        if xml_attr == attr:
            return name
    return attr


def _shape_text_frame_margins(shape: Shape) -> dict[str, int]:
    margins = dict(_TEXT_FRAME_MARGIN_DEFAULTS)
    try:
        with zipfile.ZipFile(shape._slide._presentation.path) as package:
            root = ET.fromstring(package.read(shape._slide.partname))
        shape_element = _shape_xml_element(root, shape._payload)
        body_properties = _find_text_body_properties_element(shape_element)
    except (
        AttributeError,
        FileNotFoundError,
        KeyError,
        ET.ParseError,
        zipfile.BadZipFile,
    ):
        return margins
    if body_properties is None:
        return margins
    for attr in _TEXT_FRAME_MARGIN_DEFAULTS:
        raw = body_properties.attrib.get(attr)
        if raw is None:
            continue
        try:
            margins[attr] = int(raw)
        except ValueError:
            continue
    return margins


def _text_frame_word_wrap(shape: Shape) -> bool | None:
    if "text_frame_word_wrap" in shape._payload:
        value = shape._payload["text_frame_word_wrap"]
        return None if value is None else bool(value)
    try:
        with zipfile.ZipFile(shape._slide._presentation.path) as package:
            root = ET.fromstring(package.read(shape._slide.partname))
        shape_element = _shape_xml_element(root, shape._payload)
        body_properties = _find_text_body_properties_element(shape_element)
    except (
        AttributeError,
        FileNotFoundError,
        KeyError,
        ET.ParseError,
        zipfile.BadZipFile,
    ):
        return None
    if body_properties is None:
        return None
    raw = body_properties.attrib.get("wrap")
    if raw == "square":
        value = True
    elif raw == "none":
        value = False
    else:
        value = None
    shape._payload["text_frame_word_wrap"] = value
    return value


def _text_frame_vertical_anchor(shape: Shape) -> Any:
    if "text_frame_vertical_anchor" in shape._payload:
        return _text_frame_vertical_anchor_value(
            shape._payload["text_frame_vertical_anchor"]
        )
    try:
        with zipfile.ZipFile(shape._slide._presentation.path) as package:
            root = ET.fromstring(package.read(shape._slide.partname))
        shape_element = _shape_xml_element(root, shape._payload)
        body_properties = _find_text_body_properties_element(shape_element)
    except (
        AttributeError,
        FileNotFoundError,
        KeyError,
        ET.ParseError,
        zipfile.BadZipFile,
    ):
        return None
    value = None if body_properties is None else body_properties.attrib.get("anchor")
    shape._payload["text_frame_vertical_anchor"] = value
    return _text_frame_vertical_anchor_value(value)


def _text_frame_auto_size(shape: Shape) -> Any:
    if "text_frame_auto_size" in shape._payload:
        return _text_frame_auto_size_value(shape._payload["text_frame_auto_size"])
    try:
        with zipfile.ZipFile(shape._slide._presentation.path) as package:
            root = ET.fromstring(package.read(shape._slide.partname))
        shape_element = _shape_xml_element(root, shape._payload)
        body_properties = _find_text_body_properties_element(shape_element)
    except (
        AttributeError,
        FileNotFoundError,
        KeyError,
        ET.ParseError,
        zipfile.BadZipFile,
    ):
        return None
    value = None
    if body_properties is not None:
        for child in body_properties:
            local_name = _xml_local_name(child.tag)
            if local_name in _TEXT_FRAME_AUTO_SIZE_TAGS:
                value = local_name
                break
    shape._payload["text_frame_auto_size"] = value
    return _text_frame_auto_size_value(value)


def _best_fit_text_font_size(
    shape: Shape,
    font_family: str,
    max_size: int,
    bold: bool,
    italic: bool,
    font_file: str | None,
) -> int:
    try:
        max_point_size = int(max_size)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"max_size must be an integer, got {max_size!r}") from exc
    if max_point_size < 1:
        raise ValueError("max_size must be greater than zero")
    try:
        if font_file is None:
            from pptx.text.fonts import FontFiles

            font_file = FontFiles.find(str(font_family), bold, italic)
        from pptx.text.layout import TextFitter

        return int(
            TextFitter.best_fit_font_size(
                shape.text,
                _text_frame_extents(shape),
                max_point_size,
                str(font_file),
            )
        )
    except Exception:
        return max_point_size


def _text_frame_extents(shape: Shape) -> tuple[int, int]:
    transform = _shape_transform(shape)
    return (
        max(
            0,
            transform["cx"]
            - _text_frame_margin(shape, "lIns")
            - _text_frame_margin(shape, "rIns"),
        ),
        max(
            0,
            transform["cy"]
            - _text_frame_margin(shape, "tIns")
            - _text_frame_margin(shape, "bIns"),
        ),
    )


def _set_text_frame_fit_payload(
    shape: Shape,
    font_family: str,
    size_emu: int,
    bold: bool,
    italic: bool,
) -> None:
    shape._payload["text_frame_word_wrap"] = True
    shape._payload["text_frame_auto_size"] = "noAutofit"
    run_bold = _shape_paragraph_run_bold(shape)
    run_italic = _shape_paragraph_run_italic(shape)
    run_size = _shape_paragraph_run_font_size(shape)
    run_name = _shape_paragraph_run_font_name(shape)
    for paragraph_index, runs in enumerate(_shape_paragraph_run_matrix(shape)):
        for run_index in range(len(runs)):
            run_bold[paragraph_index][run_index] = bold
            run_italic[paragraph_index][run_index] = italic
            run_size[paragraph_index][run_index] = size_emu
            run_name[paragraph_index][run_index] = str(font_family)


def _paragraph_spacing_value(shape: Shape, paragraph_index: int, attr: str) -> Any:
    values = _shape_paragraph_spacing_values(shape, f"paragraph_{attr}")
    value = values[paragraph_index]
    if value is None:
        value = _shape_text_paragraph_spacing(shape, paragraph_index, attr)
        values[paragraph_index] = value
    if value is None:
        return None
    if attr == "line_spacing" and isinstance(value, float):
        return value
    return _centipoints_value(_emu_to_centipoints(int(value)))


def _set_paragraph_spacing_payload(
    shape: Shape,
    paragraph_index: int,
    attr: str,
    value: int | float | None,
) -> None:
    values = _shape_paragraph_spacing_values(shape, f"paragraph_{attr}")
    current = values[paragraph_index]
    if current is None:
        current = _shape_text_paragraph_spacing(shape, paragraph_index, attr)
        values[paragraph_index] = current
    if current == value:
        return
    values[paragraph_index] = value
    _queue_paragraph_spacing_state(shape, paragraph_index, attr, value)
