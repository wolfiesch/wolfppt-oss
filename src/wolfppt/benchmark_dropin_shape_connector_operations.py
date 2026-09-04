"""Connector drop-in benchmark operation implementations."""

from __future__ import annotations

from pathlib import Path
from time import perf_counter
from typing import Any

from .benchmark_cases import CONNECTOR_EXPECTED_TRANSFORM
from .benchmark_dropin_actions import (
    _apply_python_pptx_dropin_add_connector,
    _apply_python_pptx_dropin_add_deeper_nested_group_connector,
    _apply_python_pptx_dropin_add_group_connector,
    _apply_python_pptx_dropin_add_nested_group_connector,
    _apply_python_pptx_dropin_connector_connection_edit,
    _apply_python_pptx_dropin_connector_line_style_edit,
    _apply_python_pptx_dropin_existing_connector_connection_edit,
    _apply_python_pptx_dropin_group_connector_connection_edit,
    _apply_wolfppt_dropin_add_connector,
    _apply_wolfppt_dropin_add_deeper_nested_group_connector,
    _apply_wolfppt_dropin_add_group_connector,
    _apply_wolfppt_dropin_add_nested_group_connector,
    _apply_wolfppt_dropin_connector_connection_edit,
    _apply_wolfppt_dropin_connector_line_style_edit,
    _apply_wolfppt_dropin_existing_connector_connection_edit,
    _apply_wolfppt_dropin_group_connector_connection_edit,
)
from .benchmark_dropin_shape_details import (
    _dropin_add_connector_details,
    _dropin_add_deeper_nested_group_connector_details,
    _dropin_add_group_connector_details,
    _dropin_add_nested_group_connector_details,
    _dropin_connector_connection_edit_details,
    _dropin_connector_line_style_edit_details,
    _dropin_existing_connector_connection_edit_details,
    _dropin_group_connector_connection_edit_details,
)
from .benchmark_runtime import elapsed_ms as _elapsed_ms, stamp as _stamp
from .presentation import Presentation as WolfPresentation

TOP_LEVEL_CONNECTOR_FIXTURES = (
    "text_basic/title_body_bullets",
    "workloads/mixed_real_world_deck",
    "workloads/mixed_deal_review_deck",
)


def _bench_python_pptx_dropin_add_connector(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-add-connector-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_connector(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_connector_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_connector(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-add-connector-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_connector(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_connector_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_group_connector(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-add-group-connector-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_group_connector(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_group_connector_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_group_connector(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-add-group-connector-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_group_connector(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_group_connector_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_nested_group_connector(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-add-nested-group-connector-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_nested_group_connector(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_nested_group_connector_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_nested_group_connector(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-add-nested-group-connector-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_nested_group_connector(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_nested_group_connector_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_deeper_nested_group_connector(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-add-deeper-nested-group-connector-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_deeper_nested_group_connector(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_deeper_nested_group_connector_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_deeper_nested_group_connector(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-add-deeper-nested-group-connector-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_deeper_nested_group_connector(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_deeper_nested_group_connector_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_connector_connection_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-connector-connection-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_connector_connection_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_connector_connection_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_connector_connection_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-connector-connection-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_connector_connection_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_connector_connection_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_group_connector_connection_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-group-connector-connection-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_group_connector_connection_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_group_connector_connection_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_group_connector_connection_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-group-connector-connection-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_group_connector_connection_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_group_connector_connection_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_existing_connector_connection_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-existing-connector-connection-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_existing_connector_connection_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_existing_connector_connection_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_existing_connector_connection_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-existing-connector-connection-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_existing_connector_connection_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_existing_connector_connection_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_connector_line_style_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    source = _dropin_connector_style_source_deck(fixture_id, fixture_path, tmp_path)
    out = tmp_path / (
        f"python-pptx-dropin-connector-line-style-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(source))
    _apply_python_pptx_dropin_connector_line_style_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_connector_line_style_edit_details(
        fixture_id,
        source,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_connector_line_style_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    source = _dropin_connector_style_source_deck(fixture_id, fixture_path, tmp_path)
    out = tmp_path / (
        f"wolfppt-dropin-connector-line-style-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(source)
    _apply_wolfppt_dropin_connector_line_style_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_connector_line_style_edit_details(
        fixture_id,
        source,
        out,
        validate_openxml,
    )


def _dropin_connector_style_source_deck(
    fixture_id: str,
    fixture_path: Path,
    tmp_path: Path,
) -> Path:
    if fixture_id not in TOP_LEVEL_CONNECTOR_FIXTURES:
        raise RuntimeError(
            f"drop-in connector line-style benchmark does not support fixture {fixture_id}"
        )
    from pptx import Presentation as PythonPptxPresentation
    from pptx.enum.shapes import MSO_CONNECTOR

    source = tmp_path / f"dropin-connector-source-{fixture_path.stem}-{_stamp()}.pptx"
    prs = PythonPptxPresentation(str(fixture_path))
    prs.slides[0].shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        CONNECTOR_EXPECTED_TRANSFORM["x"],
        CONNECTOR_EXPECTED_TRANSFORM["y"],
        CONNECTOR_EXPECTED_TRANSFORM["x"] + CONNECTOR_EXPECTED_TRANSFORM["cx"],
        CONNECTOR_EXPECTED_TRANSFORM["y"] + CONNECTOR_EXPECTED_TRANSFORM["cy"],
    )
    prs.save(str(source))
    return source
