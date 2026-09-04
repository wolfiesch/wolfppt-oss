"""Common drop-in benchmark operation implementations."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from time import perf_counter
from typing import Any

from .benchmark_dropin_actions import (
    _apply_python_pptx_dropin_core_properties_edit,
    _apply_python_pptx_dropin_create_notes_slide,
    _apply_python_pptx_dropin_edit,
    _apply_python_pptx_dropin_mixed_workload_edit,
    _apply_python_pptx_dropin_multi_edit,
    _apply_python_pptx_dropin_notes_background_edit,
    _apply_python_pptx_dropin_notes_hyperlink_edit,
    _apply_python_pptx_dropin_notes_paragraph_run_edit,
    _apply_python_pptx_dropin_notes_placeholder_clone,
    _apply_python_pptx_dropin_notes_text_edit,
    _apply_python_pptx_dropin_notes_text_frame_flow_edit,
    _apply_python_pptx_dropin_slide_background_edit,
    _apply_python_pptx_dropin_slide_layout_remove,
    _apply_python_pptx_dropin_slide_size_edit,
    _apply_python_pptx_dropin_slide_name_edit,
    _apply_python_pptx_dropin_template_background_edit,
    _apply_wolfppt_dropin_core_properties_edit,
    _apply_wolfppt_dropin_create_notes_slide,
    _apply_wolfppt_dropin_edit,
    _apply_wolfppt_dropin_mixed_workload_edit,
    _apply_wolfppt_dropin_multi_edit,
    _apply_wolfppt_dropin_notes_background_edit,
    _apply_wolfppt_dropin_notes_hyperlink_edit,
    _apply_wolfppt_dropin_notes_paragraph_run_edit,
    _apply_wolfppt_dropin_notes_placeholder_clone,
    _apply_wolfppt_dropin_notes_text_edit,
    _apply_wolfppt_dropin_notes_text_frame_flow_edit,
    _apply_wolfppt_dropin_slide_background_edit,
    _apply_wolfppt_dropin_slide_layout_remove,
    _apply_wolfppt_dropin_slide_size_edit,
    _apply_wolfppt_dropin_slide_name_edit,
    _apply_wolfppt_dropin_template_background_edit,
)
from .benchmark_dropin_common_details import (
    _dropin_core_properties_edit_details,
    _dropin_edit_details,
    _dropin_mixed_workload_details,
    _dropin_multi_edit_details,
    _dropin_slide_background_edit_details,
    _dropin_slide_layout_remove_details,
    _dropin_slide_size_edit_details,
    _dropin_slide_name_edit_details,
    _dropin_template_background_edit_details,
)
from .benchmark_dropin_notes_details import (
    _dropin_create_notes_slide_details,
    _dropin_notes_background_edit_details,
    _dropin_notes_hyperlink_edit_details,
    _dropin_notes_paragraph_run_edit_details,
    _dropin_notes_placeholder_clone_details,
    _dropin_notes_text_edit_details,
    _dropin_notes_text_frame_flow_edit_details,
)
from .benchmark_runtime import elapsed_ms as _elapsed_ms, stamp as _stamp
from .presentation import Presentation as WolfPresentation


def _bench_python_pptx_dropin_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = (
        tmp_path
        / f"python-pptx-dropin-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_edit_details(
        fixture_id, fixture_path, out, validate_openxml
    )


def _bench_wolfppt_facade_dropin_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = (
        tmp_path / f"wolfppt-dropin-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_edit_details(
        fixture_id, fixture_path, out, validate_openxml
    )


def _bench_python_pptx_dropin_file_like_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-file-like-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    source_stream = BytesIO(fixture_path.read_bytes())
    source_stream.seek(source_stream.getbuffer().nbytes)
    output_stream = BytesIO()
    start = perf_counter()
    prs = PythonPptxPresentation(source_stream)
    _apply_python_pptx_dropin_edit(fixture_id, prs)
    prs.save(output_stream)
    out.write_bytes(output_stream.getvalue())
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_edit_details(
        fixture_id, fixture_path, out, validate_openxml
    )


def _bench_wolfppt_facade_dropin_file_like_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-file-like-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    source_stream = BytesIO(fixture_path.read_bytes())
    source_stream.seek(source_stream.getbuffer().nbytes)
    output_stream = BytesIO()
    start = perf_counter()
    prs = WolfPresentation(source_stream)
    _apply_wolfppt_dropin_edit(fixture_id, prs)
    prs.save(output_stream)
    out.write_bytes(output_stream.getvalue())
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_edit_details(
        fixture_id, fixture_path, out, validate_openxml
    )


def _bench_python_pptx_dropin_multi_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = (
        tmp_path
        / f"python-pptx-dropin-multi-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_multi_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_multi_edit_details(
        fixture_id, fixture_path, out, validate_openxml
    )


def _bench_wolfppt_facade_dropin_multi_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = (
        tmp_path
        / f"wolfppt-dropin-multi-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_multi_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_multi_edit_details(
        fixture_id, fixture_path, out, validate_openxml
    )


def _bench_python_pptx_dropin_mixed_workload_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-mixed-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_mixed_workload_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_mixed_workload_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_mixed_workload_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-mixed-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_mixed_workload_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_mixed_workload_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_notes_text_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-notes-text-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_notes_text_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_notes_text_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_notes_text_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-notes-text-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_notes_text_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_notes_text_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_notes_text_frame_flow_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-notes-text-frame-flow-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_notes_text_frame_flow_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_notes_text_frame_flow_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_notes_text_frame_flow_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-notes-text-frame-flow-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_notes_text_frame_flow_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_notes_text_frame_flow_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_notes_background_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-notes-background-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_notes_background_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_notes_background_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_notes_background_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-notes-background-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_notes_background_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_notes_background_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_notes_hyperlink_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-notes-hyperlink-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_notes_hyperlink_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_notes_hyperlink_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_notes_hyperlink_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-notes-hyperlink-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_notes_hyperlink_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_notes_hyperlink_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_create_notes_slide(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-create-notes-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_create_notes_slide(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_create_notes_slide_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_create_notes_slide(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-create-notes-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_create_notes_slide(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_create_notes_slide_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_notes_paragraph_run_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        "python-pptx-dropin-notes-paragraph-run-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_notes_paragraph_run_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_notes_paragraph_run_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_notes_paragraph_run_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-notes-paragraph-run-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_notes_paragraph_run_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_notes_paragraph_run_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_notes_placeholder_clone(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        "python-pptx-dropin-notes-placeholder-clone-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_notes_placeholder_clone(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_notes_placeholder_clone_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_notes_placeholder_clone(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-notes-placeholder-clone-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_notes_placeholder_clone(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_notes_placeholder_clone_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_slide_name_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-slide-name-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_slide_name_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_slide_name_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_slide_name_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-slide-name-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_slide_name_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_slide_name_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_slide_size_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-slide-size-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_slide_size_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_slide_size_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_slide_size_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-slide-size-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_slide_size_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_slide_size_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_slide_layout_remove(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-slide-layout-remove-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_slide_layout_remove(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_slide_layout_remove_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_slide_layout_remove(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-slide-layout-remove-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_slide_layout_remove(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_slide_layout_remove_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_slide_background_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-slide-background-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_slide_background_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_slide_background_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_slide_background_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-slide-background-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_slide_background_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_slide_background_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_template_background_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-template-background-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_template_background_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_template_background_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_template_background_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-template-background-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_template_background_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_template_background_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_core_properties_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-core-properties-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_core_properties_edit(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_core_properties_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_core_properties_edit(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-core-properties-{fixture_path.stem}-"
        f"{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_core_properties_edit(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_core_properties_edit_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )
