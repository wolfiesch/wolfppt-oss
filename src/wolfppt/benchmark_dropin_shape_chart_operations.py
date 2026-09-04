"""Chart shape drop-in benchmark operation implementations."""

from __future__ import annotations

from pathlib import Path
from time import perf_counter
from typing import Any

from .benchmark_dropin_actions import (
    _apply_python_pptx_dropin_add_bar_chart,
    _apply_python_pptx_dropin_add_bubble_3d_chart,
    _apply_python_pptx_dropin_add_bubble_chart,
    _apply_python_pptx_dropin_add_chart,
    _apply_python_pptx_dropin_add_chart_template_family,
    _apply_python_pptx_dropin_add_hierarchical_chart,
    _apply_python_pptx_dropin_add_line_chart,
    _apply_python_pptx_dropin_add_pie_chart,
    _apply_python_pptx_dropin_add_xy_scatter_chart,
    _apply_python_pptx_dropin_add_xy_scatter_template_family,
    _apply_wolfppt_dropin_add_bar_chart,
    _apply_wolfppt_dropin_add_bubble_3d_chart,
    _apply_wolfppt_dropin_add_bubble_chart,
    _apply_wolfppt_dropin_add_chart,
    _apply_wolfppt_dropin_add_chart_template_family,
    _apply_wolfppt_dropin_add_hierarchical_chart,
    _apply_wolfppt_dropin_add_line_chart,
    _apply_wolfppt_dropin_add_pie_chart,
    _apply_wolfppt_dropin_add_xy_scatter_chart,
    _apply_wolfppt_dropin_add_xy_scatter_template_family,
)
from .benchmark_dropin_shape_details import (
    _dropin_add_chart_details,
    _dropin_add_chart_template_family_details,
    _dropin_add_xy_scatter_template_family_details,
)
from .benchmark_cases import (
    ADD_BUBBLE_3D_CHART_EXPECTED,
    ADD_BUBBLE_CHART_EXPECTED,
    ADD_HIERARCHICAL_CHART_EXPECTED,
    ADD_XY_SCATTER_CHART_EXPECTED,
)
from .benchmark_runtime import elapsed_ms as _elapsed_ms, stamp as _stamp
from .chart_axis_ids import (
    normalize_chart_axis_ids_in_package as _normalize_chart_axis_ids,
)
from .presentation import Presentation as WolfPresentation


def _bench_python_pptx_dropin_add_chart(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_python_pptx_dropin_add_chart_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="chart",
        expected_chart_type="COLUMN_CLUSTERED (51)",
        apply=_apply_python_pptx_dropin_add_chart,
    )


def _bench_wolfppt_facade_dropin_add_chart(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_wolfppt_facade_dropin_add_chart_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="chart",
        expected_chart_type="COLUMN_CLUSTERED (51)",
        apply=_apply_wolfppt_dropin_add_chart,
    )


def _bench_python_pptx_dropin_add_hierarchical_chart(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    expected_metadata = _add_hierarchical_chart_expected_metadata()
    return _bench_python_pptx_dropin_add_chart_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="hierarchical-chart",
        expected_chart_type=expected_metadata["chart_type"],
        expected_chart_metadata=expected_metadata,
        expected_hierarchy_level_count=ADD_HIERARCHICAL_CHART_EXPECTED[
            "hierarchy_level_count"
        ],
        apply=_apply_python_pptx_dropin_add_hierarchical_chart,
    )


def _bench_wolfppt_facade_dropin_add_hierarchical_chart(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    expected_metadata = _add_hierarchical_chart_expected_metadata()
    return _bench_wolfppt_facade_dropin_add_chart_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="hierarchical-chart",
        expected_chart_type=expected_metadata["chart_type"],
        expected_chart_metadata=expected_metadata,
        expected_hierarchy_level_count=ADD_HIERARCHICAL_CHART_EXPECTED[
            "hierarchy_level_count"
        ],
        apply=_apply_wolfppt_dropin_add_hierarchical_chart,
    )


def _bench_python_pptx_dropin_add_bar_chart(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_python_pptx_dropin_add_chart_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="bar-chart",
        expected_chart_type="BAR_CLUSTERED (57)",
        apply=_apply_python_pptx_dropin_add_bar_chart,
    )


def _bench_wolfppt_facade_dropin_add_bar_chart(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_wolfppt_facade_dropin_add_chart_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="bar-chart",
        expected_chart_type="BAR_CLUSTERED (57)",
        apply=_apply_wolfppt_dropin_add_bar_chart,
    )


def _bench_python_pptx_dropin_add_line_chart(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_python_pptx_dropin_add_chart_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="line-chart",
        expected_chart_type="LINE_MARKERS (65)",
        apply=_apply_python_pptx_dropin_add_line_chart,
    )


def _bench_wolfppt_facade_dropin_add_line_chart(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_wolfppt_facade_dropin_add_chart_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="line-chart",
        expected_chart_type="LINE_MARKERS (65)",
        apply=_apply_wolfppt_dropin_add_line_chart,
    )


def _bench_python_pptx_dropin_add_pie_chart(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_python_pptx_dropin_add_chart_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="pie-chart",
        expected_chart_type="PIE (5)",
        apply=_apply_python_pptx_dropin_add_pie_chart,
    )


def _bench_wolfppt_facade_dropin_add_pie_chart(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_wolfppt_facade_dropin_add_chart_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="pie-chart",
        expected_chart_type="PIE (5)",
        apply=_apply_wolfppt_dropin_add_pie_chart,
    )


def _bench_python_pptx_dropin_add_xy_scatter_chart(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_python_pptx_dropin_add_chart_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="xy-scatter-chart",
        expected_chart_type="XY_SCATTER (-4169)",
        expected_chart_metadata=ADD_XY_SCATTER_CHART_EXPECTED,
        apply=_apply_python_pptx_dropin_add_xy_scatter_chart,
    )


def _bench_wolfppt_facade_dropin_add_xy_scatter_chart(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_wolfppt_facade_dropin_add_chart_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="xy-scatter-chart",
        expected_chart_type="XY_SCATTER (-4169)",
        expected_chart_metadata=ADD_XY_SCATTER_CHART_EXPECTED,
        apply=_apply_wolfppt_dropin_add_xy_scatter_chart,
    )


def _bench_python_pptx_dropin_add_bubble_chart(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_python_pptx_dropin_add_chart_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="bubble-chart",
        expected_chart_type="BUBBLE (15)",
        expected_chart_metadata=ADD_BUBBLE_CHART_EXPECTED,
        apply=_apply_python_pptx_dropin_add_bubble_chart,
    )


def _bench_wolfppt_facade_dropin_add_bubble_chart(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_wolfppt_facade_dropin_add_chart_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="bubble-chart",
        expected_chart_type="BUBBLE (15)",
        expected_chart_metadata=ADD_BUBBLE_CHART_EXPECTED,
        apply=_apply_wolfppt_dropin_add_bubble_chart,
    )


def _bench_python_pptx_dropin_add_bubble_3d_chart(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_python_pptx_dropin_add_chart_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="bubble-3d-chart",
        expected_chart_type="BUBBLE_THREE_D_EFFECT (87)",
        expected_chart_metadata=ADD_BUBBLE_3D_CHART_EXPECTED,
        apply=_apply_python_pptx_dropin_add_bubble_3d_chart,
    )


def _bench_wolfppt_facade_dropin_add_bubble_3d_chart(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_wolfppt_facade_dropin_add_chart_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="bubble-3d-chart",
        expected_chart_type="BUBBLE_THREE_D_EFFECT (87)",
        expected_chart_metadata=ADD_BUBBLE_3D_CHART_EXPECTED,
        apply=_apply_wolfppt_dropin_add_bubble_3d_chart,
    )


def _bench_python_pptx_dropin_add_chart_template_family(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_python_pptx_dropin_add_chart_template_family_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        apply=_apply_python_pptx_dropin_add_chart_template_family,
    )


def _bench_wolfppt_facade_dropin_add_chart_template_family(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_wolfppt_facade_dropin_add_chart_template_family_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        apply=_apply_wolfppt_dropin_add_chart_template_family,
    )


def _bench_python_pptx_dropin_add_xy_scatter_template_family(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_python_pptx_dropin_add_chart_template_family_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="xy-scatter-template-family",
        apply=_apply_python_pptx_dropin_add_xy_scatter_template_family,
        details=_dropin_add_xy_scatter_template_family_details,
    )


def _bench_wolfppt_facade_dropin_add_xy_scatter_template_family(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    return _bench_wolfppt_facade_dropin_add_chart_template_family_variant(
        fixture_id,
        fixture_path,
        expected_path,
        tmp_path,
        validate_openxml,
        variant="xy-scatter-template-family",
        apply=_apply_wolfppt_dropin_add_xy_scatter_template_family,
        details=_dropin_add_xy_scatter_template_family_details,
    )


def _bench_python_pptx_dropin_add_chart_variant(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
    *,
    variant: str,
    expected_chart_type: str,
    apply: Any,
    expected_chart_metadata: dict[str, Any] | None = None,
    expected_hierarchy_level_count: int | None = None,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-add-{variant}-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    apply(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    _normalize_chart_axis_ids(out)
    return elapsed_ms, _dropin_add_chart_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
        expected_chart_type=expected_chart_type,
        expected_chart_metadata=expected_chart_metadata,
        expected_hierarchy_level_count=expected_hierarchy_level_count,
    )


def _bench_python_pptx_dropin_add_chart_template_family_variant(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
    *,
    apply: Any,
    variant: str = "chart-template-family",
    details: Any = _dropin_add_chart_template_family_details,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    out = tmp_path / (
        f"python-pptx-dropin-add-{variant}-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    apply(fixture_id, prs)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    _normalize_chart_axis_ids(out)
    return elapsed_ms, details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_chart_variant(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
    *,
    variant: str,
    expected_chart_type: str,
    apply: Any,
    expected_chart_metadata: dict[str, Any] | None = None,
    expected_hierarchy_level_count: int | None = None,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-add-{variant}-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    apply(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_chart_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
        expected_chart_type=expected_chart_type,
        expected_chart_metadata=expected_chart_metadata,
        expected_hierarchy_level_count=expected_hierarchy_level_count,
    )


def _add_hierarchical_chart_expected_metadata() -> dict[str, Any]:
    return {
        "chart_type": ADD_HIERARCHICAL_CHART_EXPECTED["chart_type"],
        "categories": ADD_HIERARCHICAL_CHART_EXPECTED["categories"],
        "series": ADD_HIERARCHICAL_CHART_EXPECTED["series"],
    }


def _bench_wolfppt_facade_dropin_add_chart_template_family_variant(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
    *,
    apply: Any,
    variant: str = "chart-template-family",
    details: Any = _dropin_add_chart_template_family_details,
) -> tuple[float, dict[str, Any]]:
    out = tmp_path / (
        f"wolfppt-dropin-add-{variant}-"
        f"{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    apply(fixture_id, prs)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )
