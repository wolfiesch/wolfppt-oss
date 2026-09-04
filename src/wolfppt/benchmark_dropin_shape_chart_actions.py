"""Chart-add drop-in benchmark edit actions for shape collections."""

from __future__ import annotations

from typing import Any

from .benchmark_cases import (
    ADD_BUBBLE_CHART_EXPECTED,
    ADD_CHART_EXPECTED,
    ADD_CHART_EXPECTED_TRANSFORM,
    ADD_HIERARCHICAL_CHART_EXPECTED,
    ADD_XY_SCATTER_CHART_EXPECTED,
)
from .benchmark_dropin_common_actions import _shape_collection_at
from .chart_adds import _CHART_TEMPLATE_NAME_TO_ID, _XY_CHART_TEMPLATE_NAME_TO_ID
from .presentation import Presentation as WolfPresentation

ADD_CHART_TEMPLATE_FAMILY_TYPES = tuple(sorted(_CHART_TEMPLATE_NAME_TO_ID))
ADD_XY_CHART_TEMPLATE_FAMILY_TYPES = tuple(sorted(_XY_CHART_TEMPLATE_NAME_TO_ID))
TOP_LEVEL_CREATION_FIXTURES = (
    "text_basic/title_body_bullets",
    "workloads/mixed_real_world_deck",
    "workloads/mixed_deal_review_deck",
)


def _apply_python_pptx_dropin_add_chart(fixture_id: str, prs: Any) -> None:
    _apply_python_pptx_dropin_add_chart_type(
        fixture_id,
        prs,
        "COLUMN_CLUSTERED",
    )


def _apply_wolfppt_dropin_add_chart(fixture_id: str, prs: WolfPresentation) -> None:
    _apply_wolfppt_dropin_add_chart_type(
        fixture_id,
        prs,
        "COLUMN_CLUSTERED",
    )


def _apply_python_pptx_dropin_add_hierarchical_chart(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_dropin_add_hierarchical_chart(fixture_id, prs)


def _apply_wolfppt_dropin_add_hierarchical_chart(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_dropin_add_hierarchical_chart(fixture_id, prs)


def _apply_python_pptx_dropin_add_bar_chart(fixture_id: str, prs: Any) -> None:
    _apply_python_pptx_dropin_add_chart_type(fixture_id, prs, "BAR_CLUSTERED")


def _apply_wolfppt_dropin_add_bar_chart(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_wolfppt_dropin_add_chart_type(fixture_id, prs, "BAR_CLUSTERED")


def _apply_python_pptx_dropin_add_line_chart(fixture_id: str, prs: Any) -> None:
    _apply_python_pptx_dropin_add_chart_type(fixture_id, prs, "LINE_MARKERS")


def _apply_wolfppt_dropin_add_line_chart(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_wolfppt_dropin_add_chart_type(fixture_id, prs, "LINE_MARKERS")


def _apply_python_pptx_dropin_add_pie_chart(fixture_id: str, prs: Any) -> None:
    _apply_python_pptx_dropin_add_chart_type(fixture_id, prs, "PIE")


def _apply_wolfppt_dropin_add_pie_chart(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_wolfppt_dropin_add_chart_type(fixture_id, prs, "PIE")


def _apply_python_pptx_dropin_add_xy_scatter_chart(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_dropin_add_xy_scatter_chart(fixture_id, prs)


def _apply_wolfppt_dropin_add_xy_scatter_chart(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_dropin_add_xy_scatter_chart(fixture_id, prs)


def _apply_python_pptx_dropin_add_bubble_chart(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_dropin_add_bubble_chart(fixture_id, prs, "BUBBLE")


def _apply_wolfppt_dropin_add_bubble_chart(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_dropin_add_bubble_chart(fixture_id, prs, "BUBBLE")


def _apply_python_pptx_dropin_add_bubble_3d_chart(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_dropin_add_bubble_chart(fixture_id, prs, "BUBBLE_THREE_D_EFFECT")


def _apply_wolfppt_dropin_add_bubble_3d_chart(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_dropin_add_bubble_chart(fixture_id, prs, "BUBBLE_THREE_D_EFFECT")


def _apply_python_pptx_dropin_add_chart_template_family(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_python_pptx_dropin_add_chart_types(
        fixture_id,
        prs,
        ADD_CHART_TEMPLATE_FAMILY_TYPES,
    )


def _apply_wolfppt_dropin_add_chart_template_family(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_wolfppt_dropin_add_chart_types(
        fixture_id,
        prs,
        ADD_CHART_TEMPLATE_FAMILY_TYPES,
    )


def _apply_python_pptx_dropin_add_xy_scatter_template_family(
    fixture_id: str,
    prs: Any,
) -> None:
    _apply_dropin_add_xy_scatter_types(
        fixture_id,
        prs,
        ADD_XY_CHART_TEMPLATE_FAMILY_TYPES,
    )


def _apply_wolfppt_dropin_add_xy_scatter_template_family(
    fixture_id: str,
    prs: WolfPresentation,
) -> None:
    _apply_dropin_add_xy_scatter_types(
        fixture_id,
        prs,
        ADD_XY_CHART_TEMPLATE_FAMILY_TYPES,
    )


def _apply_python_pptx_dropin_add_chart_types(
    fixture_id: str,
    prs: Any,
    chart_type_names: tuple[str, ...],
) -> None:
    _apply_dropin_add_chart_types(fixture_id, prs, chart_type_names)


def _apply_wolfppt_dropin_add_chart_types(
    fixture_id: str,
    prs: WolfPresentation,
    chart_type_names: tuple[str, ...],
) -> None:
    _apply_dropin_add_chart_types(fixture_id, prs, chart_type_names)


def _apply_dropin_add_xy_scatter_chart(fixture_id: str, prs: Any) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-xy-scatter benchmark does not support fixture {fixture_id}"
        )
    from pptx.chart.data import XyChartData
    from pptx.enum.chart import XL_CHART_TYPE

    chart_data = XyChartData()
    for series in ADD_XY_SCATTER_CHART_EXPECTED["series"]:
        xy_series = chart_data.add_series(series["name"])
        for index, value in enumerate(series["values"], start=1):
            xy_series.add_data_point(float(index), value)
    _shape_collection_at(prs, 0).add_chart(
        XL_CHART_TYPE.XY_SCATTER,
        ADD_CHART_EXPECTED_TRANSFORM["x"],
        ADD_CHART_EXPECTED_TRANSFORM["y"],
        ADD_CHART_EXPECTED_TRANSFORM["cx"],
        ADD_CHART_EXPECTED_TRANSFORM["cy"],
        chart_data,
    )


def _apply_dropin_add_bubble_chart(
    fixture_id: str,
    prs: Any,
    chart_type_name: str,
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-bubble-chart benchmark does not support fixture {fixture_id}"
        )
    from pptx.chart.data import BubbleChartData
    from pptx.enum.chart import XL_CHART_TYPE

    chart_data = BubbleChartData()
    for series_index, series in enumerate(ADD_BUBBLE_CHART_EXPECTED["series"]):
        bubble_series = chart_data.add_series(series["name"])
        bubble_sizes = ADD_BUBBLE_CHART_EXPECTED["bubble_sizes"][series_index]
        for index, value in enumerate(series["values"], start=1):
            bubble_series.add_data_point(
                float(index),
                value,
                bubble_sizes[index - 1],
            )
    _shape_collection_at(prs, 0).add_chart(
        getattr(XL_CHART_TYPE, chart_type_name),
        ADD_CHART_EXPECTED_TRANSFORM["x"],
        ADD_CHART_EXPECTED_TRANSFORM["y"],
        ADD_CHART_EXPECTED_TRANSFORM["cx"],
        ADD_CHART_EXPECTED_TRANSFORM["cy"],
        chart_data,
    )


def _apply_dropin_add_hierarchical_chart(fixture_id: str, prs: Any) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-hierarchical-chart benchmark does not support fixture {fixture_id}"
        )
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    chart_data = CategoryChartData()
    fy24 = chart_data.categories.add_category("FY24")
    q1 = fy24.add_sub_category("Q1")
    q1.add_sub_category("Jan")
    q1.add_sub_category("Feb")
    q2 = fy24.add_sub_category("Q2")
    q2.add_sub_category("Mar")
    fy25 = chart_data.categories.add_category("FY25")
    q1 = fy25.add_sub_category("Q1")
    q1.add_sub_category("Jan")
    for series in ADD_HIERARCHICAL_CHART_EXPECTED["series"]:
        chart_data.add_series(series["name"], series["values"])
    _shape_collection_at(prs, 0).add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        ADD_CHART_EXPECTED_TRANSFORM["x"],
        ADD_CHART_EXPECTED_TRANSFORM["y"],
        ADD_CHART_EXPECTED_TRANSFORM["cx"],
        ADD_CHART_EXPECTED_TRANSFORM["cy"],
        chart_data,
    )


def _apply_dropin_add_chart_types(
    fixture_id: str,
    prs: Any,
    chart_type_names: tuple[str, ...],
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-chart family benchmark does not support fixture {fixture_id}"
        )
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    for index, chart_type_name in enumerate(chart_type_names):
        chart_data = CategoryChartData()
        chart_data.categories = ["Q1", "Q2", "Q3"]
        chart_data.add_series("Revenue", (10, 14, 18))
        _shape_collection_at(prs, 0).add_chart(
            getattr(XL_CHART_TYPE, chart_type_name),
            Inches(0.4 + (index % 4) * 2.1),
            Inches(0.4 + (index // 4) * 1.1),
            Inches(1.8),
            Inches(1.0),
            chart_data,
        )


def _apply_dropin_add_xy_scatter_types(
    fixture_id: str,
    prs: Any,
    chart_type_names: tuple[str, ...],
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-xy-scatter family benchmark does not support fixture {fixture_id}"
        )
    from pptx.chart.data import XyChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    for index, chart_type_name in enumerate(chart_type_names):
        chart_data = XyChartData()
        first = chart_data.add_series("Revenue")
        first.add_data_point(1.0, 10.0)
        first.add_data_point(2.0, 14.0)
        second = chart_data.add_series("Cost")
        second.add_data_point(1.0, 6.0)
        second.add_data_point(2.0, 7.0)
        _shape_collection_at(prs, 0).add_chart(
            getattr(XL_CHART_TYPE, chart_type_name),
            Inches(0.4 + (index % 3) * 2.4),
            Inches(0.4 + (index // 3) * 1.3),
            Inches(2.1),
            Inches(1.1),
            chart_data,
        )


def _apply_python_pptx_dropin_add_chart_type(
    fixture_id: str,
    prs: Any,
    chart_type_name: str,
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-chart benchmark does not support fixture {fixture_id}"
        )
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    chart_data = CategoryChartData()
    chart_data.categories = ADD_CHART_EXPECTED["categories"]
    for series in ADD_CHART_EXPECTED["series"]:
        chart_data.add_series(series["name"], series["values"])
    _shape_collection_at(prs, 0).add_chart(
        getattr(XL_CHART_TYPE, chart_type_name),
        ADD_CHART_EXPECTED_TRANSFORM["x"],
        ADD_CHART_EXPECTED_TRANSFORM["y"],
        ADD_CHART_EXPECTED_TRANSFORM["cx"],
        ADD_CHART_EXPECTED_TRANSFORM["cy"],
        chart_data,
    )


def _apply_wolfppt_dropin_add_chart_type(
    fixture_id: str,
    prs: WolfPresentation,
    chart_type_name: str,
) -> None:
    if fixture_id not in TOP_LEVEL_CREATION_FIXTURES:
        raise RuntimeError(
            f"drop-in add-chart benchmark does not support fixture {fixture_id}"
        )
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    chart_data = CategoryChartData()
    chart_data.categories = ADD_CHART_EXPECTED["categories"]
    for series in ADD_CHART_EXPECTED["series"]:
        chart_data.add_series(series["name"], series["values"])
    _shape_collection_at(prs, 0).add_chart(
        getattr(XL_CHART_TYPE, chart_type_name),
        ADD_CHART_EXPECTED_TRANSFORM["x"],
        ADD_CHART_EXPECTED_TRANSFORM["y"],
        ADD_CHART_EXPECTED_TRANSFORM["cx"],
        ADD_CHART_EXPECTED_TRANSFORM["cy"],
        chart_data,
    )
