"""Enterprise implementation-readiness workload fixture builder."""

from __future__ import annotations

from io import BytesIO


def add_enterprise_implementation_pack(prs, slide) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches

    _add_textbox(
        slide,
        "Enterprise Implementation Readiness Pack",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    _add_textbox(
        slide,
        "Data migration, integration milestones, customer risks, and owner actions.",
        Inches(0.7),
        Inches(1.05),
        Inches(8.2),
        Inches(0.8),
    )
    slide.shapes.add_picture(
        BytesIO(_png_1x1()),
        Inches(8.85),
        Inches(0.35),
        Inches(0.4),
        Inches(0.4),
    )
    slide.notes_slide.notes_text_frame.text = (
        "Synthetic enterprise implementation pack for real-corpus preservation evidence."
    )

    plan_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        plan_slide,
        "Implementation Workplan",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    plan = plan_slide.shapes.add_table(
        6,
        5,
        Inches(0.45),
        Inches(1.0),
        Inches(9.1),
        Inches(3.0),
    ).table
    for row_idx, row in enumerate(
        [
            ["Workstream", "Owner", "Status", "Due", "Dependency"],
            ["Data mapping", "Solutions", "Green", "Jun 07", "Field inventory"],
            ["SSO setup", "IT", "Yellow", "Jun 10", "IdP metadata"],
            ["ERP connector", "Engineering", "Yellow", "Jun 14", "Sandbox access"],
            ["Training plan", "Enablement", "Green", "Jun 18", "Role roster"],
            ["Go-live cutover", "PMO", "Red", "Jun 25", "Rollback plan"],
        ]
    ):
        for col_idx, value in enumerate(row):
            plan.cell(row_idx, col_idx).text = value
    plan_slide.notes_slide.notes_text_frame.text = (
        "Implementation workplan table should preserve owners, due dates, and dependency text."
    )

    migration_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        migration_slide,
        "Data Migration Burnup",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    chart_data = CategoryChartData()
    chart_data.categories = ["Week 1", "Week 2", "Week 3", "Week 4"]
    chart_data.add_series("Mapped fields", (48, 92, 138, 171))
    chart_data.add_series("Validated fields", (21, 57, 104, 149))
    migration_slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.7),
        Inches(1.0),
        Inches(8.4),
        Inches(4.0),
        chart_data,
    )
    migration_slide.notes_slide.notes_text_frame.text = (
        "Migration chart workbook and field-validation notes must survive round-trip."
    )

    risk_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        risk_slide,
        "Go-Live Risk Register",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    risks = risk_slide.shapes.add_table(
        5,
        4,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(2.45),
    ).table
    for row_idx, row in enumerate(
        [
            ["Risk", "Impact", "Owner", "Mitigation"],
            ["Late sandbox access", "High", "Engineering", "Escalate IT path"],
            ["Duplicate account IDs", "Medium", "Data", "Preload match rules"],
            ["Cutover freeze conflict", "High", "PMO", "Move change window"],
            ["Low admin attendance", "Medium", "Enablement", "Add office hours"],
        ]
    ):
        for col_idx, value in enumerate(row):
            risks.cell(row_idx, col_idx).text = value
    risk_slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(1.0),
        Inches(3.75),
        Inches(8.4),
        Inches(3.75),
    )
    risk_slide.notes_slide.notes_text_frame.text = (
        "Risk-register table and connector line should be preserved."
    )

    owners_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        owners_slide,
        "Launch Desk Owners",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    group = owners_slide.shapes.add_group_shape()
    for idx, (label, owner) in enumerate(
        (
            ("Data triage", "Solutions"),
            ("Cutover desk", "PMO"),
            ("Connector desk", "Engineering"),
            ("Training desk", "Enablement"),
        )
    ):
        tile = group.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.25 + (idx % 2) * 2.85),
            Inches(0.25 + (idx // 2) * 1.05),
            Inches(2.5),
            Inches(0.78),
        )
        tile.text = f"{label}\n{owner}"
    owners_slide.notes_slide.notes_text_frame.text = (
        "Grouped implementation-owner tiles should keep child text and transforms."
    )

    appendix_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        appendix_slide,
        "Appendix: Implementation Evidence Index",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    appendix = appendix_slide.shapes.add_table(
        5,
        3,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(2.6),
    ).table
    for row_idx, row in enumerate(
        [
            ["Artifact", "Source", "Owner"],
            ["Field inventory", "field_mapping.xlsx", "Solutions"],
            ["SSO checklist", "sso_readiness.docx", "IT"],
            ["Cutover plan", "cutover_plan.xlsx", "PMO"],
            ["Training roster", "training_roster.csv", "Enablement"],
        ]
    ):
        for col_idx, value in enumerate(row):
            appendix.cell(row_idx, col_idx).text = value


def _add_textbox(slide, text: str, left, top, width, height):
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.text = text
    return shape


def _png_1x1() -> bytes:
    return (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
        b"\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde\x00"
        b"\x00\x00\x0cIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02"
        b"\xfeA\xde\xfc\x8d\x00\x00\x00\x00IEND\xaeB`\x82"
    )
