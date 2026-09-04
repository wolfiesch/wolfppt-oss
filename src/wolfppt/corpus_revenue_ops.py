"""Revenue-operations forecast workload fixture builder."""

from __future__ import annotations

from io import BytesIO


def add_revenue_ops_forecast_pack(prs, slide) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches

    _add_textbox(
        slide,
        "Revenue Operations Forecast Pack",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    _add_textbox(
        slide,
        "Pipeline coverage, commit quality, forecast risks, and owner actions.",
        Inches(0.7),
        Inches(1.05),
        Inches(8.2),
        Inches(0.8),
    )
    slide.shapes.add_picture(
        BytesIO(_png_1x1()),
        Inches(8.85),
        Inches(0.35),
        Inches(0.4),
        Inches(0.4),
    )
    slide.notes_slide.notes_text_frame.text = (
        "Synthetic revenue operations forecast pack for real-corpus preservation evidence."
    )

    forecast_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        forecast_slide,
        "Forecast Commit by Segment",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    forecast = forecast_slide.shapes.add_table(
        6,
        5,
        Inches(0.45),
        Inches(1.0),
        Inches(9.1),
        Inches(3.0),
    ).table
    for row_idx, row in enumerate(
        [
            ["Segment", "Pipeline", "Commit", "Best case", "Risk"],
            ["Enterprise", "$42.0m", "$31.5m", "$37.2m", "Procurement"],
            ["Mid-market", "$18.4m", "$12.7m", "$15.8m", "Champion gap"],
            ["Strategic", "$55.1m", "$39.0m", "$47.6m", "Security review"],
            ["Channel", "$9.6m", "$5.1m", "$7.4m", "Partner timing"],
            ["Total", "$125.1m", "$88.3m", "$108.0m", "Coverage 3.1x"],
        ]
    ):
        for col_idx, value in enumerate(row):
            forecast.cell(row_idx, col_idx).text = value
    forecast_slide.notes_slide.notes_text_frame.text = (
        "Forecast table should preserve commit, best-case, and risk text exactly."
    )

    trend_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        trend_slide,
        "Pipeline and Commit Trend",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    chart_data = CategoryChartData()
    chart_data.categories = ["Jan", "Feb", "Mar", "Apr", "May"]
    chart_data.add_series("Pipeline", (84, 92, 101, 113, 125))
    chart_data.add_series("Commit", (61, 67, 74, 82, 88))
    trend_slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.7),
        Inches(1.0),
        Inches(8.4),
        Inches(4.0),
        chart_data,
    )
    trend_slide.notes_slide.notes_text_frame.text = (
        "Line chart workbook and forecast-trend notes must survive round-trip."
    )

    coverage_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        coverage_slide,
        "Coverage and Forecast Risk",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    coverage = coverage_slide.shapes.add_table(
        5,
        4,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(2.45),
    ).table
    for row_idx, row in enumerate(
        [
            ["Region", "Coverage", "Slippage", "Action"],
            ["North America", "3.4x", "$4.1m", "Legal unblock"],
            ["EMEA", "2.8x", "$2.7m", "Executive mapping"],
            ["APAC", "2.2x", "$1.6m", "Partner assist"],
            ["LATAM", "3.0x", "$0.9m", "Pricing approval"],
        ]
    ):
        for col_idx, value in enumerate(row):
            coverage.cell(row_idx, col_idx).text = value
    coverage_slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(1.0),
        Inches(3.75),
        Inches(8.4),
        Inches(3.75),
    )
    coverage_slide.notes_slide.notes_text_frame.text = (
        "Coverage table and connector line should be preserved."
    )

    actions_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        actions_slide,
        "Forecast Actions and Owners",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    group = actions_slide.shapes.add_group_shape()
    for idx, (label, owner) in enumerate(
        (
            ("Deal desk escalation", "RevOps"),
            ("Security fast track", "CISO"),
            ("Partner pull-forward", "Channels"),
            ("Forecast hygiene", "Sales Ops"),
        )
    ):
        tile = group.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.25 + (idx % 2) * 2.85),
            Inches(0.25 + (idx // 2) * 1.05),
            Inches(2.5),
            Inches(0.78),
        )
        tile.text = f"{label}\n{owner}"
    actions_slide.notes_slide.notes_text_frame.text = (
        "Grouped revenue-operations action tiles should keep child text and transforms."
    )

    appendix_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        appendix_slide,
        "Appendix: Forecast Evidence Index",
        Inches(0.5),
        Inches(0.35),
        Inches(8.7),
        Inches(0.6),
    )
    appendix = appendix_slide.shapes.add_table(
        5,
        3,
        Inches(0.7),
        Inches(1.0),
        Inches(8.2),
        Inches(2.6),
    ).table
    for row_idx, row in enumerate(
        [
            ["Artifact", "Source", "Owner"],
            ["Pipeline snapshot", "pipeline_snapshot.xlsx", "RevOps"],
            ["Commit bridge", "commit_bridge.xlsx", "Sales Ops"],
            ["Risk register", "forecast_risks.csv", "Revenue Leadership"],
            ["Deal desk queue", "deal_desk_export.csv", "Finance"],
        ]
    ):
        for col_idx, value in enumerate(row):
            appendix.cell(row_idx, col_idx).text = value


def _add_textbox(slide, text: str, left, top, width, height):
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.text = text
    return shape


def _png_1x1() -> bytes:
    return (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
        b"\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde\x00"
        b"\x00\x00\x0cIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02"
        b"\xfeA\xde\xfc\x8d\x00\x00\x00\x00IEND\xaeB`\x82"
    )
