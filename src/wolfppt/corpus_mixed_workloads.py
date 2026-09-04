"""Mixed synthetic workload fixture builders."""

from __future__ import annotations

from io import BytesIO

from .corpus_slide_helpers import add_textbox as _add_textbox
from .corpus_slide_helpers import png_1x1 as _png_1x1


def add_mixed_real_world_deck(prs, slide) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    _add_textbox(
        slide,
        "Quarterly Business Review",
        Inches(0.6),
        Inches(0.4),
        Inches(8.0),
        Inches(0.7),
    )
    _add_textbox(
        slide,
        "Revenue growth, margin bridge, operating risks, and next-quarter actions.",
        Inches(0.8),
        Inches(1.3),
        Inches(7.8),
        Inches(1.0),
    )
    slide.shapes.add_picture(
        BytesIO(_png_1x1()),
        Inches(8.6),
        Inches(0.45),
        Inches(0.45),
        Inches(0.45),
    )
    slide.notes_slide.notes_text_frame.text = (
        "Open with the preservation claim, then walk through the operating metrics."
    )

    table_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        table_slide,
        "Operating Metrics",
        Inches(0.6),
        Inches(0.4),
        Inches(8.0),
        Inches(0.6),
    )
    table_shape = table_slide.shapes.add_table(
        5,
        4,
        Inches(0.6),
        Inches(1.2),
        Inches(8.8),
        Inches(2.6),
    )
    table = table_shape.table
    rows = [
        ["Metric", "Q1", "Q2", "Q3"],
        ["Revenue", "$10.0m", "$12.4m", "$14.1m"],
        ["Gross margin", "54%", "56%", "58%"],
        ["Bookings", "$11.2m", "$13.0m", "$15.5m"],
        ["Retention", "92%", "93%", "95%"],
    ]
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            table.cell(row_idx, col_idx).text = value

    chart_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        chart_slide,
        "Revenue and Margin Trend",
        Inches(0.6),
        Inches(0.4),
        Inches(8.0),
        Inches(0.6),
    )
    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2", "Q3", "Q4"]
    chart_data.add_series("Revenue", (10.0, 12.4, 14.1, 16.2))
    chart_data.add_series("Gross Margin", (5.4, 6.9, 8.2, 9.6))
    chart_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.8),
        Inches(1.1),
        Inches(7.8),
        Inches(4.1),
        chart_data,
    )

    risks_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        risks_slide,
        "Risk Register",
        Inches(0.6),
        Inches(0.4),
        Inches(8.0),
        Inches(0.6),
    )
    for idx, label in enumerate(
        ("Data quality", "Renewal timing", "Vendor concentration")
    ):
        shape = risks_slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8),
            Inches(1.2 + idx * 1.1),
            Inches(4.6),
            Inches(0.7),
        )
        shape.text = label

    actions_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        actions_slide,
        "Next Actions",
        Inches(0.6),
        Inches(0.4),
        Inches(8.0),
        Inches(0.6),
    )
    body = _add_textbox(
        actions_slide,
        "Validate source workbooks",
        Inches(0.8),
        Inches(1.2),
        Inches(7.8),
        Inches(1.4),
    )
    body.text_frame.add_paragraph().text = "Reconcile management adjustments"
    body.text_frame.add_paragraph().text = "Prepare board-ready appendix"
    actions_slide.notes_slide.notes_text_frame.text = (
        "Close with specific owner/action/date follow-through."
    )


def add_mixed_deal_review_deck(prs, slide) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    _add_textbox(
        slide,
        "QoE Deal Review",
        Inches(0.6),
        Inches(0.4),
        Inches(8.0),
        Inches(0.7),
    )
    _add_textbox(
        slide,
        "Adjusted EBITDA bridge, revenue quality, diligence risks, and close plan.",
        Inches(0.8),
        Inches(1.3),
        Inches(7.8),
        Inches(1.0),
    )
    slide.notes_slide.notes_text_frame.text = (
        "Open with scope, source files, and the working capital sensitivity."
    )

    table_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        table_slide,
        "Diligence KPIs",
        Inches(0.6),
        Inches(0.4),
        Inches(8.0),
        Inches(0.6),
    )
    table_shape = table_slide.shapes.add_table(
        5,
        4,
        Inches(0.6),
        Inches(1.2),
        Inches(8.8),
        Inches(2.6),
    )
    table = table_shape.table
    rows = [
        ["Metric", "Q1", "Q2", "Q3"],
        ["Revenue", "$9.8m", "$11.6m", "$13.7m"],
        ["Adj. EBITDA", "$1.6m", "$2.1m", "$2.7m"],
        ["NWC", "$2.0m", "$2.3m", "$2.5m"],
        ["Churn", "5.1%", "4.8%", "4.4%"],
    ]
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            table.cell(row_idx, col_idx).text = value

    chart_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        chart_slide,
        "Revenue Quality Trend",
        Inches(0.6),
        Inches(0.4),
        Inches(8.0),
        Inches(0.6),
    )
    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2", "Q3", "Q4"]
    chart_data.add_series("Revenue", (9.8, 11.6, 13.7, 15.4))
    chart_data.add_series("Gross Margin", (4.9, 6.0, 7.4, 8.3))
    chart_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.8),
        Inches(1.1),
        Inches(7.8),
        Inches(4.1),
        chart_data,
    )

    risks_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        risks_slide,
        "Diligence Risks",
        Inches(0.6),
        Inches(0.4),
        Inches(8.0),
        Inches(0.6),
    )
    for idx, label in enumerate(
        ("Customer concentration", "Deferred revenue cut-off", "Quality of data")
    ):
        shape = risks_slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8),
            Inches(1.2 + idx * 1.1),
            Inches(4.6),
            Inches(0.7),
        )
        shape.text = label

    actions_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(
        actions_slide,
        "Close Plan",
        Inches(0.6),
        Inches(0.4),
        Inches(8.0),
        Inches(0.6),
    )
    body = _add_textbox(
        actions_slide,
        "Validate seller trial balance",
        Inches(0.8),
        Inches(1.2),
        Inches(7.8),
        Inches(1.4),
    )
    body.text_frame.add_paragraph().text = "Tie workpapers to source schedules"
    body.text_frame.add_paragraph().text = "Prepare diligence committee appendix"
    actions_slide.notes_slide.notes_text_frame.text = (
        "Close with the buyer diligence owner and final data request dates."
    )
