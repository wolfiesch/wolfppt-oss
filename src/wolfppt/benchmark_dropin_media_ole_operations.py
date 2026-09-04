"""OLE media drop-in benchmark operation implementations."""

from __future__ import annotations

from pathlib import Path
from time import perf_counter
from typing import Any

from .benchmark_dropin_media_actions import (
    _apply_python_pptx_dropin_add_deeper_nested_group_ole_object,
    _apply_python_pptx_dropin_add_group_ole_object,
    _apply_python_pptx_dropin_add_nested_group_ole_object,
    _apply_python_pptx_dropin_add_ole_object,
    _apply_python_pptx_dropin_add_ole_object_file_like,
    _apply_wolfppt_dropin_add_deeper_nested_group_ole_object,
    _apply_wolfppt_dropin_add_group_ole_object,
    _apply_wolfppt_dropin_add_nested_group_ole_object,
    _apply_wolfppt_dropin_add_ole_object,
    _apply_wolfppt_dropin_add_ole_object_file_like,
)
from .benchmark_dropin_media_sources import _dropin_add_ole_object_sources
from .benchmark_dropin_shape_details import (
    _dropin_add_deeper_nested_group_ole_object_details,
    _dropin_add_group_ole_object_details,
    _dropin_add_nested_group_ole_object_details,
    _dropin_add_ole_object_details,
    _dropin_add_ole_object_file_like_details,
)
from .benchmark_runtime import elapsed_ms as _elapsed_ms, stamp as _stamp
from .presentation import Presentation as WolfPresentation


def _bench_python_pptx_dropin_add_ole_object(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    embedded_path, icon_path = _dropin_add_ole_object_sources(
        fixture_id,
        fixture_path,
        tmp_path,
    )
    out = tmp_path / (
        f"python-pptx-dropin-add-ole-object-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_ole_object(
        fixture_id,
        prs,
        embedded_path,
        icon_path,
    )
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_ole_object_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_ole_object(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    embedded_path, icon_path = _dropin_add_ole_object_sources(
        fixture_id,
        fixture_path,
        tmp_path,
    )
    out = tmp_path / (
        f"wolfppt-dropin-add-ole-object-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_ole_object(
        fixture_id,
        prs,
        embedded_path,
        icon_path,
    )
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_ole_object_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_group_ole_object(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    embedded_path, icon_path = _dropin_add_ole_object_sources(
        fixture_id,
        fixture_path,
        tmp_path,
    )
    out = tmp_path / (
        f"python-pptx-dropin-add-group-ole-object-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_group_ole_object(
        fixture_id,
        prs,
        embedded_path,
        icon_path,
    )
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_group_ole_object_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_group_ole_object(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    embedded_path, icon_path = _dropin_add_ole_object_sources(
        fixture_id,
        fixture_path,
        tmp_path,
    )
    out = tmp_path / (
        f"wolfppt-dropin-add-group-ole-object-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_group_ole_object(
        fixture_id,
        prs,
        embedded_path,
        icon_path,
    )
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_group_ole_object_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_nested_group_ole_object(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    embedded_path, icon_path = _dropin_add_ole_object_sources(
        fixture_id,
        fixture_path,
        tmp_path,
    )
    out = tmp_path / (
        f"python-pptx-dropin-add-nested-group-ole-object-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_nested_group_ole_object(
        fixture_id,
        prs,
        embedded_path,
        icon_path,
    )
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_nested_group_ole_object_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_nested_group_ole_object(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    embedded_path, icon_path = _dropin_add_ole_object_sources(
        fixture_id,
        fixture_path,
        tmp_path,
    )
    out = tmp_path / (
        f"wolfppt-dropin-add-nested-group-ole-object-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_nested_group_ole_object(
        fixture_id,
        prs,
        embedded_path,
        icon_path,
    )
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_nested_group_ole_object_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_deeper_nested_group_ole_object(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    embedded_path, icon_path = _dropin_add_ole_object_sources(
        fixture_id,
        fixture_path,
        tmp_path,
    )
    out = tmp_path / (
        f"python-pptx-dropin-add-deeper-nested-group-ole-object-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_deeper_nested_group_ole_object(
        fixture_id,
        prs,
        embedded_path,
        icon_path,
    )
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_deeper_nested_group_ole_object_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_deeper_nested_group_ole_object(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    embedded_path, icon_path = _dropin_add_ole_object_sources(
        fixture_id,
        fixture_path,
        tmp_path,
    )
    out = tmp_path / (
        f"wolfppt-dropin-add-deeper-nested-group-ole-object-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_deeper_nested_group_ole_object(
        fixture_id,
        prs,
        embedded_path,
        icon_path,
    )
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_deeper_nested_group_ole_object_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_ole_object_file_like(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    embedded_path, icon_path = _dropin_add_ole_object_sources(
        fixture_id,
        fixture_path,
        tmp_path,
    )
    out = tmp_path / (
        f"python-pptx-dropin-add-ole-object-file-like-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_ole_object_file_like(
        fixture_id,
        prs,
        embedded_path,
        icon_path,
    )
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_ole_object_file_like_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_ole_object_file_like(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    embedded_path, icon_path = _dropin_add_ole_object_sources(
        fixture_id,
        fixture_path,
        tmp_path,
    )
    out = tmp_path / (
        f"wolfppt-dropin-add-ole-object-file-like-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_ole_object_file_like(
        fixture_id,
        prs,
        embedded_path,
        icon_path,
    )
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_ole_object_file_like_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )
