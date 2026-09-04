"""Deterministic Phase 1 fixture corpus generation."""

from __future__ import annotations

import json
from dataclasses import asdict
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

from .chart_axis_ids import (
    normalize_chart_axis_ids_in_package as _normalize_chart_axis_ids,
)
from .corpus_board_packs import (
    _add_finance_board_pack,
    _add_investor_update_pack,
    _add_large_real_world_deck,
    _add_management_reporting_deck,
    _add_portfolio_ops_review_deck,
)
from .corpus_customer_success import add_customer_success_review_pack
from .corpus_earnings import add_earnings_board_appendix_pack
from .corpus_enterprise_implementation import add_enterprise_implementation_pack
from .corpus_manifest import FIXTURES, FixtureSpec
from .corpus_injections import (
    inject_chart_style_parts as _inject_chart_style_parts,
    inject_grouped_shape as _inject_grouped_shape,
    inject_legacy_comment as _inject_legacy_comment,
    inject_ole_object as _inject_ole_object,
    inject_slide_timing as _inject_slide_timing,
    inject_slide_transition as _inject_slide_transition,
    write_pptm_from_pptx as _write_pptm_from_pptx,
)
from .corpus_mixed_workloads import (
    add_mixed_deal_review_deck as _add_mixed_deal_review_deck,
    add_mixed_real_world_deck as _add_mixed_real_world_deck,
)
from .corpus_post_merger import add_post_merger_integration_pack
from .corpus_procurement import add_procurement_vendor_risk_pack
from .corpus_product_launch import add_product_launch_readiness_pack
from .corpus_qoe_diligence import add_qoe_diligence_pack
from .corpus_revenue_ops import add_revenue_ops_forecast_pack
from .corpus_security_compliance import add_security_compliance_review_pack
from .corpus_slide_helpers import add_textbox as _add_textbox
from .corpus_slide_helpers import bmp_2x3 as _bmp_2x3
from .corpus_slide_helpers import jpeg_2x3 as _jpeg_2x3
from .corpus_slide_helpers import png_1x1 as _png_1x1
from .corpus_slide_helpers import tiff_2x3 as _tiff_2x3
from .extractor import extract_semantics


FIXTURE_ROOT = Path("fixtures/pptx")
EXPECTED_ROOT = Path("fixtures/expected/semantic_oracles")
MANIFEST_PATH = Path("fixtures/manifest.json")


def generate_corpus(
    fixture_root: Path = FIXTURE_ROOT,
    expected_root: Path = EXPECTED_ROOT,
    manifest_path: Path = MANIFEST_PATH,
) -> list[FixtureSpec]:
    """Generate deterministic minimal PPTX/PPTM fixtures and semantic oracles."""

    for fixture in FIXTURES:
        path = fixture_root / fixture.path
        path.parent.mkdir(parents=True, exist_ok=True)
        _write_fixture(path, fixture.id)

        oracle_path = expected_root / f"{fixture.id}.json"
        oracle_path.parent.mkdir(parents=True, exist_ok=True)
        oracle_path.write_text(extract_semantics(path).to_json() + "\n", encoding="utf-8")

    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(
        json.dumps([asdict(fixture) for fixture in FIXTURES], indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return list(FIXTURES)


def fixture_manifest(path: Path = MANIFEST_PATH) -> list[dict[str, object]]:
    if not path.exists():
        return []
    return json.loads(path.read_text(encoding="utf-8"))


def _write_fixture(path: Path, fixture_id: str) -> None:
    try:
        from pptx import Presentation
        from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches
    except ImportError as exc:  # pragma: no cover - exercised in lean envs
        raise RuntimeError("fixture generation requires python-pptx; run with the dev or baseline extra") from exc

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    if fixture_id == "text_basic/title_body_bullets":
        _add_textbox(slide, "WolfPPT", Inches(0.8), Inches(0.5), Inches(8.0), Inches(0.7))
        body = _add_textbox(slide, "Fast PPTX", Inches(0.9), Inches(1.5), Inches(7.5), Inches(1.2))
        body.text_frame.add_paragraph().text = "Lossless first"
    elif fixture_id == "slides/two_slide_text":
        _add_textbox(slide, "Jump Source", Inches(0.8), Inches(0.5), Inches(8.0), Inches(0.7))
        target_slide = prs.slides.add_slide(prs.slide_layouts[6])
        _add_textbox(
            target_slide,
            "Jump Target",
            Inches(0.8),
            Inches(0.5),
            Inches(8.0),
            Inches(0.7),
        )
    elif fixture_id == "tables/simple_table":
        table_shape = slide.shapes.add_table(2, 2, Inches(1.0), Inches(1.0), Inches(6.0), Inches(1.0))
        table = table_shape.table
        table.cell(0, 0).text = "Metric"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Slides"
        table.cell(1, 1).text = "1"
    elif fixture_id == "workloads/multi_edit_table":
        table_shape = slide.shapes.add_table(6, 6, Inches(0.5), Inches(0.5), Inches(9.0), Inches(4.0))
        table = table_shape.table
        for row_idx in range(6):
            for col_idx in range(6):
                table.cell(row_idx, col_idx).text = f"R{row_idx}C{col_idx}"
    elif fixture_id == "workloads/multi_format_runs":
        shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9.0), Inches(4.0))
        paragraph = shape.text_frame.paragraphs[0]
        for run_idx in range(48):
            run = paragraph.add_run()
            run.text = f"Run {run_idx:02d} "
    elif fixture_id == "workloads/mixed_real_world_deck":
        _add_mixed_real_world_deck(prs, slide)
    elif fixture_id == "workloads/mixed_deal_review_deck":
        _add_mixed_deal_review_deck(prs, slide)
    elif fixture_id == "workloads/large_real_world_deck":
        _add_large_real_world_deck(prs, slide)
    elif fixture_id == "workloads/portfolio_ops_review_deck":
        _add_portfolio_ops_review_deck(prs, slide)
    elif fixture_id == "workloads/management_reporting_deck":
        _add_management_reporting_deck(prs, slide)
    elif fixture_id == "workloads/finance_board_pack":
        _add_finance_board_pack(prs, slide)
    elif fixture_id == "workloads/investor_update_pack":
        _add_investor_update_pack(prs, slide)
    elif fixture_id == "workloads/qoe_diligence_pack":
        add_qoe_diligence_pack(prs, slide)
    elif fixture_id == "workloads/customer_success_review_pack":
        add_customer_success_review_pack(prs, slide)
    elif fixture_id == "workloads/security_compliance_review_pack":
        add_security_compliance_review_pack(prs, slide)
    elif fixture_id == "workloads/revenue_ops_forecast_pack":
        add_revenue_ops_forecast_pack(prs, slide)
    elif fixture_id == "workloads/product_launch_readiness_pack":
        add_product_launch_readiness_pack(prs, slide)
    elif fixture_id == "workloads/enterprise_implementation_pack":
        add_enterprise_implementation_pack(prs, slide)
    elif fixture_id == "workloads/procurement_vendor_risk_pack":
        add_procurement_vendor_risk_pack(prs, slide)
    elif fixture_id == "workloads/post_merger_integration_pack":
        add_post_merger_integration_pack(prs, slide)
    elif fixture_id == "workloads/earnings_board_appendix_pack":
        add_earnings_board_appendix_pack(prs, slide)
    elif fixture_id == "shapes/grouped_shapes":
        pass
    elif fixture_id == "notes/speaker_notes":
        _add_textbox(slide, "Speaker Notes", Inches(0.8), Inches(0.6), Inches(8.0), Inches(0.8))
        _add_textbox(slide, "Narration lives beside the slide.", Inches(0.9), Inches(1.5), Inches(7.5), Inches(0.8))
        slide.notes_slide.notes_text_frame.text = "Remember to mention preservation before rendering."
    elif fixture_id == "side_parts/legacy_comments":
        _add_textbox(slide, "Comment Fixture", Inches(0.8), Inches(0.6), Inches(8.0), Inches(0.8))
    elif fixture_id == "media/png_picture":
        slide.shapes.add_picture(BytesIO(_png_1x1()), Inches(1.0), Inches(1.0), Inches(1.0), Inches(1.0))
    elif fixture_id == "media/image_formats":
        _add_textbox(slide, "Image Format Fixture", Inches(0.5), Inches(0.3), Inches(8.0), Inches(0.5))
        with TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            image_inputs = [
                ("photo.jfif", _jpeg_2x3(), Inches(0.6)),
                ("diagram.bmp", _bmp_2x3(), Inches(2.3)),
                ("scan.tif", _tiff_2x3(), Inches(4.0)),
            ]
            for filename, payload, left in image_inputs:
                image_path = tmp_path / filename
                image_path.write_bytes(payload)
                slide.shapes.add_picture(
                    str(image_path),
                    left,
                    Inches(1.1),
                    Inches(1.2),
                    Inches(1.0),
                )
    elif fixture_id == "charts/bar_chart":
        chart_data = CategoryChartData()
        chart_data.categories = ["Q1", "Q2"]
        chart_data.add_series("Revenue", (10, 14))
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1.0),
            Inches(1.0),
            Inches(6.0),
            Inches(3.5),
            chart_data,
        )
    elif fixture_id == "charts/styled_bar_chart":
        chart_data = CategoryChartData()
        chart_data.categories = ["Q1", "Q2"]
        chart_data.add_series("Revenue", (10, 14))
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1.0),
            Inches(1.0),
            Inches(6.0),
            Inches(3.5),
            chart_data,
        )
    elif fixture_id == "charts/multi_series_chart":
        chart_data = CategoryChartData()
        chart_data.categories = ["Q1", "Q2", "Q3"]
        chart_data.add_series("Revenue", (10, 14, 17))
        chart_data.add_series("Margin", (2, 3, 5))
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1.0),
            Inches(1.0),
            Inches(6.0),
            Inches(3.5),
            chart_data,
        )
    elif fixture_id == "charts/pie_chart":
        chart_data = CategoryChartData()
        chart_data.categories = ["Software", "Services", "Support"]
        chart_data.add_series("Revenue Mix", (45, 35, 20))
        slide.shapes.add_chart(
            XL_CHART_TYPE.PIE,
            Inches(1.0),
            Inches(1.0),
            Inches(6.0),
            Inches(3.5),
            chart_data,
        )
    elif fixture_id == "charts/xy_scatter_chart":
        chart_data = XyChartData()
        series = chart_data.add_series("Baseline")
        series.add_data_point(1.0, 2.0)
        series.add_data_point(2.0, 4.0)
        slide.shapes.add_chart(
            XL_CHART_TYPE.XY_SCATTER,
            Inches(1.0),
            Inches(1.0),
            Inches(6.0),
            Inches(3.5),
            chart_data,
        )
    elif fixture_id == "charts/bubble_chart":
        chart_data = BubbleChartData()
        series = chart_data.add_series("Baseline")
        series.add_data_point(1.0, 2.0, 3.0)
        series.add_data_point(2.0, 4.0, 5.0)
        slide.shapes.add_chart(
            XL_CHART_TYPE.BUBBLE,
            Inches(1.0),
            Inches(1.0),
            Inches(6.0),
            Inches(3.5),
            chart_data,
        )
    elif fixture_id == "package/ole_object":
        _add_textbox(slide, "OLE Fixture", Inches(0.8), Inches(0.6), Inches(8.0), Inches(0.8))
        _add_textbox(
            slide,
            "Preserve embedded object package.",
            Inches(0.9),
            Inches(1.5),
            Inches(7.5),
            Inches(0.8),
        )
    elif fixture_id == "effects/simple_transition":
        _add_textbox(
            slide,
            "Transition Fixture",
            Inches(0.8),
            Inches(0.6),
            Inches(8.0),
            Inches(0.8),
        )
    elif fixture_id == "effects/simple_timing":
        _add_textbox(
            slide,
            "Timing Fixture",
            Inches(0.8),
            Inches(0.6),
            Inches(8.0),
            Inches(0.8),
        )
    elif fixture_id == "package/macro_preservation":
        _add_textbox(slide, "Macro Fixture", Inches(0.8), Inches(0.6), Inches(8.0), Inches(0.8))
        _add_textbox(slide, "Preserve vbaProject.bin", Inches(0.9), Inches(1.5), Inches(7.5), Inches(0.8))
    else:  # pragma: no cover - defensive guard for future fixture additions
        raise ValueError(f"unknown fixture id: {fixture_id}")

    if path.suffix == ".pptm":
        with TemporaryDirectory() as tmp:
            pptx_path = Path(tmp) / "macro_fixture.pptx"
            prs.save(pptx_path)
            _write_pptm_from_pptx(pptx_path, path)
    else:
        prs.save(path)
        if (
            fixture_id.startswith("charts/")
            or fixture_id.startswith("workloads/mixed_")
            or fixture_id == "workloads/large_real_world_deck"
            or fixture_id == "workloads/portfolio_ops_review_deck"
            or fixture_id == "workloads/management_reporting_deck"
            or fixture_id == "workloads/finance_board_pack"
            or fixture_id == "workloads/investor_update_pack"
            or fixture_id == "workloads/qoe_diligence_pack"
            or fixture_id == "workloads/customer_success_review_pack"
            or fixture_id == "workloads/security_compliance_review_pack"
            or fixture_id == "workloads/revenue_ops_forecast_pack"
            or fixture_id == "workloads/product_launch_readiness_pack"
            or fixture_id == "workloads/enterprise_implementation_pack"
            or fixture_id == "workloads/procurement_vendor_risk_pack"
            or fixture_id == "workloads/post_merger_integration_pack"
            or fixture_id == "workloads/earnings_board_appendix_pack"
        ):
            _normalize_chart_axis_ids(path)
            if fixture_id == "charts/styled_bar_chart":
                _inject_chart_style_parts(path)
            if fixture_id in {
                "workloads/portfolio_ops_review_deck",
                "workloads/management_reporting_deck",
                "workloads/finance_board_pack",
                "workloads/investor_update_pack",
                "workloads/qoe_diligence_pack",
                "workloads/customer_success_review_pack",
                "workloads/security_compliance_review_pack",
                "workloads/revenue_ops_forecast_pack",
                "workloads/product_launch_readiness_pack",
                "workloads/enterprise_implementation_pack",
                "workloads/procurement_vendor_risk_pack",
                "workloads/post_merger_integration_pack",
                "workloads/earnings_board_appendix_pack",
            }:
                _inject_legacy_comment(path)
                _inject_ole_object(path)
        elif fixture_id == "side_parts/legacy_comments":
            _inject_legacy_comment(path)
        elif fixture_id == "effects/simple_transition":
            _inject_slide_transition(path)
        elif fixture_id == "effects/simple_timing":
            _inject_slide_timing(path)
        elif fixture_id == "shapes/grouped_shapes":
            _inject_grouped_shape(path)
        elif fixture_id == "package/ole_object":
            _inject_ole_object(path)
