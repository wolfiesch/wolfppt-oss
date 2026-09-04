"""python-pptx facade value normalization helpers."""

from __future__ import annotations

from numbers import Number
from typing import Any

from .facade_chart_values import (
    axis_crosses_value,
    axis_tick_mark_value,
    category_axis_type_value,
    data_label_position_value,
    legend_position_value,
    normalize_axis_crosses,
    normalize_axis_tick_mark,
    normalize_data_label_position,
    normalize_legend_position,
    normalize_tick_label_position,
    tick_label_position_value,
)


LANGUAGE_ID_NONE = "__wolfppt_language_id_none__"
MAX_PARAGRAPH_SPACING_EMU = 20116800
MAX_LINE_SPACING_MULTIPLE = 132.0


def normalize_rgb(value: Any) -> str:
    if isinstance(value, str):
        cleaned = value.strip().lstrip("#").upper()
        if len(cleaned) == 6 and all(char in "0123456789ABCDEF" for char in cleaned):
            return cleaned
        raise ValueError("RGB color strings must use six hexadecimal digits")
    try:
        red, green, blue = value
    except (TypeError, ValueError) as exc:
        raise TypeError("rgb must be an RGBColor, RGB tuple, or six-digit hex string") from exc
    channels = (int(red), int(green), int(blue))
    if any(channel < 0 or channel > 255 for channel in channels):
        raise ValueError("RGB channels must be between 0 and 255")
    return "{:02X}{:02X}{:02X}".format(*channels)


def normalize_paragraph_alignment(value: Any) -> str | None:
    if value is None:
        return None
    try:
        from pptx.enum.text import PP_ALIGN
    except ImportError:
        PP_ALIGN = None  # type: ignore[assignment]

    if isinstance(value, str):
        raw = value.strip()
        if not raw:
            raise ValueError("paragraph alignment must not be empty")
        for xml_value in _paragraph_alignment_xml_values():
            if raw == xml_value:
                return raw
        raw_name = raw.rsplit(".", 1)[-1].upper()
        if PP_ALIGN is not None:
            for member in PP_ALIGN:
                if member.name == raw_name and getattr(member, "xml_value", None):
                    return str(member.xml_value)
        raise ValueError(f"unsupported paragraph alignment: {value!r}")

    if PP_ALIGN is not None:
        for member in PP_ALIGN:
            try:
                matches = value == member or int(value) == int(member)
            except (TypeError, ValueError):
                matches = value == member
            if matches:
                xml_value = getattr(member, "xml_value", None)
                if not xml_value:
                    raise ValueError(f"{member} has no XML representation")
                return str(xml_value)
    raise TypeError("paragraph alignment must be a PP_ALIGN value, XML value, or None")


def coerce_paragraph_level(value: Any) -> int:
    if isinstance(value, bool):
        raise ValueError(
            f"invalid literal for int() with base 10: {str(value)!r}"
        )
    if not isinstance(value, int):
        raise TypeError(f"value must be an integral type, got {type(value)!r}")
    if value < 0 or value > 8:
        raise ValueError(f"value must be in range 0 to 8 inclusive, got {value}")
    return int(value)


def normalize_paragraph_spacing_length(value: Any, name: str) -> int | None:
    if value is None:
        return None
    if not isinstance(value, int):
        raise TypeError(f"value must be an integral type, got {type(value)!r}")
    value = int(value)
    if value < 0 or value > MAX_PARAGRAPH_SPACING_EMU:
        raise ValueError(
            f"value must be in range 0 to {MAX_PARAGRAPH_SPACING_EMU} inclusive, got {value}"
        )
    return value


def normalize_line_spacing(value: Any) -> int | float | None:
    if value is None:
        return None
    if _is_length_value(value):
        spacing = int(value)
        if spacing < 0 or spacing > MAX_PARAGRAPH_SPACING_EMU:
            raise ValueError(
                f"value must be in range 0 to {MAX_PARAGRAPH_SPACING_EMU} inclusive, got {spacing}"
            )
        return spacing
    if not isinstance(value, (int, float)):
        raise TypeError(f"value must be a number, got {type(value)!r}")
    multiple = float(value)
    if multiple < 0 or multiple > MAX_LINE_SPACING_MULTIPLE:
        raise ValueError(
            f"value must be in range 0.0 to {MAX_LINE_SPACING_MULTIPLE:.1f} inclusive, got {value}"
        )
    return multiple



def normalize_line_dash_style(value: Any) -> str | None:
    if value is None:
        return None
    try:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
    except ImportError:
        MSO_LINE_DASH_STYLE = None  # type: ignore[assignment]

    xml_values = _line_dash_style_xml_values()
    if isinstance(value, str):
        raw = value.strip()
        if raw in xml_values:
            return raw
        raw_name = raw.rsplit(".", 1)[-1].upper()
        if MSO_LINE_DASH_STYLE is not None:
            for member in MSO_LINE_DASH_STYLE:
                if member.name == raw_name and getattr(member, "xml_value", None):
                    return str(member.xml_value)
        raise ValueError(f"unsupported line dash style: {value!r}")

    if MSO_LINE_DASH_STYLE is not None:
        for member in MSO_LINE_DASH_STYLE:
            try:
                matches = value == member or int(value) == int(member)
            except (TypeError, ValueError):
                matches = value == member
            if matches:
                xml_value = getattr(member, "xml_value", None)
                if not xml_value:
                    raise ValueError(f"{member} has no XML representation")
                return str(xml_value)
    raise TypeError(
        "line dash style must be an MSO_LINE_DASH_STYLE value, XML value, or None"
    )


def normalize_pattern_type(value: Any) -> str | None:
    if value is None:
        return None
    try:
        from pptx.enum.dml import MSO_PATTERN_TYPE
    except ImportError:
        MSO_PATTERN_TYPE = None  # type: ignore[assignment]

    xml_values = _pattern_type_xml_values()
    if isinstance(value, str):
        raw = value.strip()
        if raw in xml_values:
            return raw
        raw_name = raw.rsplit(".", 1)[-1].upper()
        if MSO_PATTERN_TYPE is not None:
            for member in MSO_PATTERN_TYPE:
                if member.name == raw_name and getattr(member, "xml_value", None):
                    return str(member.xml_value)
        raise ValueError(f"unsupported pattern type: {value!r}")

    if MSO_PATTERN_TYPE is not None:
        for member in MSO_PATTERN_TYPE:
            try:
                matches = value == member or int(value) == int(member)
            except (TypeError, ValueError):
                matches = value == member
            if matches:
                xml_value = getattr(member, "xml_value", None)
                if not xml_value:
                    raise ValueError(f"{member} has no XML representation")
                return str(xml_value)
    raise TypeError(
        "pattern type must be an MSO_PATTERN_TYPE value, XML value, or None"
    )


def normalize_theme_color(value: Any) -> str:
    try:
        from pptx.enum.dml import MSO_THEME_COLOR
    except ImportError:
        MSO_THEME_COLOR = None  # type: ignore[assignment]

    if isinstance(value, str):
        raw = value.strip()
        if raw in _theme_color_xml_values():
            return raw
        raw_name = raw.rsplit(".", 1)[-1].upper()
        if MSO_THEME_COLOR is not None:
            for member in MSO_THEME_COLOR:
                if member.name == raw_name and getattr(member, "xml_value", None):
                    return str(member.xml_value)
        raise ValueError(f"unsupported theme color: {value!r}")

    if MSO_THEME_COLOR is not None:
        for member in MSO_THEME_COLOR:
            try:
                matches = value == member or int(value) == int(member)
            except (TypeError, ValueError):
                matches = value == member
            if matches:
                xml_value = getattr(member, "xml_value", None)
                if not xml_value:
                    raise ValueError(f"{member} has no XML representation")
                return str(xml_value)
    raise TypeError("theme_color must be an MSO_THEME_COLOR value or XML value")


def normalize_language_id(value: Any) -> str | None:
    try:
        from pptx.enum.lang import MSO_LANGUAGE_ID
    except ImportError:
        MSO_LANGUAGE_ID = None  # type: ignore[assignment]

    if value is None:
        return None
    if MSO_LANGUAGE_ID is None:
        if isinstance(value, str):
            return value
        if int(value) == 0:
            return None
        raise TypeError("language_id must be an MSO_LANGUAGE_ID value or None")

    language_id = MSO_LANGUAGE_ID(value)
    if language_id == MSO_LANGUAGE_ID.NONE:
        return None
    return str(language_id.xml_value)


def paragraph_alignment_value(value: str | None) -> Any:
    if value is None:
        return None
    try:
        from pptx.enum.text import PP_ALIGN
    except ImportError:
        return value
    for member in PP_ALIGN:
        if getattr(member, "xml_value", None) == value:
            return member
    return value



def optional_float(value: Any) -> float | None:
    if value is None:
        return None
    return float(value)


def normalize_optional_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, Number):
        raise TypeError(f"value must be a number or None, got {type(value)!r}")
    return float(value)


def line_dash_style_value(value: str | None) -> Any:
    if value is None:
        return None
    try:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
    except ImportError:
        return value
    for member in MSO_LINE_DASH_STYLE:
        if getattr(member, "xml_value", None) == value:
            return member
    return value


def pattern_type_value(value: str | None) -> Any:
    if value is None:
        return None
    try:
        from pptx.enum.dml import MSO_PATTERN_TYPE
    except ImportError:
        return value
    for member in MSO_PATTERN_TYPE:
        if getattr(member, "xml_value", None) == value:
            return member
    return value


def language_id_value(value: str | None) -> Any:
    if value == LANGUAGE_ID_NONE:
        value = None
    try:
        from pptx.enum.lang import MSO_LANGUAGE_ID
    except ImportError:
        return value
    if value is None:
        return MSO_LANGUAGE_ID.NONE
    for member in MSO_LANGUAGE_ID:
        if getattr(member, "xml_value", None) == value:
            return member
    return value


def fill_type_value(value: str | None) -> Any:
    if value is None:
        return None
    try:
        from pptx.enum.dml import MSO_FILL
    except ImportError:
        return value
    mapping = {
        "solid": MSO_FILL.SOLID,
        "background": MSO_FILL.BACKGROUND,
        "patterned": MSO_FILL.PATTERNED,
        "gradient": MSO_FILL.GRADIENT,
        "picture": MSO_FILL.PICTURE,
        "group": MSO_FILL.GROUP,
    }
    return mapping.get(value, value)


def color_type_value(value: str | None) -> Any:
    if value is None:
        return None
    try:
        from pptx.enum.dml import MSO_COLOR_TYPE
    except ImportError:
        return value
    if value == "rgb":
        return MSO_COLOR_TYPE.RGB
    if value == "scheme":
        return MSO_COLOR_TYPE.SCHEME
    return value


def theme_color_value(value: str) -> Any:
    try:
        from pptx.enum.dml import MSO_THEME_COLOR
    except ImportError:
        return value
    if value == "":
        return MSO_THEME_COLOR.NOT_THEME_COLOR
    for member in MSO_THEME_COLOR:
        if getattr(member, "xml_value", None) == value:
            return member
    return value


def normalize_text_frame_word_wrap(value: Any) -> bool | None:
    if value is None:
        return None
    if value in (True, False):
        return bool(value)
    raise ValueError(f"assigned value must be True, False, or None, got {value}")


def normalize_text_frame_vertical_anchor(value: Any) -> str | None:
    if value is None:
        return None
    try:
        from pptx.enum.text import MSO_ANCHOR
    except ImportError:
        MSO_ANCHOR = None  # type: ignore[assignment]

    if MSO_ANCHOR is not None:
        for member in MSO_ANCHOR:
            try:
                matches = value == member or int(value) == int(member)
            except (TypeError, ValueError):
                matches = value == member
            if matches:
                xml_value = getattr(member, "xml_value", None)
                if xml_value:
                    return str(xml_value)
    raise ValueError(f"{value!r} is not a valid MSO_VERTICAL_ANCHOR")


def text_frame_vertical_anchor_value(value: str | None) -> Any:
    if value is None:
        return None
    try:
        from pptx.enum.text import MSO_ANCHOR
    except ImportError:
        return value
    for member in MSO_ANCHOR:
        if getattr(member, "xml_value", None) == value:
            return member
    return value


def normalize_text_frame_auto_size(value: Any) -> str | None:
    if value is None:
        return None
    try:
        from pptx.enum.text import MSO_AUTO_SIZE
    except ImportError:
        MSO_AUTO_SIZE = None  # type: ignore[assignment]

    if isinstance(value, str):
        numeric = None
    elif isinstance(value, float) and not value.is_integer():
        numeric = None
    else:
        try:
            numeric = int(value)
        except (TypeError, ValueError):
            numeric = None
    if MSO_AUTO_SIZE is not None and numeric is not None:
        if numeric == int(MSO_AUTO_SIZE.NONE):
            return "noAutofit"
        if numeric == int(MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT):
            return "spAutoFit"
        if numeric == int(MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE):
            return "normAutofit"
        if numeric == int(MSO_AUTO_SIZE.MIXED):
            return None
    raise ValueError(f"{value!r} is not a valid MSO_AUTO_SIZE")


def text_frame_auto_size_value(value: str | None) -> Any:
    if value is None:
        return None
    try:
        from pptx.enum.text import MSO_AUTO_SIZE
    except ImportError:
        return value
    mapping = {
        "noAutofit": MSO_AUTO_SIZE.NONE,
        "spAutoFit": MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT,
        "normAutofit": MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE,
    }
    return mapping.get(value, value)


def rgb_value(rgb: str) -> Any:
    try:
        from pptx.dml.color import RGBColor
    except ImportError:  # pragma: no cover - lean runtime fallback
        return rgb
    return RGBColor.from_string(rgb)


def centipoints_value(value: int) -> Any:
    try:
        from pptx.util import Centipoints
    except ImportError:  # pragma: no cover - lean runtime fallback
        return centipoints_to_emu(value)
    return Centipoints(value)


def emu_to_centipoints(value: int) -> int:
    return int(round(value / 127))


def centipoints_to_emu(value: int) -> int:
    return int(value) * 127


def paragraph_spacing_xml_tag(attr: str) -> str:
    mapping = {
        "line_spacing": "lnSpc",
        "space_before": "spcBef",
        "space_after": "spcAft",
    }
    return mapping[attr]


def emu_value(value: int) -> Any:
    try:
        from pptx.util import Emu
    except ImportError:  # pragma: no cover - lean runtime fallback
        return value
    return Emu(value)


def coerce_emu(value: Any, name: str) -> int:
    try:
        return int(value)
    except (TypeError, ValueError) as exc:
        raise TypeError(f"{name} must be an integer EMU value") from exc


def coerce_crop_fraction(value: Any) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise TypeError(f"value must be a number, got {type(value)!r}")
    return float(value)


def coerce_rotation_degrees(value: Any) -> float:
    if not isinstance(value, (int, float)):
        raise TypeError(f"value must be a number, got {type(value)!r}")
    units = int(round(float(value) * 60000)) % (360 * 60000)
    return rotation_degrees_from_units(units)


def rotation_openxml_units(value: float) -> int:
    return int(round(float(value) * 60000)) % (360 * 60000)


def rotation_degrees_from_units(value: int) -> float:
    return (int(value) % (360 * 60000)) / 60000.0


def coerce_positive_int(value: Any, name: str) -> int:
    coerced = coerce_emu(value, name)
    if coerced < 1:
        raise ValueError(f"{name} must be greater than zero")
    return coerced


def _is_length_value(value: Any) -> bool:
    try:
        from pptx.util import Length
    except ImportError:
        return isinstance(value, int) and not isinstance(value, bool) and abs(value) > 132
    return isinstance(value, Length) or (
        isinstance(value, int) and not isinstance(value, bool) and abs(value) > 132
    )


def _axis_crosses_member_xml(member: Any) -> str | None:
    xml_value = getattr(member, "xml_value", None)
    return None if not xml_value else str(xml_value)


def _paragraph_alignment_xml_values() -> set[str]:
    return {"l", "ctr", "r", "just", "dist", "thaiDist", "justLow"}


def _line_dash_style_xml_values() -> set[str]:
    return {
        "solid",
        "sysDash",
        "sysDot",
        "dash",
        "dashDot",
        "lgDashDotDot",
        "lgDash",
        "lgDashDot",
    }


def _pattern_type_xml_values() -> set[str]:
    try:
        from pptx.enum.dml import MSO_PATTERN_TYPE
    except ImportError:
        return {"divot", "pct5", "pct10", "pct20", "dashHorz", "dashVert"}
    return {
        str(member.xml_value)
        for member in MSO_PATTERN_TYPE
        if getattr(member, "xml_value", None)
    }


def _theme_color_xml_values() -> set[str]:
    try:
        from pptx.enum.dml import MSO_THEME_COLOR
    except ImportError:
        return {"accent1", "accent2", "bg1", "dk1", "lt1", "tx1"}
    return {
        str(member.xml_value)
        for member in MSO_THEME_COLOR
        if getattr(member, "xml_value", None)
    }
