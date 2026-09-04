"""Media drop-in benchmark operation implementations."""

# ruff: noqa: F401

from __future__ import annotations

from pathlib import Path
from time import perf_counter
from typing import Any

from .benchmark_dropin_media_actions import (
    _apply_python_pptx_dropin_add_deeper_nested_group_picture,
    _apply_python_pptx_dropin_add_group_picture,
    _apply_python_pptx_dropin_add_movie,
    _apply_python_pptx_dropin_add_movie_file_like,
    _apply_python_pptx_dropin_add_nested_group_picture,
    _apply_python_pptx_dropin_add_picture,
    _apply_python_pptx_dropin_add_picture_auto_size,
    _apply_python_pptx_dropin_add_picture_file_like,
    _apply_python_pptx_dropin_replace_picture,
    _read_python_pptx_picture_image_metadata,
    _read_wolfppt_picture_image_metadata,
    _apply_wolfppt_dropin_add_deeper_nested_group_picture,
    _apply_wolfppt_dropin_add_group_picture,
    _apply_wolfppt_dropin_add_movie,
    _apply_wolfppt_dropin_add_movie_file_like,
    _apply_wolfppt_dropin_add_nested_group_picture,
    _apply_wolfppt_dropin_add_picture,
    _apply_wolfppt_dropin_add_picture_auto_size,
    _apply_wolfppt_dropin_add_picture_file_like,
    _apply_wolfppt_dropin_replace_picture,
)
from .benchmark_dropin_media_ole_operations import (
    _bench_python_pptx_dropin_add_deeper_nested_group_ole_object,
    _bench_python_pptx_dropin_add_group_ole_object,
    _bench_python_pptx_dropin_add_nested_group_ole_object,
    _bench_python_pptx_dropin_add_ole_object,
    _bench_python_pptx_dropin_add_ole_object_file_like,
    _bench_wolfppt_facade_dropin_add_deeper_nested_group_ole_object,
    _bench_wolfppt_facade_dropin_add_group_ole_object,
    _bench_wolfppt_facade_dropin_add_nested_group_ole_object,
    _bench_wolfppt_facade_dropin_add_ole_object,
    _bench_wolfppt_facade_dropin_add_ole_object_file_like,
)
from .benchmark_dropin_media_sources import (
    _dropin_add_movie_sources,
    _dropin_add_picture_source_image,
    _dropin_replace_picture_source_image,
)
from .benchmark_dropin_shape_details import (
    _dropin_add_deeper_nested_group_picture_details,
    _dropin_add_group_picture_details,
    _dropin_add_nested_group_picture_details,
    _dropin_add_movie_details,
    _dropin_add_movie_file_like_details,
    _dropin_add_picture_auto_size_details,
    _dropin_add_picture_details,
    _dropin_add_picture_file_like_details,
    _dropin_picture_image_inspection_details,
    _dropin_replace_picture_details,
)
from .benchmark_runtime import elapsed_ms as _elapsed_ms, stamp as _stamp
from .presentation import Presentation as WolfPresentation


def _bench_python_pptx_dropin_add_picture_auto_size(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    image_path = _dropin_add_picture_source_image(fixture_id, fixture_path, tmp_path)
    out = tmp_path / (
        f"python-pptx-dropin-add-picture-auto-size-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_picture_auto_size(fixture_id, prs, image_path)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_picture_auto_size_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_picture_auto_size(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    image_path = _dropin_add_picture_source_image(fixture_id, fixture_path, tmp_path)
    out = tmp_path / (
        f"wolfppt-dropin-add-picture-auto-size-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_picture_auto_size(fixture_id, prs, image_path)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_picture_auto_size_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_picture(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    image_path = _dropin_add_picture_source_image(fixture_id, fixture_path, tmp_path)
    out = tmp_path / (
        f"python-pptx-dropin-add-picture-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_picture(fixture_id, prs, image_path)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_picture_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_picture(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    image_path = _dropin_add_picture_source_image(fixture_id, fixture_path, tmp_path)
    out = tmp_path / (
        f"wolfppt-dropin-add-picture-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_picture(fixture_id, prs, image_path)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_picture_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_picture_image_inspection(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    metadata = _read_python_pptx_picture_image_metadata(fixture_id, prs)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_picture_image_inspection_details(
        fixture_id,
        fixture_path,
        metadata,
    )


def _bench_wolfppt_facade_dropin_picture_image_inspection(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    metadata = _read_wolfppt_picture_image_metadata(fixture_id, prs)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_picture_image_inspection_details(
        fixture_id,
        fixture_path,
        metadata,
    )


def _bench_python_pptx_dropin_add_group_picture(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    image_path = _dropin_add_picture_source_image(fixture_id, fixture_path, tmp_path)
    out = tmp_path / (
        f"python-pptx-dropin-add-group-picture-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_group_picture(fixture_id, prs, image_path)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_group_picture_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_group_picture(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    image_path = _dropin_add_picture_source_image(fixture_id, fixture_path, tmp_path)
    out = tmp_path / (
        f"wolfppt-dropin-add-group-picture-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_group_picture(fixture_id, prs, image_path)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_group_picture_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_nested_group_picture(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    image_path = _dropin_add_picture_source_image(fixture_id, fixture_path, tmp_path)
    out = tmp_path / (
        f"python-pptx-dropin-add-nested-group-picture-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_nested_group_picture(fixture_id, prs, image_path)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_nested_group_picture_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_nested_group_picture(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    image_path = _dropin_add_picture_source_image(fixture_id, fixture_path, tmp_path)
    out = tmp_path / (
        f"wolfppt-dropin-add-nested-group-picture-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_nested_group_picture(fixture_id, prs, image_path)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_nested_group_picture_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_deeper_nested_group_picture(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    image_path = _dropin_add_picture_source_image(fixture_id, fixture_path, tmp_path)
    out = tmp_path / (
        f"python-pptx-dropin-add-deeper-nested-group-picture-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_deeper_nested_group_picture(
        fixture_id,
        prs,
        image_path,
    )
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_deeper_nested_group_picture_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_deeper_nested_group_picture(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    image_path = _dropin_add_picture_source_image(fixture_id, fixture_path, tmp_path)
    out = tmp_path / (
        f"wolfppt-dropin-add-deeper-nested-group-picture-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_deeper_nested_group_picture(
        fixture_id,
        prs,
        image_path,
    )
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_deeper_nested_group_picture_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_picture_file_like(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    image_path = _dropin_add_picture_source_image(fixture_id, fixture_path, tmp_path)
    out = tmp_path / (
        f"python-pptx-dropin-add-picture-file-like-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_picture_file_like(fixture_id, prs, image_path)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_picture_file_like_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_picture_file_like(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    image_path = _dropin_add_picture_source_image(fixture_id, fixture_path, tmp_path)
    out = tmp_path / (
        f"wolfppt-dropin-add-picture-file-like-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_picture_file_like(fixture_id, prs, image_path)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_picture_file_like_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_replace_picture(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    image_path = _dropin_replace_picture_source_image(
        fixture_id,
        fixture_path,
        tmp_path,
    )
    out = tmp_path / (
        f"python-pptx-dropin-replace-picture-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_replace_picture(fixture_id, prs, image_path)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_replace_picture_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_replace_picture(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    image_path = _dropin_replace_picture_source_image(
        fixture_id,
        fixture_path,
        tmp_path,
    )
    out = tmp_path / (
        f"wolfppt-dropin-replace-picture-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_replace_picture(fixture_id, prs, image_path)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_replace_picture_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_movie(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    movie_path, poster_path = _dropin_add_movie_sources(
        fixture_id, fixture_path, tmp_path
    )
    out = tmp_path / (
        f"python-pptx-dropin-add-movie-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_movie(fixture_id, prs, movie_path, poster_path)
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_movie_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_wolfppt_facade_dropin_add_movie(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    movie_path, poster_path = _dropin_add_movie_sources(
        fixture_id, fixture_path, tmp_path
    )
    out = tmp_path / (
        f"wolfppt-dropin-add-movie-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_movie(fixture_id, prs, movie_path, poster_path)
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_movie_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )


def _bench_python_pptx_dropin_add_movie_file_like(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    from pptx import Presentation as PythonPptxPresentation

    movie_path, poster_path = _dropin_add_movie_sources(
        fixture_id, fixture_path, tmp_path
    )
    out = tmp_path / (
        f"python-pptx-dropin-add-movie-file-like-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = PythonPptxPresentation(str(fixture_path))
    _apply_python_pptx_dropin_add_movie_file_like(
        fixture_id,
        prs,
        movie_path,
        poster_path,
    )
    prs.save(str(out))
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_movie_file_like_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )

def _bench_wolfppt_facade_dropin_add_movie_file_like(
    fixture_id: str,
    fixture_path: Path,
    expected_path: Path,
    tmp_path: Path,
    validate_openxml: bool,
) -> tuple[float, dict[str, Any]]:
    movie_path, poster_path = _dropin_add_movie_sources(
        fixture_id, fixture_path, tmp_path
    )
    out = tmp_path / (
        f"wolfppt-dropin-add-movie-file-like-{fixture_path.stem}-{_stamp()}{fixture_path.suffix}"
    )
    start = perf_counter()
    prs = WolfPresentation(fixture_path)
    _apply_wolfppt_dropin_add_movie_file_like(
        fixture_id,
        prs,
        movie_path,
        poster_path,
    )
    prs.save(out)
    elapsed_ms = _elapsed_ms(start)
    return elapsed_ms, _dropin_add_movie_file_like_details(
        fixture_id,
        fixture_path,
        out,
        validate_openxml,
    )
