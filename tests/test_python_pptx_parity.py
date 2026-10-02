"""Edits that must land where python-pptx lands them."""

from __future__ import annotations

import io
import re
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree as ET

import pptx
import pytest
from pptx.chart.data import BubbleChartData, CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

import wolfppt

SS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def _run_texts(paragraph) -> list[str]:
    return [run.text for run in paragraph.runs]


def _canonical(value: str) -> str:
    try:
        return str(Decimal(value).normalize())
    except InvalidOperation:
        return value


@pytest.fixture
def mixed_run_deck(tmp_path: Path) -> tuple[Path, int]:
    """A title paragraph shaped like many real decks: an empty first run, a
    line break, text runs, then a slide-number field, plus a second paragraph."""
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.add_run().text = ""
    paragraph.add_line_break()
    paragraph.add_run().text = "Title"
    paragraph.add_run().text = " tail"
    field = paragraph._p.makeelement(
        qn("a:fld"),
        {"id": "{6C5D3F1E-0000-4000-8000-000000000001}", "type": "slidenum"},
    )
    field.append(field.makeelement(qn("a:t"), {}))
    field[0].text = "7"
    paragraph._p.append(field)
    box.text_frame.add_paragraph().add_run().text = "Second"
    path = tmp_path / "runs.pptx"
    prs.save(path)
    return path, box.shape_id


def test_run_reads_match_python_pptx(mixed_run_deck: tuple[Path, int]) -> None:
    path, shape_id = mixed_run_deck
    reference = pptx.Presentation(path).slides[0].shapes[0].text_frame
    shape = next(
        s for s in wolfppt.Presentation(path).slides[0].shapes if s.shape_id == shape_id
    )

    assert [_run_texts(p) for p in shape.text_frame.paragraphs] == [
        _run_texts(p) for p in reference.paragraphs
    ]
    assert shape.text_frame.text == reference.text


@pytest.mark.parametrize("run_index", [0, 1, 2])
def test_run_edit_lands_on_the_python_pptx_run(
    mixed_run_deck: tuple[Path, int], tmp_path: Path, run_index: int
) -> None:
    path, shape_id = mixed_run_deck
    presentation = wolfppt.Presentation(path)
    shape = next(s for s in presentation.slides[0].shapes if s.shape_id == shape_id)
    shape.text_frame.paragraphs[0].runs[run_index].text = "EDIT"
    output = tmp_path / "edited.pptx"
    presentation.save(output)

    expected = pptx.Presentation(path)
    expected.slides[0].shapes[0].text_frame.paragraphs[0].runs[run_index].text = "EDIT"
    expected_frame = expected.slides[0].shapes[0].text_frame
    actual_frame = pptx.Presentation(output).slides[0].shapes[0].text_frame

    assert [_run_texts(p) for p in actual_frame.paragraphs] == [
        _run_texts(p) for p in expected_frame.paragraphs
    ]
    assert actual_frame.text == expected_frame.text


def _embedded_workbook(package: zipfile.ZipFile, chart_index: int) -> zipfile.ZipFile:
    names = sorted(n for n in package.namelist() if n.startswith("ppt/embeddings/"))
    return zipfile.ZipFile(io.BytesIO(package.read(names[chart_index])))


def _workbook_cells(workbook: zipfile.ZipFile) -> dict[str, str]:
    shared = [
        "".join(t.text or "" for t in si.iter(f"{SS}t"))
        for si in ET.fromstring(workbook.read("xl/sharedStrings.xml")).iter(f"{SS}si")
    ]
    cells = {}
    for cell in ET.fromstring(workbook.read("xl/worksheets/sheet1.xml")).iter(f"{SS}c"):
        value = cell.findtext(f"{SS}v") or ""
        cells[cell.get("r")] = shared[int(value)] if cell.get("t") == "s" else value
    return cells


def _chart_formula_values(chart_xml: bytes) -> dict[str, str]:
    """Map each cell a chart formula references to the chart's cached value."""
    values = {}
    for reference in ET.fromstring(chart_xml).iter():
        if reference.tag not in {f"{C}strRef", f"{C}numRef"}:
            continue
        match = re.fullmatch(
            r"Sheet1!\$([A-Z]+)\$(\d+)(?::\$[A-Z]+\$(\d+))?",
            reference.findtext(f"{C}f"),
        )
        assert match is not None
        column, first, last = (
            match.group(1),
            int(match.group(2)),
            int(match.group(3) or match.group(2)),
        )
        points = {
            int(pt.get("idx")): pt.findtext(f"{C}v") for pt in reference.iter(f"{C}pt")
        }
        for offset, row in enumerate(range(first, last + 1)):
            values[f"{column}{row}"] = points[offset]
    return values


def test_chart_replace_data_updates_embedded_workbook_in_place(tmp_path: Path) -> None:
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    category = CategoryChartData()
    category.categories = ["North", "South", "West"]
    category.add_series("Revenue", (10, 20, 30))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(4), Inches(3), category
    )
    bubble = BubbleChartData()
    series = bubble.add_series("Pipeline")
    for point in ((1, 6, 4), (2, 9, 8), (3, 12, 12)):
        series.add_data_point(*point)
    slide.shapes.add_chart(
        XL_CHART_TYPE.BUBBLE, 0, Inches(3), Inches(4), Inches(3), bubble
    )
    source = tmp_path / "charts.pptx"
    prs.save(source)

    new_category = CategoryChartData()
    new_category.categories = ["East", "42", "84"]
    new_category.add_series("Revenue", (42, 84.5, 126))
    new_bubble = BubbleChartData()
    new_series = new_bubble.add_series("Pipeline")
    for point in ((5, 13, 21), (34, 55, 89), (8, 13, 21)):
        new_series.add_data_point(*point)

    presentation = wolfppt.Presentation(source)
    presentation.slides[0].shapes[0].chart.replace_data(new_category)
    presentation.slides[0].shapes[1].chart.replace_data(new_bubble)
    output = tmp_path / "replaced.pptx"
    presentation.save(output)

    expected = pptx.Presentation(source)
    expected.slides[0].shapes[0].chart.replace_data(new_category)
    expected.slides[0].shapes[1].chart.replace_data(new_bubble)
    expected_path = tmp_path / "expected.pptx"
    expected.save(expected_path)

    with (
        zipfile.ZipFile(source) as before,
        zipfile.ZipFile(output) as after,
        zipfile.ZipFile(expected_path) as reference,
    ):
        charts = sorted(
            n for n in after.namelist() if re.fullmatch(r"ppt/charts/chart\d+\.xml", n)
        )
        for index, chart in enumerate(charts):
            original = _embedded_workbook(before, index)
            updated = _embedded_workbook(after, index)
            cells = _workbook_cells(updated)
            # The workbook holds exactly what python-pptx writes, and every
            # cell a chart formula reads agrees with the chart's cache.
            assert cells == _workbook_cells(_embedded_workbook(reference, index))
            for coordinate, value in _chart_formula_values(after.read(chart)).items():
                assert _canonical(cells[coordinate]) == _canonical(value)
            # Everything but the cell values survives: parts, theme, styles,
            # column widths and sheet view.
            assert updated.namelist() == original.namelist()
            for name in ("xl/styles.xml", "xl/theme/theme1.xml", "xl/workbook.xml"):
                assert updated.read(name) == original.read(name)
            strip = re.compile(rb"<v>[^<]*</v>")
            assert strip.sub(
                b"", updated.read("xl/worksheets/sheet1.xml")
            ) == strip.sub(b"", original.read("xl/worksheets/sheet1.xml"))


def test_chart_replace_data_with_less_data_clears_stale_workbook_cells(
    tmp_path: Path,
) -> None:
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    category = CategoryChartData()
    category.categories = ["North", "South", "West"]
    category.add_series("Revenue", (10, 20, 30))
    category.add_series("Cost", (1, 2, 3))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(4), Inches(3), category
    )
    source = tmp_path / "chart.pptx"
    prs.save(source)

    smaller = CategoryChartData()
    smaller.categories = ["East", "West"]
    smaller.add_series("Revenue", (5, 6))

    presentation = wolfppt.Presentation(source)
    presentation.slides[0].shapes[0].chart.replace_data(smaller)
    output = tmp_path / "replaced.pptx"
    presentation.save(output)

    expected = pptx.Presentation(source)
    expected.slides[0].shapes[0].chart.replace_data(smaller)
    expected_path = tmp_path / "expected.pptx"
    expected.save(expected_path)

    with (
        zipfile.ZipFile(source) as before,
        zipfile.ZipFile(output) as after,
        zipfile.ZipFile(expected_path) as reference,
    ):
        updated = _embedded_workbook(after, 0)
        cells = {ref: value for ref, value in _workbook_cells(updated).items() if value}
        # The dropped category row and series column hold no stale values.
        assert cells == _workbook_cells(_embedded_workbook(reference, 0))
        original = _embedded_workbook(before, 0)
        for name in ("xl/styles.xml", "xl/theme/theme1.xml", "xl/workbook.xml"):
            assert updated.read(name) == original.read(name)
